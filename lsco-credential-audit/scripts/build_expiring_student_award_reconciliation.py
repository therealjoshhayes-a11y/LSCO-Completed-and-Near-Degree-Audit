from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path.cwd()

TARGET_CATALOG = "2021-2022"
TARGET_START_YEAR = 2021

REPORTING = (
    ROOT
    / "data"
    / "processed"
    / "reporting"
    / "additional_awards_clean_query_20260803"
)

# COMPLETE + catalog-eligible alternatives, before one lineage version is selected.
ELIGIBLE = REPORTING / "05_all_eligible_catalog_alternatives.csv"

# Official Banner awards successfully normalized to our canonical credential lineages.
OFFICIAL = REPORTING / "08_official_awards_normalized.csv"

# Student × modeled lineage pairs already represented by an official award,
# including governed parent/child relationships.
REPRESENTATION = REPORTING / "09_official_award_representation_pairs.csv"

IDENTITY = (
    ROOT
    / "data"
    / "raw"
    / "student_exports"
    / "banner_course_history_actual.csv"
)

OUTDIR = (
    ROOT
    / "data"
    / "processed"
    / "reporting"
    / "expiring_2021_2022_student_reconciliation_20260811"
)

OUTFILE = OUTDIR / "2021_2022_EXPIRING_STUDENT_AWARD_RECONCILIATION.xlsx"


def fail(msg: str) -> None:
    raise RuntimeError(msg)


def norm_id(v: object) -> str:
    return str(v).strip()


def norm_text(v: object) -> str:
    if pd.isna(v):
        return ""
    return re.sub(
        r"\s+",
        " ",
        str(v).strip().upper()
    )


def catalog_start(v: object) -> int:
    m = re.search(r"(20\d{2})", str(v))
    return int(m.group(1)) if m else -1


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def find_col(columns, candidates, required=True, label="column"):
    lookup = {
        str(c).strip().lower(): c
        for c in columns
    }

    for candidate in candidates:
        key = candidate.lower()
        if key in lookup:
            return lookup[key]

    if required:
        fail(
            f"Could not find {label}. "
            f"Tried {candidates}. "
            f"Available columns: {list(columns)}"
        )

    return None


def choose_display_label(row, candidates):
    for col in candidates:
        if col and col in row.index:
            value = str(row[col]).strip()
            if value:
                return value
    return ""


def join_unique(values):
    clean = sorted(
        {
            str(v).strip()
            for v in values
            if str(v).strip()
        }
    )
    return " | ".join(clean)


print("=" * 110)
print("2021-2022 EXPIRING STUDENT AWARD RECONCILIATION")
print("=" * 110)

for path in [ELIGIBLE, OFFICIAL, REPRESENTATION, IDENTITY]:
    if not path.exists():
        fail(f"Required source missing: {path}")

OUTDIR.mkdir(parents=True, exist_ok=True)


# =============================================================================
# 1. FIND EXPIRING 2021-2022 COMPLETIONS
#
# Start from ALL COMPLETE + catalog-eligible alternatives.
#
# Keep student × lineage where:
#   - 2021-2022 is COMPLETE + eligible
#   - NO later catalog is COMPLETE + eligible for that same lineage
#
# Award status plays NO ROLE in this selection.
# =============================================================================

eligible = pd.read_csv(
    ELIGIBLE,
    dtype=str,
    low_memory=False
).fillna("")

student_col = find_col(
    eligible.columns,
    ["student_id"],
    label="student ID"
)

catalog_col = find_col(
    eligible.columns,
    ["catalog_year"],
    label="catalog year"
)

lineage_col = find_col(
    eligible.columns,
    [
        "canonical_lineage",
        "credential_lineage",
        "detected_lineage"
    ],
    label="canonical lineage"
)

title_col = find_col(
    eligible.columns,
    [
        "credential_title",
        "program_title",
        "title"
    ],
    required=False
)

cred_id_col = find_col(
    eligible.columns,
    [
        "credential_id",
        "credential_base_id"
    ],
    required=False
)

if "audit_status" in eligible.columns:
    eligible = eligible[
        eligible["audit_status"]
        .astype(str)
        .str.strip()
        .str.upper()
        .eq("COMPLETE")
    ].copy()

eligible["_student_id"] = eligible[student_col].map(norm_id)
eligible["_lineage"] = eligible[lineage_col].map(norm_text)
eligible["_catalog_start"] = eligible[catalog_col].map(catalog_start)

target = eligible[
    eligible[catalog_col]
    .astype(str)
    .str.strip()
    .eq(TARGET_CATALOG)
].copy()

if target.empty:
    fail("No eligible COMPLETE rows found for 2021-2022.")

target_keys = (
    target[
        ["_student_id", "_lineage"]
    ]
    .drop_duplicates()
)

later_keys = (
    eligible[
        eligible["_catalog_start"] > TARGET_START_YEAR
    ][
        ["_student_id", "_lineage"]
    ]
    .drop_duplicates()
)

later_keys["_later_complete_eligible"] = True

expiring_keys = target_keys.merge(
    later_keys,
    on=["_student_id", "_lineage"],
    how="left",
    validate="one_to_one"
)

later_mask = (
    expiring_keys["_later_complete_eligible"]
    .fillna(False)
    .astype(bool)
)

expiring_keys = expiring_keys[
    ~later_mask
][
    ["_student_id", "_lineage"]
].copy()

expiring = target.merge(
    expiring_keys,
    on=["_student_id", "_lineage"],
    how="inner",
    validate="many_to_one"
)

# One row per student × lineage.
expiring = (
    expiring
    .sort_values(
        ["_student_id", "_lineage"]
    )
    .drop_duplicates(
        ["_student_id", "_lineage"],
        keep="first"
    )
    .reset_index(drop=True)
)

if title_col:
    expiring["_target_credential"] = (
        expiring[title_col]
        .astype(str)
        .str.strip()
    )
else:
    expiring["_target_credential"] = (
        expiring["_lineage"]
        .str.replace("_", " ", regex=False)
        .str.title()
    )

if cred_id_col:
    blank_title = expiring["_target_credential"].eq("")
    expiring.loc[
        blank_title,
        "_target_credential"
    ] = expiring.loc[
        blank_title,
        cred_id_col
    ].astype(str).str.strip()

EXPIRING_LINEAGES = len(expiring)
EXPIRING_STUDENTS = expiring["_student_id"].nunique()


# =============================================================================
# 2. LOAD ENGINE'S KNOWN OFFICIAL AWARDS
# =============================================================================

official = pd.read_csv(
    OFFICIAL,
    dtype=str,
    low_memory=False
).fillna("")

official_student = find_col(
    official.columns,
    ["student_id"],
    label="official-award student ID"
)

official["_student_id"] = official[official_student].map(norm_id)

needed_students = set(expiring["_student_id"])

official = official[
    official["_student_id"].isin(needed_students)
].copy()

# Human-readable preference order.
award_title_candidates = [
    find_col(
        official.columns,
        [
            "credential_title",
            "award_title",
            "program_title",
            "degree_title"
        ],
        required=False
    ),
    find_col(
        official.columns,
        ["credential_id"],
        required=False
    ),
    find_col(
        official.columns,
        ["canonical_lineage"],
        required=False
    ),
    find_col(
        official.columns,
        ["institutional_lineage"],
        required=False
    ),
    find_col(
        official.columns,
        ["Curr1ProgramCode"],
        required=False
    ),
    find_col(
        official.columns,
        ["Major1Code"],
        required=False
    ),
    find_col(
        official.columns,
        ["DegreeCode"],
        required=False
    )
]

award_title_candidates = [
    c for c in award_title_candidates if c
]

if not award_title_candidates:
    fail(
        "08_official_awards_normalized.csv has no usable "
        "credential display field."
    )

official["_known_award_label"] = official.apply(
    lambda row: choose_display_label(
        row,
        award_title_candidates
    ),
    axis=1
)

known_awards = (
    official.groupby(
        "_student_id",
        as_index=False
    )["_known_award_label"]
    .agg(join_unique)
    .rename(
        columns={
            "_known_award_label": "Known Awards"
        }
    )
)


# =============================================================================
# 3. USE EXISTING ENGINE REPRESENTATION PAIRS TO CLASSIFY TARGETS
#
# If student × target lineage exists in 09, the engine already knows that
# an official award represents that modeled credential.
#
# Otherwise, the target credential belongs in Unawarded.
# =============================================================================

representation = pd.read_csv(
    REPRESENTATION,
    dtype=str,
    low_memory=False
).fillna("")

rep_student = find_col(
    representation.columns,
    ["student_id"],
    label="representation student ID"
)

rep_detected = find_col(
    representation.columns,
    ["detected_lineage"],
    label="representation detected lineage"
)

representation["_student_id"] = (
    representation[rep_student].map(norm_id)
)
representation["_lineage"] = (
    representation[rep_detected].map(norm_text)
)

represented_keys = (
    representation[
        ["_student_id", "_lineage"]
    ]
    .drop_duplicates()
)

represented_keys["_known_award_represents_target"] = True

expiring = expiring.merge(
    represented_keys,
    on=["_student_id", "_lineage"],
    how="left",
    validate="one_to_one"
)

represented_mask = (
    expiring["_known_award_represents_target"]
    .fillna(False)
    .astype(bool)
)

expiring["_unawarded"] = ""

expiring.loc[
    ~represented_mask,
    "_unawarded"
] = expiring.loc[
    ~represented_mask,
    "_target_credential"
]

KNOWN_TARGETS = int(represented_mask.sum())
UNAWARDED_TARGETS = int((~represented_mask).sum())

if KNOWN_TARGETS + UNAWARDED_TARGETS != EXPIRING_LINEAGES:
    fail("Award-status partition failed.")


# =============================================================================
# 4. ONE ROW PER STUDENT: AGGREGATE UNAWARDED TARGETS
# =============================================================================

unawarded_by_student = (
    expiring.groupby(
        "_student_id",
        as_index=False
    )["_unawarded"]
    .agg(join_unique)
    .rename(
        columns={
            "_unawarded": "Unawarded"
        }
    )
)

expiring_targets_by_student = (
    expiring.groupby(
        "_student_id",
        as_index=False
    )["_target_credential"]
    .agg(join_unique)
    .rename(
        columns={
            "_target_credential":
                "Expiring 2021-2022 Completed Credential(s)"
        }
    )
)


# =============================================================================
# 5. IDENTITY + MAJOR
# =============================================================================

identity_header = pd.read_csv(
    IDENTITY,
    nrows=0
)

id_col = find_col(
    identity_header.columns,
    ["StudenID", "StudentID", "student_id"],
    label="Banner ID"
)

first_col = find_col(
    identity_header.columns,
    ["FirstName", "first_name"],
    label="first name"
)

last_col = find_col(
    identity_header.columns,
    ["LastName", "last_name"],
    label="last name"
)

major_col = find_col(
    identity_header.columns,
    ["StudentMajor", "student_major", "major"],
    label="declared major"
)

term_col = find_col(
    identity_header.columns,
    ["Term", "term", "term_taken"],
    required=False
)

usecols = [
    id_col,
    first_col,
    last_col,
    major_col
]

if term_col:
    usecols.append(term_col)

identity = pd.read_csv(
    IDENTITY,
    usecols=usecols,
    dtype=str,
    low_memory=False
).fillna("")

identity["_student_id"] = identity[id_col].map(norm_id)

identity = identity[
    identity["_student_id"].isin(needed_students)
].copy()

identity["_row_order"] = range(len(identity))

if term_col:
    identity["_term_sort"] = pd.to_numeric(
        identity[term_col],
        errors="coerce"
    )
else:
    identity["_term_sort"] = identity["_row_order"]


def latest_nonblank(frame, value_col):
    temp = frame[
        frame[value_col]
        .astype(str)
        .str.strip()
        .ne("")
    ].copy()

    temp = temp.sort_values(
        [
            "_student_id",
            "_term_sort",
            "_row_order"
        ]
    )

    return (
        temp.groupby(
            "_student_id",
            as_index=False
        )
        .tail(1)[
            ["_student_id", value_col]
        ]
    )


latest_first = latest_nonblank(identity, first_col)
latest_last = latest_nonblank(identity, last_col)
latest_major = latest_nonblank(identity, major_col)

students = (
    pd.DataFrame(
        {
            "_student_id":
                sorted(needed_students)
        }
    )
    .merge(
        latest_last,
        on="_student_id",
        how="left",
        validate="one_to_one"
    )
    .merge(
        latest_first,
        on="_student_id",
        how="left",
        validate="one_to_one"
    )
    .merge(
        latest_major,
        on="_student_id",
        how="left",
        validate="one_to_one"
    )
    .merge(
        known_awards,
        on="_student_id",
        how="left",
        validate="one_to_one"
    )
    .merge(
        expiring_targets_by_student,
        on="_student_id",
        how="left",
        validate="one_to_one"
    )
    .merge(
        unawarded_by_student,
        on="_student_id",
        how="left",
        validate="one_to_one"
    )
    .fillna("")
)

output = pd.DataFrame(
    {
        "Banner ID":
            students["_student_id"],
        "Last Name":
            students[last_col],
        "First Name":
            students[first_col],
        "Declared Major":
            students[major_col],
        "Catalog Year":
            TARGET_CATALOG,
        "Expiring 2021-2022 Completed Credential(s)":
            students[
                "Expiring 2021-2022 Completed Credential(s)"
            ],
        "Known Awards":
            students["Known Awards"],
        "Unawarded":
            students["Unawarded"]
    }
)

output = output.sort_values(
    [
        "Last Name",
        "First Name",
        "Banner ID"
    ],
    kind="stable"
).reset_index(drop=True)

if len(output) != EXPIRING_STUDENTS:
    fail(
        "Final student row count does not equal "
        "unique expiring students."
    )

STUDENTS_WITH_UNAWARDED = int(
    output["Unawarded"]
    .astype(str)
    .str.strip()
    .ne("")
    .sum()
)

STUDENTS_ALL_TARGETS_AWARDED = (
    len(output) - STUDENTS_WITH_UNAWARDED
)


# =============================================================================
# 6. CHECKS
# =============================================================================

checks = [
    (
        "Expiring student × credential lineages",
        EXPIRING_LINEAGES,
        "PASS"
    ),
    (
        "Unique students",
        EXPIRING_STUDENTS,
        "PASS"
    ),
    (
        "Target credentials represented by known award",
        KNOWN_TARGETS,
        "PASS"
    ),
    (
        "Target credentials truly unawarded",
        UNAWARDED_TARGETS,
        "PASS"
    ),
    (
        "Partition: known + unawarded = expiring",
        f"{KNOWN_TARGETS} + {UNAWARDED_TARGETS} = {EXPIRING_LINEAGES}",
        (
            "PASS"
            if KNOWN_TARGETS + UNAWARDED_TARGETS
            == EXPIRING_LINEAGES
            else "FAIL"
        )
    ),
    (
        "Students with at least one unawarded target",
        STUDENTS_WITH_UNAWARDED,
        "PASS"
    ),
    (
        "Students whose expiring targets are all known awards",
        STUDENTS_ALL_TARGETS_AWARDED,
        "PASS"
    )
]

if any(status == "FAIL" for _, _, status in checks):
    fail("Control check failed.")


# =============================================================================
# 7. TWO-SHEET WORKBOOK
# =============================================================================

GREEN = "00573F"
WHITE = "FFFFFF"
LIGHT_GREEN = "E8F1ED"
LIGHT_ORANGE = "FCE9D7"

wb = Workbook()

summary = wb.active
summary.title = "CONTROL_SUMMARY"
summary.sheet_view.showGridLines = False

detail = wb.create_sheet("STUDENT_RECONCILIATION")
detail.sheet_view.showGridLines = False

summary.merge_cells("A1:D1")
summary["A1"] = "2021–2022 Expiring Credential Award Reconciliation"
summary["A1"].fill = PatternFill("solid", fgColor=GREEN)
summary["A1"].font = Font(
    bold=True,
    color=WHITE,
    size=16
)

summary.merge_cells("A2:D2")
summary["A2"] = (
    "One row per student. Known Awards are drawn from the engine's "
    "normalized official-award dataset. Unawarded contains only expiring "
    "2021–2022 completed credentials not represented by a known award."
)
summary["A2"].alignment = Alignment(
    wrap_text=True
)

r = 4

metadata = [
    ("Run timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ("Target catalog", TARGET_CATALOG),
    (
        "Selection rule",
        "2021-2022 COMPLETE + catalog-eligible with no later COMPLETE + eligible version of same lineage"
    ),
    (
        "Award-status rule",
        "Existing engine representation pairs; no new matching logic"
    )
]

for label, value in metadata:
    summary.cell(r, 1).value = label
    summary.cell(r, 1).font = Font(bold=True)
    summary.merge_cells(
        start_row=r,
        start_column=2,
        end_row=r,
        end_column=4
    )
    summary.cell(r, 2).value = value
    summary.cell(r, 2).alignment = Alignment(
        wrap_text=True
    )
    r += 1

r += 1

summary.merge_cells(
    start_row=r,
    start_column=1,
    end_row=r,
    end_column=4
)
summary.cell(r, 1).value = "FUNNEL / CHECKS"
summary.cell(r, 1).fill = PatternFill(
    "solid",
    fgColor=GREEN
)
summary.cell(r, 1).font = Font(
    bold=True,
    color=WHITE
)

r += 1

for name, value, status in checks:
    summary.cell(r, 1).value = name
    summary.cell(r, 2).value = value
    summary.cell(r, 3).value = status

    if "truly unawarded" in name.lower():
        summary.cell(r, 1).fill = PatternFill(
            "solid",
            fgColor=LIGHT_ORANGE
        )
        summary.cell(r, 2).fill = PatternFill(
            "solid",
            fgColor=LIGHT_ORANGE
        )

    r += 1

r += 1

summary.merge_cells(
    start_row=r,
    start_column=1,
    end_row=r,
    end_column=4
)
summary.cell(r, 1).value = "SOURCE CHECKSUMS"
summary.cell(r, 1).fill = PatternFill(
    "solid",
    fgColor=GREEN
)
summary.cell(r, 1).font = Font(
    bold=True,
    color=WHITE
)

r += 1

for name, path in [
    ("Eligible catalog alternatives", ELIGIBLE),
    ("Normalized official awards", OFFICIAL),
    ("Award representation pairs", REPRESENTATION),
    ("Identity / major source", IDENTITY)
]:
    summary.cell(r, 1).value = name
    summary.cell(r, 2).value = str(path.relative_to(ROOT))
    summary.merge_cells(
        start_row=r,
        start_column=3,
        end_row=r,
        end_column=4
    )
    summary.cell(r, 3).value = sha256_file(path)
    summary.cell(r, 3).font = Font(
        name="Consolas",
        size=8
    )
    r += 1

summary.column_dimensions["A"].width = 43
summary.column_dimensions["B"].width = 45
summary.column_dimensions["C"].width = 38
summary.column_dimensions["D"].width = 38


# Detail
detail.append(list(output.columns))

for row in output.itertuples(
    index=False,
    name=None
):
    detail.append(list(row))

for cell in detail[1]:
    cell.fill = PatternFill(
        "solid",
        fgColor=GREEN
    )
    cell.font = Font(
        bold=True,
        color=WHITE
    )
    cell.alignment = Alignment(
        wrap_text=True,
        vertical="center"
    )

detail.freeze_panes = "A2"

if len(output):
    ref = (
        f"A1:"
        f"{get_column_letter(len(output.columns))}"
        f"{len(output) + 1}"
    )

    table = Table(
        displayName="StudentAwardReconciliation",
        ref=ref
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False
    )

    detail.add_table(table)

widths = {
    "A": 15,
    "B": 22,
    "C": 20,
    "D": 34,
    "E": 15,
    "F": 48,
    "G": 60,
    "H": 48
}

for col, width in widths.items():
    detail.column_dimensions[col].width = width

for row in detail.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

for cell in detail["A"][1:]:
    cell.number_format = "@"

wb.save(OUTFILE)


# =============================================================================
# 8. STUDENT-BLIND CONSOLE
# =============================================================================

print()
print("=" * 110)
print("RESULTS — STUDENT-BLIND")
print("=" * 110)

print(
    f"Expiring 2021-2022 student × credential lineages: "
    f"{EXPIRING_LINEAGES:,}"
)

print(
    f"Unique students: "
    f"{EXPIRING_STUDENTS:,}"
)

print(
    f"Expiring target credentials already represented by known award: "
    f"{KNOWN_TARGETS:,}"
)

print(
    f"Expiring target credentials truly unawarded: "
    f"{UNAWARDED_TARGETS:,}"
)

print(
    f"Students with at least one truly unawarded credential: "
    f"{STUDENTS_WITH_UNAWARDED:,}"
)

print(
    f"Students whose expiring credentials are all already known awards: "
    f"{STUDENTS_ALL_TARGETS_AWARDED:,}"
)

print()
print("=" * 110)
print("OUTPUT")
print("=" * 110)

print(OUTFILE.resolve())

print()
print(
    "FERPA: workbook contains student-identifying information. "
    "Keep it local."
)

