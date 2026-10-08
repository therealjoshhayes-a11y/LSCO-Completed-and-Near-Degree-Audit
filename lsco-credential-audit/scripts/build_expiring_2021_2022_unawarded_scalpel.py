from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


# =============================================================================
# 2021-2022 EXPIRING CATALOG — UNAWARDED CREDENTIAL SCALPEL QUERY
#
# TARGET
# ------
# Validated unawarded 2021-2022 student × credential-lineage opportunities
# MINUS any same-student × same-lineage completion that remains COMPLETE and
# catalog-eligible under ANY later catalog in the modeled stack.
#
# FERPA
# -----
# This output intentionally contains student identifiers.
# IT REMAINS LOCAL TO LSCO.
# Do not upload the resulting workbook outside the authorized environment.
# =============================================================================


ROOT = Path.cwd()

TARGET_CATALOG = "2021-2022"
TARGET_START_YEAR = 2021
EXPIRATION_DATE = "2026-08-31"

FINAL_DIR = (
    ROOT
    / "data"
    / "processed"
    / "reporting"
    / "final_additional_award_numbers_20260803"
)

FINAL_DETAIL = (
    FINAL_DIR
    / "01_final_additional_awards_student_detail.csv"
)

CLEAN_DIR = (
    ROOT
    / "data"
    / "processed"
    / "reporting"
    / "additional_awards_clean_query_20260803"
)

ALL_ELIGIBLE_ALTERNATIVES = (
    CLEAN_DIR
    / "05_all_eligible_catalog_alternatives.csv"
)

IDENTITY_SOURCE = (
    ROOT
    / "data"
    / "raw"
    / "student_exports"
    / "banner_course_history_actual.csv"
)

OUTPUT_DIR = (
    ROOT
    / "data"
    / "processed"
    / "reporting"
    / "expiring_2021_2022_unawarded_20260811"
)

OUTPUT_XLSX = (
    OUTPUT_DIR
    / "2021_2022_EXPIRING_UNAWARDED_CREDENTIAL_REVIEW.xlsx"
)

MISSING_DIAGNOSTIC = (
    OUTPUT_DIR
    / "98_MISSING_IDENTITY_OR_MAJOR_LOCAL_ONLY.csv"
)

WORKBOOK_HASH_FILE = (
    OUTPUT_DIR
    / "WORKBOOK_SHA256.txt"
)


# =============================================================================
# HELPERS
# =============================================================================


def fail(message: str) -> None:
    raise RuntimeError(message)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def normalize_id(value: object) -> str:
    return str(value).strip()


def normalize_lineage(value: object) -> str:
    return (
        str(value)
        .strip()
        .upper()
        .replace(" ", "_")
    )


def catalog_start_year(value: object) -> int:
    match = re.search(r"(20\d{2})", str(value))
    if not match:
        return -1
    return int(match.group(1))


def find_column(
    columns,
    candidates,
    required=True,
    label="column",
):
    lookup = {
        str(c).strip().lower(): c
        for c in columns
    }

    for candidate in candidates:
        key = candidate.lower()
        if key in lookup:
            return lookup[key]

    if required:
        fail(
            f"Could not find {label}. "
            f"Tried: {candidates}. "
            f"Available columns: {list(columns)}"
        )

    return None


def find_lineage_column(df: pd.DataFrame) -> str:
    return find_column(
        df.columns,
        [
            "credential_lineage",
            "detected_lineage",
            "canonical_lineage",
        ],
        required=True,
        label="credential lineage column",
    )


def latest_nonblank_value(
    frame: pd.DataFrame,
    value_column: str,
) -> pd.DataFrame:
    temp = frame[
        frame[value_column]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
    ].copy()

    temp = temp.sort_values(
        ["student_id", "_term_sort", "_row_order"],
        ascending=[True, True, True],
    )

    return (
        temp.groupby(
            "student_id",
            as_index=False,
            dropna=False,
        )
        .tail(1)
        [
            [
                "student_id",
                value_column,
            ]
        ]
        .reset_index(drop=True)
    )


# =============================================================================
# PRE-FLIGHT
# =============================================================================


print("=" * 110)
print("2021-2022 EXPIRING CATALOG — SCALPEL QUERY")
print("=" * 110)

required_files = [
    FINAL_DETAIL,
    ALL_ELIGIBLE_ALTERNATIVES,
    IDENTITY_SOURCE,
]

for path in required_files:
    if not path.exists():
        fail(f"Required source not found: {path}")

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

print(f"Frozen final detail: {FINAL_DETAIL}")
print(f"Later-catalog universe: {ALL_ELIGIBLE_ALTERNATIVES}")
print(f"Identity/major source: {IDENTITY_SOURCE}")
print()


# =============================================================================
# 1. FROZEN FINAL UNAWARDED SET
# =============================================================================


final_detail = pd.read_csv(
    FINAL_DETAIL,
    dtype=str,
    low_memory=False,
).fillna("")

for required in [
    "student_id",
    "catalog_year",
]:
    if required not in final_detail.columns:
        fail(
            f"Frozen final detail missing required column: {required}"
        )

final_lineage_col = find_lineage_column(
    final_detail
)

final_detail["student_id"] = (
    final_detail["student_id"]
    .map(normalize_id)
)

final_detail["_lineage"] = (
    final_detail[final_lineage_col]
    .map(normalize_lineage)
)

target = final_detail[
    final_detail["catalog_year"]
    .astype(str)
    .str.strip()
    .eq(TARGET_CATALOG)
].copy()

if target.empty:
    fail(
        f"No frozen final unawarded rows found for {TARGET_CATALOG}."
    )

if target["student_id"].eq("").any():
    fail("Blank student_id found in target frozen detail.")

if target["_lineage"].eq("").any():
    fail("Blank credential lineage found in target frozen detail.")

target_key_duplicates = (
    target.duplicated(
        subset=[
            "student_id",
            "_lineage",
        ],
        keep=False,
    )
)

if target_key_duplicates.any():
    dup_count = int(
        target_key_duplicates.sum()
    )
    fail(
        "Frozen target is not unique at "
        "student_id × credential_lineage. "
        f"Duplicate rows: {dup_count}"
    )

START_ROWS = len(target)
START_STUDENTS = target["student_id"].nunique()


# =============================================================================
# 2. ALL VALIDATED LATER-CATALOG COMPLETE + ELIGIBLE ALTERNATIVES
#
# This artifact was produced from the COMPLETE universe after application of
# the catalog-floor eligibility rules. Stage 05 intentionally preserves ALL
# qualifying catalog alternatives before the one-row-per-lineage selection.
# =============================================================================


alternatives = pd.read_csv(
    ALL_ELIGIBLE_ALTERNATIVES,
    dtype=str,
    low_memory=False,
).fillna("")

for required in [
    "student_id",
    "catalog_year",
]:
    if required not in alternatives.columns:
        fail(
            "All-eligible-alternatives file missing "
            f"required column: {required}"
        )

alt_lineage_col = find_lineage_column(
    alternatives
)

alternatives["student_id"] = (
    alternatives["student_id"]
    .map(normalize_id)
)

alternatives["_lineage"] = (
    alternatives[alt_lineage_col]
    .map(normalize_lineage)
)

alternatives["_catalog_start"] = (
    alternatives["catalog_year"]
    .map(catalog_start_year)
)

# If audit_status is present, require COMPLETE explicitly.
if "audit_status" in alternatives.columns:
    status_values = (
        alternatives["audit_status"]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    alternatives = alternatives[
        status_values.eq("COMPLETE")
    ].copy()

later = alternatives[
    alternatives["_catalog_start"]
    .gt(TARGET_START_YEAR)
].copy()

# Reduce immediately to only students/lineages in the expiring set.
target_keys = target[
    [
        "student_id",
        "_lineage",
    ]
].drop_duplicates()

later = later.merge(
    target_keys,
    on=[
        "student_id",
        "_lineage",
    ],
    how="inner",
    validate="many_to_one",
)

later_keys = later[
    [
        "student_id",
        "_lineage",
    ]
].drop_duplicates()

later_keys["_later_qualifies"] = True


# =============================================================================
# 3. HARD ANTI-JOIN
#
# KEEP ONLY 2021-2022 UNAWARDED ROWS THAT HAVE ZERO LATER CATALOG MATCHES
# WITHIN THE SAME STUDENT × CANONICAL CREDENTIAL LINEAGE.
# =============================================================================


scored = target.merge(
    later_keys,
    on=[
        "student_id",
        "_lineage",
    ],
    how="left",
    validate="one_to_one",
)

later_qualifies_mask = (
    scored["_later_qualifies"]
    .fillna(False)
    .astype(bool)
)

excluded = scored[
    later_qualifies_mask
].copy()

expiring = scored[
    ~later_qualifies_mask
].copy()

EXCLUDED_ROWS = len(excluded)
EXCLUDED_STUDENTS = (
    excluded["student_id"].nunique()
    if not excluded.empty
    else 0
)

FINAL_ROWS = len(expiring)
FINAL_STUDENTS = expiring["student_id"].nunique()

if START_ROWS - EXCLUDED_ROWS != FINAL_ROWS:
    fail(
        "Funnel reconciliation failed: "
        f"{START_ROWS} - {EXCLUDED_ROWS} != {FINAL_ROWS}"
    )

# Hard survivor check.
survivor_check = (
    expiring[
        [
            "student_id",
            "_lineage",
        ]
    ]
    .merge(
        later_keys[
            [
                "student_id",
                "_lineage",
            ]
        ],
        on=[
            "student_id",
            "_lineage",
        ],
        how="inner",
    )
)

LATER_SURVIVORS = len(
    survivor_check
)

if LATER_SURVIVORS != 0:
    fail(
        "SCALPEL FAILURE: final output still contains "
        f"{LATER_SURVIVORS} student-lineage rows with "
        "a later qualifying catalog."
    )


# =============================================================================
# 4. LOCAL ARGOS IDENTITY + DECLARED MAJOR
#
# Source columns previously used in this project:
# StudenID, FirstName, LastName, StudentMajor, Term
#
# Declared Major = latest nonblank StudentMajor by academic Term in the Argos
# course-history extract.
# =============================================================================


identity_header = pd.read_csv(
    IDENTITY_SOURCE,
    nrows=0,
)

id_col = find_column(
    identity_header.columns,
    ["StudenID", "StudentID", "student_id"],
    label="Banner/student ID column",
)

first_col = find_column(
    identity_header.columns,
    ["FirstName", "first_name"],
    label="first-name column",
)

last_col = find_column(
    identity_header.columns,
    ["LastName", "last_name"],
    label="last-name column",
)

major_col = find_column(
    identity_header.columns,
    ["StudentMajor", "student_major", "major"],
    label="declared-major column",
)

term_col = find_column(
    identity_header.columns,
    ["Term", "term", "term_taken", "term_sort"],
    label="term column",
)

identity_raw = pd.read_csv(
    IDENTITY_SOURCE,
    usecols=[
        id_col,
        first_col,
        last_col,
        major_col,
        term_col,
    ],
    dtype=str,
    low_memory=False,
).fillna("")

identity_raw = identity_raw.rename(
    columns={
        id_col: "student_id",
        first_col: "first_name",
        last_col: "last_name",
        major_col: "declared_major",
        term_col: "term",
    }
)

identity_raw["student_id"] = (
    identity_raw["student_id"]
    .map(normalize_id)
)

needed_students = set(
    expiring["student_id"]
    .astype(str)
)

identity_raw = identity_raw[
    identity_raw["student_id"]
    .isin(needed_students)
].copy()

identity_raw["_row_order"] = range(
    len(identity_raw)
)

identity_raw["_term_sort"] = pd.to_numeric(
    identity_raw["term"],
    errors="coerce",
)

if (
    not identity_raw.empty
    and identity_raw["_term_sort"].isna().all()
):
    fail(
        "Could not parse any Argos Term values numerically."
    )

# Latest available first name, last name, and declared major independently.
latest_first = latest_nonblank_value(
    identity_raw,
    "first_name",
)

latest_last = latest_nonblank_value(
    identity_raw,
    "last_name",
)

latest_major = latest_nonblank_value(
    identity_raw,
    "declared_major",
)

identity = (
    pd.DataFrame(
        {
            "student_id":
                sorted(needed_students)
        }
    )
    .merge(
        latest_first,
        on="student_id",
        how="left",
        validate="one_to_one",
    )
    .merge(
        latest_last,
        on="student_id",
        how="left",
        validate="one_to_one",
    )
    .merge(
        latest_major,
        on="student_id",
        how="left",
        validate="one_to_one",
    )
    .fillna("")
)

missing_identity = identity[
    identity["student_id"].eq("")
    | identity["first_name"].astype(str).str.strip().eq("")
    | identity["last_name"].astype(str).str.strip().eq("")
    | identity["declared_major"].astype(str).str.strip().eq("")
].copy()

if not missing_identity.empty:
    missing_identity.to_csv(
        MISSING_DIAGNOSTIC,
        index=False,
    )

    fail(
        "Identity/major join is incomplete. "
        f"{len(missing_identity):,} students are missing "
        "Banner ID, name, or declared major. "
        f"LOCAL FERPA diagnostic written to: {MISSING_DIAGNOSTIC}"
    )


# =============================================================================
# 5. BUILD PURE OPERATIONAL LIST
# =============================================================================


credential_title_col = find_column(
    expiring.columns,
    [
        "credential_title",
        "program_title",
        "title",
    ],
    required=False,
)

credential_id_col = find_column(
    expiring.columns,
    [
        "credential_id",
        "credential_base_id",
    ],
    required=False,
)

completion_term_col = find_column(
    expiring.columns,
    [
        "award_term_taken",
        "modeled_completion_term",
        "completion_term",
        "completion_term_taken",
        "award_term",
    ],
    required=False,
)

if credential_title_col:
    expiring["_credential_title"] = (
        expiring[credential_title_col]
        .astype(str)
        .str.strip()
    )
else:
    expiring["_credential_title"] = ""

missing_titles = (
    expiring["_credential_title"]
    .eq("")
)

expiring.loc[
    missing_titles,
    "_credential_title",
] = (
    expiring.loc[
        missing_titles,
        "_lineage",
    ]
    .str.replace("_", " ", regex=False)
    .str.title()
)

if credential_id_col:
    expiring["_credential_id"] = (
        expiring[credential_id_col]
        .astype(str)
        .str.strip()
    )
else:
    expiring["_credential_id"] = ""

if completion_term_col:
    expiring["_completion_term"] = (
        expiring[completion_term_col]
        .astype(str)
        .str.strip()
    )
else:
    expiring["_completion_term"] = ""

operational = (
    expiring
    .merge(
        identity,
        on="student_id",
        how="left",
        validate="many_to_one",
    )
)

operational = pd.DataFrame(
    {
        "Banner ID":
            operational["student_id"],
        "Last Name":
            operational["last_name"],
        "First Name":
            operational["first_name"],
        "Declared Major":
            operational["declared_major"],
        "Unawarded Credential":
            operational["_credential_title"],
        "Credential ID":
            operational["_credential_id"],
        "Catalog Year":
            operational["catalog_year"],
        "Modeled Completion Term":
            operational["_completion_term"],
    }
)

operational = (
    operational
    .sort_values(
        [
            "Unawarded Credential",
            "Last Name",
            "First Name",
            "Banner ID",
        ],
        kind="stable",
    )
    .reset_index(drop=True)
)

if len(operational) != FINAL_ROWS:
    fail(
        "Operational-list row count changed during identity join."
    )

if not operational["Catalog Year"].eq(TARGET_CATALOG).all():
    fail(
        "Operational output contains a catalog year "
        f"other than {TARGET_CATALOG}."
    )

FINAL_DUPLICATES = int(
    operational.duplicated(
        subset=[
            "Banner ID",
            "Unawarded Credential",
        ]
    ).sum()
)

# The authoritative duplicate test remains student × lineage above.
# This display-level check is informational because titles can theoretically collide.


# =============================================================================
# 6. STUDENT-BLIND AGGREGATES
# =============================================================================


credential_counts = (
    operational.groupby(
        "Unawarded Credential",
        dropna=False,
    )
    .agg(
        Expiring_Award_Opportunities=(
            "Banner ID",
            "size",
        ),
        Distinct_Students=(
            "Banner ID",
            "nunique",
        ),
    )
    .reset_index()
    .sort_values(
        [
            "Expiring_Award_Opportunities",
            "Unawarded Credential",
        ],
        ascending=[
            False,
            True,
        ],
    )
    .reset_index(drop=True)
)

students_with_multiple = (
    operational.groupby("Banner ID")
    .size()
    .gt(1)
    .sum()
)

logical_csv = operational.to_csv(
    index=False,
    lineterminator="\n",
)

LOGICAL_LIST_SHA256 = (
    sha256_text(logical_csv)
)


# =============================================================================
# 7. CONTROL CHECKS
# =============================================================================


checks = [
    (
        "Funnel reconciliation: Start - Excluded = Final",
        f"{START_ROWS} - {EXCLUDED_ROWS} = {FINAL_ROWS}",
        "PASS"
        if START_ROWS - EXCLUDED_ROWS == FINAL_ROWS
        else "FAIL",
    ),
    (
        "Later-catalog survivor overlap",
        LATER_SURVIVORS,
        "PASS"
        if LATER_SURVIVORS == 0
        else "FAIL",
    ),
    (
        "Final student × lineage duplicate keys",
        0,
        "PASS",
    ),
    (
        "Missing Banner ID",
        int(
            operational["Banner ID"]
            .astype(str)
            .str.strip()
            .eq("")
            .sum()
        ),
        "PASS"
        if not operational["Banner ID"]
        .astype(str)
        .str.strip()
        .eq("")
        .any()
        else "FAIL",
    ),
    (
        "Missing student name",
        int(
            (
                operational["First Name"]
                .astype(str)
                .str.strip()
                .eq("")
                |
                operational["Last Name"]
                .astype(str)
                .str.strip()
                .eq("")
            ).sum()
        ),
        "PASS"
        if not (
            operational["First Name"]
            .astype(str)
            .str.strip()
            .eq("")
            |
            operational["Last Name"]
            .astype(str)
            .str.strip()
            .eq("")
        ).any()
        else "FAIL",
    ),
    (
        "Missing declared major",
        int(
            operational["Declared Major"]
            .astype(str)
            .str.strip()
            .eq("")
            .sum()
        ),
        "PASS"
        if not operational["Declared Major"]
        .astype(str)
        .str.strip()
        .eq("")
        .any()
        else "FAIL",
    ),
    (
        f"All final rows are {TARGET_CATALOG}",
        FINAL_ROWS,
        "PASS"
        if operational["Catalog Year"]
        .eq(TARGET_CATALOG)
        .all()
        else "FAIL",
    ),
]


if any(
    status == "FAIL"
    for _, _, status in checks
):
    fail(
        "One or more final control checks failed."
    )


# =============================================================================
# 8. CREATE TWO-SHEET WORKBOOK
# =============================================================================


LSCO_GREEN = "00573F"
LSCO_ORANGE = "E57200"
LIGHT_GREEN = "E8F1ED"
LIGHT_ORANGE = "FCE9D7"
WHITE = "FFFFFF"
DARK = "222222"
LIGHT_GRAY = "E7E7E7"
MID_GRAY = "666666"

thin_gray = Side(
    style="thin",
    color="D9D9D9",
)

wb = Workbook()

summary = wb.active
summary.title = "CONTROL_SUMMARY"
summary.sheet_view.showGridLines = False

detail = wb.create_sheet(
    "EXPIRING_AWARDS"
)

detail.sheet_view.showGridLines = False


# -----------------------------------------------------------------------------
# CONTROL SUMMARY
# -----------------------------------------------------------------------------


summary.merge_cells("A1:D1")
summary["A1"] = (
    "2021–2022 Expiring Catalog — "
    "Unawarded Credential Review"
)

summary["A1"].fill = PatternFill(
    "solid",
    fgColor=LSCO_GREEN,
)
summary["A1"].font = Font(
    color=WHITE,
    bold=True,
    size=16,
)
summary["A1"].alignment = Alignment(
    vertical="center",
)
summary.row_dimensions[1].height = 28

summary["A2"] = (
    "SCALPEL FILTER: Retains only frozen 2021–2022 unawarded "
    "student × credential-lineage opportunities with NO later "
    "COMPLETE + catalog-eligible version of the same lineage."
)
summary.merge_cells("A2:D2")
summary["A2"].font = Font(
    italic=True,
    color=MID_GRAY,
    size=10,
)
summary["A2"].alignment = Alignment(
    wrap_text=True,
    vertical="top",
)
summary.row_dimensions[2].height = 34


def section_title(row: int, text: str):
    summary.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=4,
    )
    cell = summary.cell(row, 1)
    cell.value = text
    cell.fill = PatternFill(
        "solid",
        fgColor=LSCO_GREEN,
    )
    cell.font = Font(
        color=WHITE,
        bold=True,
        size=11,
    )


def kv_row(row: int, label: str, value, note=""):
    summary.cell(row, 1).value = label
    summary.cell(row, 1).font = Font(
        bold=True,
        color=DARK,
    )

    summary.cell(row, 2).value = value

    summary.merge_cells(
        start_row=row,
        start_column=3,
        end_row=row,
        end_column=4,
    )
    summary.cell(row, 3).value = note
    summary.cell(row, 3).font = Font(
        color=MID_GRAY,
        italic=True,
        size=9,
    )
    summary.cell(row, 3).alignment = Alignment(
        wrap_text=True,
    )

    for col in range(1, 5):
        summary.cell(row, col).border = Border(
            bottom=thin_gray,
        )


row = 4

section_title(row, "RUN METADATA")
row += 1

kv_row(
    row,
    "Run timestamp",
    datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    ),
    "Local LSCO execution time.",
)
row += 1

kv_row(
    row,
    "Target catalog",
    TARGET_CATALOG,
    f"Catalog opportunity expires after {EXPIRATION_DATE}.",
)
row += 1

kv_row(
    row,
    "Qualification definition",
    "Same student + same canonical credential lineage",
    (
        "A later catalog removes urgency only when the validated "
        "all-eligible-alternatives artifact contains a later "
        "COMPLETE/catalog-eligible row for that same lineage."
    ),
)
row += 1

kv_row(
    row,
    "Identity / major rule",
    "Latest nonblank Argos value by Term",
    (
        "Names and StudentMajor are read locally from "
        "banner_course_history_actual.csv."
    ),
)
row += 2


section_title(row, "AGGREGATE FUNNEL")
row += 1

funnel_rows = [
    (
        "Starting 2021–2022 validated unawarded opportunities",
        START_ROWS,
        "Frozen final award package.",
    ),
    (
        "Starting unique students",
        START_STUDENTS,
        "",
    ),
    (
        "Excluded: same lineage qualifies under later catalog",
        EXCLUDED_ROWS,
        "Hard anti-join exclusion.",
    ),
    (
        "Unique students represented in exclusions",
        EXCLUDED_STUDENTS,
        "",
    ),
    (
        "FINAL expiring award opportunities",
        FINAL_ROWS,
        "Requires attention before catalog expiration.",
    ),
    (
        "FINAL unique students",
        FINAL_STUDENTS,
        "",
    ),
    (
        "Students with >1 expiring credential",
        int(students_with_multiple),
        "",
    ),
]

for label, value, note in funnel_rows:
    kv_row(
        row,
        label,
        int(value),
        note,
    )

    if label.startswith("FINAL"):
        summary.cell(row, 1).fill = PatternFill(
            "solid",
            fgColor=LIGHT_ORANGE,
        )
        summary.cell(row, 2).fill = PatternFill(
            "solid",
            fgColor=LIGHT_ORANGE,
        )
        summary.cell(row, 1).font = Font(
            bold=True,
            color=LSCO_GREEN,
        )
        summary.cell(row, 2).font = Font(
            bold=True,
            color=LSCO_GREEN,
        )

    row += 1

row += 1


section_title(row, "CONTROL CHECKS")
row += 1

summary.cell(row, 1).value = "Check"
summary.cell(row, 2).value = "Observed"
summary.cell(row, 3).value = "Status"

for c in range(1, 4):
    summary.cell(row, c).fill = PatternFill(
        "solid",
        fgColor=LIGHT_GREEN,
    )
    summary.cell(row, c).font = Font(
        bold=True,
        color=LSCO_GREEN,
    )

row += 1

for check_name, observed, status in checks:
    summary.cell(row, 1).value = check_name
    summary.cell(row, 2).value = observed
    summary.cell(row, 3).value = status

    summary.cell(row, 3).font = Font(
        bold=True,
        color=LSCO_GREEN
        if status == "PASS"
        else "C00000",
    )

    for c in range(1, 4):
        summary.cell(row, c).border = Border(
            bottom=thin_gray,
        )

    row += 1

row += 1


section_title(row, "SOURCE PROVENANCE & CHECKSUMS")
row += 1

checksum_rows = [
    (
        "Frozen final unawarded source",
        str(FINAL_DETAIL.relative_to(ROOT)),
        sha256_file(FINAL_DETAIL),
    ),
    (
        "Later-catalog COMPLETE/eligible universe",
        str(
            ALL_ELIGIBLE_ALTERNATIVES
            .relative_to(ROOT)
        ),
        sha256_file(
            ALL_ELIGIBLE_ALTERNATIVES
        ),
    ),
    (
        "Argos identity/major source",
        str(
            IDENTITY_SOURCE
            .relative_to(ROOT)
        ),
        sha256_file(
            IDENTITY_SOURCE
        ),
    ),
    (
        "Query script",
        str(
            Path(__file__).resolve()
            .relative_to(ROOT)
        ),
        sha256_file(
            Path(__file__).resolve()
        ),
    ),
    (
        "Logical EXPIRING_AWARDS list",
        "Deterministic UTF-8 CSV representation of Sheet 2",
        LOGICAL_LIST_SHA256,
    ),
]

summary.cell(row, 1).value = "Source"
summary.cell(row, 2).value = "Path / Definition"
summary.cell(row, 3).value = "SHA-256"
summary.merge_cells(
    start_row=row,
    start_column=3,
    end_row=row,
    end_column=4,
)

for c in range(1, 5):
    summary.cell(row, c).fill = PatternFill(
        "solid",
        fgColor=LIGHT_GREEN,
    )
    summary.cell(row, c).font = Font(
        bold=True,
        color=LSCO_GREEN,
    )

row += 1

for source, path_text, digest in checksum_rows:
    summary.cell(row, 1).value = source
    summary.cell(row, 2).value = path_text

    summary.merge_cells(
        start_row=row,
        start_column=3,
        end_row=row,
        end_column=4,
    )
    summary.cell(row, 3).value = digest
    summary.cell(row, 3).font = Font(
        name="Consolas",
        size=8,
    )

    for c in range(1, 5):
        summary.cell(row, c).border = Border(
            bottom=thin_gray,
        )

    row += 1

row += 1


section_title(row, "FINAL COUNTS BY UNAWARDED CREDENTIAL")
row += 1

summary.cell(row, 1).value = "Unawarded Credential"
summary.cell(row, 2).value = "Opportunities"
summary.cell(row, 3).value = "Distinct Students"

for c in range(1, 4):
    summary.cell(row, c).fill = PatternFill(
        "solid",
        fgColor=LIGHT_GREEN,
    )
    summary.cell(row, c).font = Font(
        bold=True,
        color=LSCO_GREEN,
    )

row += 1

for _, record in credential_counts.iterrows():
    summary.cell(row, 1).value = (
        record["Unawarded Credential"]
    )
    summary.cell(row, 2).value = int(
        record["Expiring_Award_Opportunities"]
    )
    summary.cell(row, 3).value = int(
        record["Distinct_Students"]
    )

    for c in range(1, 4):
        summary.cell(row, c).border = Border(
            bottom=thin_gray,
        )

    row += 1


summary.column_dimensions["A"].width = 46
summary.column_dimensions["B"].width = 42
summary.column_dimensions["C"].width = 39
summary.column_dimensions["D"].width = 39

for summary_row in summary.iter_rows():
    for cell in summary_row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

summary.freeze_panes = "A4"


# -----------------------------------------------------------------------------
# PURE OPERATIONAL LIST — NOTHING BUT THE LIST
# -----------------------------------------------------------------------------


headers = list(
    operational.columns
)

detail.append(headers)

for record in operational.itertuples(
    index=False,
    name=None,
):
    detail.append(list(record))

for cell in detail[1]:
    cell.fill = PatternFill(
        "solid",
        fgColor=LSCO_GREEN,
    )
    cell.font = Font(
        bold=True,
        color=WHITE,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

detail.row_dimensions[1].height = 30
detail.freeze_panes = "A2"
detail.auto_filter.ref = detail.dimensions

# Excel table.
if FINAL_ROWS > 0:
    table_ref = (
        f"A1:"
        f"{get_column_letter(len(headers))}"
        f"{FINAL_ROWS + 1}"
    )

    table = Table(
        displayName="ExpiringAwardsTable",
        ref=table_ref,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    detail.add_table(table)

widths = {
    "A": 15,
    "B": 22,
    "C": 20,
    "D": 34,
    "E": 42,
    "F": 30,
    "G": 15,
    "H": 24,
}

for col, width in widths.items():
    detail.column_dimensions[col].width = width

for row_cells in detail.iter_rows():
    for cell in row_cells:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

# Keep Banner IDs text-formatted.
for cell in detail["A"][1:]:
    cell.number_format = "@"

# Light bottom borders make the pure list readable without clutter.
for data_row in detail.iter_rows(
    min_row=2,
):
    for cell in data_row:
        cell.border = Border(
            bottom=thin_gray,
        )


# =============================================================================
# 9. SAVE + HASH WORKBOOK
# =============================================================================


wb.save(
    OUTPUT_XLSX
)

WORKBOOK_SHA256 = sha256_file(
    OUTPUT_XLSX
)

WORKBOOK_HASH_FILE.write_text(
    "\n".join(
        [
            "2021-2022 EXPIRING UNAWARDED CREDENTIAL REVIEW",
            "=" * 72,
            "",
            f"Workbook: {OUTPUT_XLSX.name}",
            f"SHA-256: {WORKBOOK_SHA256}",
            "",
            (
                "NOTE: The workbook contains FERPA-protected "
                "student-identifying information and remains local."
            ),
        ]
    ),
    encoding="utf-8",
)


# =============================================================================
# 10. CONSOLE — STUDENT-BLIND ONLY
# =============================================================================


print()
print("=" * 110)
print("SCALPEL FUNNEL — STUDENT-BLIND")
print("=" * 110)

print(
    f"Starting {TARGET_CATALOG} validated unawarded opportunities: "
    f"{START_ROWS:,}"
)
print(
    f"Starting unique students: "
    f"{START_STUDENTS:,}"
)
print(
    "Excluded because same student / same credential lineage "
    f"qualifies under a later catalog: {EXCLUDED_ROWS:,}"
)
print(
    f"Unique students represented in exclusions: "
    f"{EXCLUDED_STUDENTS:,}"
)
print(
    f"FINAL EXPIRING OPPORTUNITIES: "
    f"{FINAL_ROWS:,}"
)
print(
    f"FINAL UNIQUE STUDENTS: "
    f"{FINAL_STUDENTS:,}"
)
print(
    f"Students with >1 expiring credential: "
    f"{int(students_with_multiple):,}"
)

print()
print("=" * 110)
print("CONTROL CHECKS")
print("=" * 110)

for name, observed, status in checks:
    print(
        f"{status:4s} | {name}: {observed}"
    )

print()
print("=" * 110)
print("FINAL COUNTS BY CREDENTIAL")
print("=" * 110)

if credential_counts.empty:
    print("No expiring credentials remain after the later-catalog anti-join.")
else:
    print(
        credential_counts.to_string(
            index=False
        )
    )

print()
print("=" * 110)
print("CHECKSUMS")
print("=" * 110)

print(
    f"Logical EXPIRING_AWARDS SHA-256: "
    f"{LOGICAL_LIST_SHA256}"
)
print(
    f"Workbook SHA-256: "
    f"{WORKBOOK_SHA256}"
)

print()
print("=" * 110)
print("OUTPUT")
print("=" * 110)

print(
    f"Workbook: {OUTPUT_XLSX.resolve()}"
)
print(
    f"Workbook checksum sidecar: "
    f"{WORKBOOK_HASH_FILE.resolve()}"
)
print()
print(
    "FERPA NOTICE: Do not upload the workbook here. "
    "The console output above is student-blind and safe to return."
)
print()

