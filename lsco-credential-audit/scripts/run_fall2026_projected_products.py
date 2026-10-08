from __future__ import annotations
"""Products 4 and 5: Fall 2026 projected path-aware curriculum audit.

Run from lsco-credential-audit root:
  python -u .\\scripts\\run_fall2026_projected_products.py --limit-students 5
  python -u .\\scripts\\run_fall2026_projected_products.py

Local FERPA records only. Uses exactly the established path-aware engine and
recomputes both earned and prospective course allocations. Fall grades are
provisional C for curricular matching; never used as evidence of actual GPA,
graduation or an official award. Catalog and award governance still apply.
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT/"scripts"))
sys.path.insert(0,str(ROOT/"src"))
import run_fall_2026_incremental_audit as base
import profile_refresh_delta as refresh
from catalog_temporal_policy import catalog_temporal_fields, EVALUATION_DATE

OUT=ROOT/"data/processed/reporting/fall2026_projected_products_v2"
COLUMNS=["Banner ID","Last Name","First Name","Declared Major","Catalog Year",
 "Credential ID","Whole Classes Remaining","Missing Requirement 1","Allowed Courses 1",
 "Missing Requirement 2","Allowed Courses 2","Missing Requirement 3","Allowed Courses 3",
 "Fall Courses","Projected Term","Status / Qualification"]

def resolve_missing(row, requirements, support, occupied):
    """Resolve one selected path requirement into whole-course obligations.

    No SCH subtraction. A group with min_required=2 and one applied course
    leaves one obligation; compound courses are counted independently.
    Elective semantics are deliberately held for explicit course-hour solving.
    """
    rid=str(row["requirement_id"])
    group=requirements[requirements.requirement_id.astype(str).eq(rid)]
    if group.empty:
        return None,"MISSING_REQUIREMENT_DEFINITION"
    types=set(group.option_type.astype(str).str.upper().str.strip())
    if "ELECTIVE" in types:
        return None,"ELECTIVE_WHOLE_CLASS_SOLVER_REQUIRED"
    alternatives=support["compound_alternatives"].get(rid,[])
    if alternatives:
        viable=[]
        for alternative in alternatives:
            courses=[base.engine.normalize_course_code(x) for x in alternative]
            remaining=[c for c in courses if c not in occupied]
            viable.append(remaining)
        if not viable: return None,"COMPOUND_NO_VALID_ALTERNATIVE"
        fewest=min(map(len,viable))
        best=sorted(tuple(x) for x in viable if len(x)==fewest)
        # Different minimal compound paths have different class combinations,
        # not equivalent one-class options. Hold ambiguity for review.
        if len(set(best))!=1:return None,"COMPOUND_ALTERNATIVES_REQUIRE_REVIEW"
        return [(rid,course) for course in best[0]],"OK"
    candidates=set()
    for _,option in group.iterrows():
        kind=str(option.option_type).upper().strip()
        value=base.engine.normalize_course_code(option.option_value)
        if kind=="COURSE":
            candidates.add(value)
        elif kind=="CORE_BUCKET":
            bucket=base.engine.normalize_bucket_name(value)
            if not bucket:return None,"CORE_BUCKET_UNRECOGNIZED"
            key=(str(row["catalog_year"]),str(bucket))
            if key not in support["core_lookup"]:
                return None,"CORE_BUCKET_LOOKUP_MISSING"
            candidates.update(support["core_lookup"][key])
        else:
            return None,"UNSUPPORTED_OPTION_TYPE_"+kind
    candidates={x for x in candidates if x and x not in occupied}
    if not candidates:return None,"NO_UNUSED_ALLOWED_COURSES"
    required=int(float(group.iloc[0]["min_required"]))
    matched={base.engine.normalize_course_code(x) for x in
        str(row.get("matched_options","")).split(";") if x.strip()}
    needed=max(0,required-len(matched))
    if needed==0:
        return None,"UNMET_BUT_NO_ADDITIONAL_COURSE_NEEDED_REVIEW"
    # Multiple obligations in one catalog requirement cannot be counted
    # as one merely because they share an ID.
    return [(rid,"; ".join(sorted(candidates)))]*needed,"OK"

def resolve_plan(detail, requirements, support):
    counted=chosen(detail.to_dict("records"))
    applied=set()
    for text in counted.loc[counted.status.eq("MET"),"matched_options"].fillna(""):
        applied.update(base.engine.normalize_course_code(x) for x in str(text).split(";") if x.strip())
    needed=[]
    for _,row in counted.loc[~counted.status.eq("MET")].iterrows():
        if str(row.status).startswith("UNRESOLVED"):
            return None,"ENGINE_UNRESOLVED_"+str(row.status)
        items,reason=resolve_missing(row,requirements,support,applied)
        if items is None:return None,reason
        needed.extend(items)
    return needed,"OK"

def chosen(rows):
    df=pd.DataFrame(rows)
    selected=df["path_group_id"].fillna("").astype(str).str.strip().eq("") | df["path_selected"].fillna("").astype(str).str.upper().eq("TRUE")
    return df[selected].copy()

def one_audit(sid,year,cred,req,courses,support):
    lookup=base.engine.completed_course_lookup(courses)
    rows=base.engine.audit_student_credential(
        student_id=sid,catalog_year=year,credential_id=cred,
        credential_requirements=req,course_lookup=lookup,
        core_lookup=support["core_lookup"],elective_rules=support["elective_rules"],
        academic_course_lookup=support["academic_course_lookup"],
        compound_alternatives=support["compound_alternatives"],
        alternative_path_lookup=support["alternative_path_lookup"])
    return chosen(rows)

def write_sheet(wb,name,head,records):
    sh=wb.create_sheet(name);sh.append(head)
    for record in records:sh.append(record)
    sh.freeze_panes="A2"
    sh.auto_filter.ref=f"A1:{get_column_letter(len(head))}{len(records)+1}"
    sh.sheet_view.showGridLines=False
    for c in sh[1]:
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="145A32")
    for col in range(1,len(head)+1):
        sh.column_dimensions[get_column_letter(col)].width=(50 if "Courses" in head[col-1] or "Requirement" in head[col-1] else 23)
    return sh

def make_book(path,records,title):
    wb=Workbook();wb.active.title="CONTROL SUMMARY"
    sh=wb.active
    for row in [
        ["LSCO Fall 2026",title],
        ["Scope","Fall 2026 (202690) registered students only"],
        ["Source","Path-aware six-year catalog reallocation, projected successful Fall courses"],
        ["Total recommendation rows",len(records)],
        ["Distinct students",len({r[0] for r in records})],
        ["Condition","Curricular projection only; final grade, residency, GPA and official-award checks pending"],
        ["As-of",str(EVALUATION_DATE)],
    ]:sh.append(row)
    groups=defaultdict(list)
    for row in records:groups[str(row[4])].append(row)
    for year,items in sorted(groups.items()):
        write_sheet(wb,year,COLUMNS,sorted(items,key=lambda x:(x[1],x[2],x[0],x[5])))
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--limit-students",type=int,default=None)
    parser.add_argument("--force",action="store_true")
    args=parser.parse_args()
    assert Path.cwd().resolve()==ROOT,"Run from lsco-credential-audit root"
    requirements=pd.read_csv(base.REQUIREMENTS,dtype=str).fillna("")
    base.configure_engine_paths()
    support=base.load_engine_support(requirements)
    state=base.load_authoritative_attempt_state()
    fall=state["current_fall"].copy()
    assert set(fall["grade"].astype(str).str.strip())=={""},"Expected blank Fall final grades; check source"
    fall_ids=set(fall.student_id.astype(str))
    assert len(fall_ids)>=8000,"Fall cohort unexpectedly small; verify term 202690"
    eligible=base.build_refreshed_eligibility(state["affected"],sorted(requirements.catalog_year.unique()))
    eligible=eligible[eligible.catalog_eligible.eq(True)&eligible.student_id.isin(fall_ids)]
    catalog_pairs=eligible[["student_id","catalog_year"]].drop_duplicates()
    # Explicit policy screen: no historical expiry or source-unknown catalogs.
    # Catalog 2026-27 analyst-approved variance is retained with disclosure.
    allowed=[]
    for year in catalog_pairs.catalog_year.unique():
        status=catalog_temporal_fields(year)
        if status.get("conferral_temporal_eligibility","") in ("PASS_VERIFIED_POLICY","PASS_ANALYST_APPROVED_VARIANCE"):
            allowed.append(year)
    catalog_pairs=catalog_pairs[catalog_pairs.catalog_year.isin(allowed)]
    reqgroups={(str(y),str(c)):g.copy() for (y,c),g in requirements.groupby(["catalog_year","credential_id"])}
    targets=defaultdict(list)
    for sid,year in catalog_pairs.itertuples(index=False,name=None):
        targets[str(sid)].extend((str(year),cred,g) for (y,cred),g in reqgroups.items() if y==str(year))
    passing=base.build_passing_course_view(state["affected"])
    students=sorted(fall_ids)
    if args.limit_students is not None:students=students[:args.limit_students]
    smoke=args.limit_students is not None
    OUT.mkdir(parents=True,exist_ok=True)
    chunks=OUT/"student_chunks";chunks.mkdir(exist_ok=True)
    print(f"Fall cohort: {len(fall_ids)}; eligible target pairs: {sum(len(v) for v in targets.values())}",flush=True)
    for idx,sid in enumerate(students,1):
        out=chunks/(sid.replace("/","_")+".json")
        if out.exists() and not args.force:continue
        earned=passing[passing.student_id.eq(sid)].copy()
        enroll=fall[fall.student_id.eq(sid)].copy()
        synthetic=enroll.copy()
        synthetic["grade"]="C"
        synthetic["source_dataset"]="FALL_2026_PROJECTED_C_NOT_EARNED"
        # Normalize prospective attempts through the very same engine pathway.
        projected_raw=pd.concat([state["affected"][state["affected"].student_id.eq(sid) & ~pd.to_numeric(state["affected"].term_sort,errors="coerce").eq(refresh.FALL_2026)],synthetic])
        projected=base.build_passing_course_view(projected_raw)
        info=enroll.iloc[0]
        fields=[sid,str(info.get("last_name","")),str(info.get("first_name","")),str(info.get("student_major",""))]
        fc="; ".join(sorted(set(base.normalize_course_code(x.subject,x.course_number) for x in enroll.itertuples())))
        result=[]
        unresolved_plans=[]
        for year,cred,req in targets.get(sid,[]):
            pre=one_audit(sid,year,cred,req,earned,support)
            post=one_audit(sid,year,cred,req,projected,support)
            deficits,reason=resolve_plan(post,req,support)
            if deficits is None:
                unresolved_plans.append({"year":year,"credential":cred,"reason":reason})
                continue
            if len(deficits)>3:continue
            before_met=bool(pre.status.eq("MET").all())
            if len(deficits)==0:
                if before_met:continue # already complete before Fall, not newly projected
                status="PROJECTED_CURRICULAR_COMPLETE_NOT_CONFERRED"
            else:
                status="MISSING_WHOLE_CLASS_OBLIGATIONS"
            padded=(deficits+[("","")]*3)[:3]
            result.append({"student":sid,"year":year,"credential":cred,"remaining":len(deficits),
                "fields":fields,"fall_courses":fc,"deficits":padded,"status":status,
                "note":"Projection assumes C in active Fall courses; academic governance pending"})
        tmp=out.with_suffix(".tmp")
        tmp.write_text(json.dumps({"plans":result,"unresolved":unresolved_plans}),encoding="utf-8")
        tmp.replace(out)
        if idx%25==0:print(f"Processed {idx}/{len(students)} Fall students",flush=True)
    # Do not publish partial outputs as final.
    present={p.stem for p in chunks.glob("*.json")}
    expected={sid.replace("/","_") for sid in students}
    assert expected.issubset(present),"Missing chunk results; cannot publish"
    p4=[];best={};provisional={};held_students=set();held_plans=Counter()
    for sid in students:
        payload=json.loads((chunks/(sid.replace("/","_")+".json")).read_text(encoding="utf-8"))
        rows=payload["plans"]
        if payload["unresolved"]:
            held_students.add(sid)
            held_plans.update(r["reason"] for r in payload["unresolved"])
        if any(r["remaining"]==0 for r in rows):
            for r in rows:
                if r["remaining"]==0:p4.append(r)
            continue
        matches=[r for r in rows if 1<=r["remaining"]<=3]
        if matches:
            low=min(x["remaining"] for x in matches)
            ties=[r for r in matches if r["remaining"]==low]
            if sid in held_students:
                provisional[sid]=ties
            else:
                best[sid]=ties
    def layout(r):
        d=r["deficits"]
        return r["fields"][:]+[r["year"],r["credential"],r["remaining"],
            d[0][0],d[0][1],d[1][0],d[1][1],d[2][0],d[2][1],
            r["fall_courses"],"202690",r["status"]]
    near=[r for rows in best.values() for r in rows]
    candidate_review=[r for rows in provisional.values() for r in rows]
    workbook_dir=OUT/("SMOKE_TEST_NOT_FOR_RELEASE" if smoke else "FULL_COHORT")
    make_book(workbook_dir/"RESTRICTED_Product4_Fall2026_Projected_Curriculum_Completions.xlsx",
              [layout(r) for r in p4],"New projected curricular completions")
    make_book(workbook_dir/"RESTRICTED_Product5_Fall2026_Near_Completers_1_to_3_Classes.xlsx",
              [layout(r) for r in near],"Closest curricula missing 1–3 whole classes")
    make_book(workbook_dir/"RESTRICTED_Product5_PROVISIONAL_Candidates_Not_Ranked_For_Release.xlsx",
              [layout(r) for r in candidate_review],"Validated candidates; other credential plans unresolved")
    print(f"Provisional candidate plans: {len(candidate_review)} / students: {len(provisional)}")
    print(f"Product 4 plans: {len(p4)} / students: {len({r['student'] for r in p4})}")
    print(f"Product 5 plans: {len(near)} / students: {len(best)}")
    print(f"Students held from nearest-plan ranking due to unresolved options: {len(held_students)}")
    print("Unresolved reasons:",dict(held_plans))
    print("Whole-class counts use min_required or compound lengths; elective SCH-to-course solver is NOT implemented.")
    print("Do not release Product 5 as complete while held students exist.")
    print("Files:",workbook_dir)
    if smoke:print("SMOKE TEST ONLY — not a full cohort release.")

if __name__=="__main__":main()
