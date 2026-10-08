from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path.cwd()

TARGET_CATALOG = "2021-2022"
TARGET_START_YEAR = 2021

FINAL_DETAIL = ROOT / (
    "data/processed/reporting/"
    "final_additional_award_numbers_20260803/"
    "01_final_additional_awards_student_detail.csv"
)

LATER_ALTS = ROOT / (
    "data/processed/reporting/"
    "additional_awards_clean_query_20260803/"
    "05_all_eligible_catalog_alternatives.csv"
)

IDENTITY_SOURCE = ROOT / (
    "data/raw/student_exports/"
    "banner_course_history_actual.csv"
)

OUTDIR = ROOT / (
    "data/processed/reporting/"
    "expiring_2021_2022_award_reconciliation_20260811"
)

OUTFILE = OUTDIR / (
    "2021_2022_EXPIRING_AWARD_RECONCILIATION.xlsx"
)


def fail(msg):
    raise RuntimeError(msg)


def norm(v):
    return re.sub(
        r"[^A-Z0-9]+",
        "_",
        str(v).strip().upper()
    ).strip("_")


def norm_id(v):
    return str(v).strip()


def cat_year(v):
    m = re.search(r"(20\d{2})", str(v))
    return int(m.group(1)) if m else -1


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b""
        ):
            h.update(chunk)
    return h.hexdigest()


def find_col(
    cols,
    aliases,
    required=True,
    label="column"
):
    lookup = {
        str(c).strip().lower(): c
        for c in cols
    }

    for alias in aliases:
        if alias.lower() in lookup:
            return lookup[alias.lower()]

    if required:
        fail(
            f"Missing {label}. "
            f"Tried {aliases}. "
            f"Available: {list(cols)}"
        )

    return None


def lineage_col(df):
    return find_col(
        df.columns,
        [
            "credential_lineage",
            "detected_lineage",
            "canonical_lineage"
        ],
        True,
        "credential lineage"
    )


def strongest_nonblank(
    df,
    student_col,
    value_col,
    term_col=None
):
    t = df.copy()

    t = t[
        t[value_col]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
    ].copy()

    if term_col:
        t["_term_sort"] = pd.to_numeric(
            t[term_col],
            errors="coerce"
        )
    else:
        t["_term_sort"] = range(len(t))

    t["_row"] = range(len(t))

    t = t.sort_values(
        [student_col, "_term_sort", "_row"]
    )

    return (
        t.groupby(student_col, as_index=False)
        .tail(1)[[student_col, value_col]]
    )


# =====================================================================
# 0. PREFLIGHT
# =====================================================================

print("=" * 105)
print("2021-2022 EXPIRING AWARD — OFFICIAL AWARD RECONCILIATION")
print("=" * 105)

for p in [
    FINAL_DETAIL,
    LATER_ALTS,
    IDENTITY_SOURCE
]:
    if not p.exists():
        fail(f"Missing required source: {p}")

OUTDIR.mkdir(
    parents=True,
    exist_ok=True
)


# =====================================================================
# 1. RECONSTRUCT EXACT PRIOR EXPIRING SET
# =====================================================================

final = pd.read_csv(
    FINAL_DETAIL,
    dtype=str,
    low_memory=False
).fillna("")

flin = lineage_col(final)

final["student_id"] = (
    final["student_id"].map(norm_id)
)
final["_lineage"] = (
    final[flin].map(norm)
)

target = final[
    final["catalog_year"]
    .astype(str)
    .str.strip()
    .eq(TARGET_CATALOG)
].copy()

if target.empty:
    fail("No 2021-2022 frozen candidate rows.")

alts = pd.read_csv(
    LATER_ALTS,
    dtype=str,
    low_memory=False
).fillna("")

alin = lineage_col(alts)

alts["student_id"] = (
    alts["student_id"].map(norm_id)
)
alts["_lineage"] = (
    alts[alin].map(norm)
)
alts["_cy"] = (
    alts["catalog_year"].map(cat_year)
)

if "audit_status" in alts.columns:
    alts = alts[
        alts["audit_status"]
        .astype(str)
        .str.upper()
        .eq("COMPLETE")
    ].copy()

later = alts[
    alts["_cy"] > TARGET_START_YEAR
].copy()

later_keys = (
    later[
        ["student_id", "_lineage"]
    ]
    .drop_duplicates()
)

later_keys["_later"] = True

scored = target.merge(
    later_keys,
    on=["student_id", "_lineage"],
    how="left"
)

later_mask = (
    scored["_later"]
    .fillna(False)
    .astype(bool)
)

prior_expiring = scored[
    ~later_mask
].copy()

print(
    f"Prior expiring rows reconstructed: "
    f"{len(prior_expiring):,}"
)
print(
    f"Prior expiring students: "
    f"{prior_expiring.student_id.nunique():,}"
)
print()


# =====================================================================
# 2. DISCOVER OFFICIAL AWARD HISTORY SOURCE
#
# We deliberately do NOT trust the downstream unawarded flag.
# Search repo CSV headers for a real award-history style source.
# =====================================================================

ID_ALIASES = [
    "student_id",
    "studentid",
    "studenid",
    "banner_id",
    "pidm",
    "student_pidm"
]

LINEAGE_ALIASES = [
    "credential_lineage",
    "canonical_lineage",
    "detected_lineage"
]

AWARD_ID_ALIASES = [
    "credential_id",
    "credential_base_id",
    "award_code",
    "degree_code",
    "degree",
    "program_code",
    "curriculum_code"
]

AWARD_TITLE_ALIASES = [
    "credential_title",
    "award_title",
    "degree_title",
    "program_title",
    "degree_desc",
    "award_description"
]

AWARD_EVIDENCE_ALIASES = [
    "award_term",
    "awarded_term",
    "degree_term",
    "graduation_term",
    "conferral_term",
    "award_date",
    "degree_date",
    "graduation_date",
    "conferral_date"
]


def has_alias(cols, aliases):
    low = {
        str(c).strip().lower()
        for c in cols
    }
    return any(
        a.lower() in low
        for a in aliases
    )


candidates = []

search_roots = [
    ROOT / "data/raw",
    ROOT / "data/interim",
    ROOT / "data/processed"
]

for base in search_roots:
    if not base.exists():
        continue

    for p in base.rglob("*.csv"):

        # Exclude our known downstream candidate/report files.
        plow = str(p).lower()

        if (
            "additional_award_numbers" in plow
            or "additional_awards_clean_query" in plow
            or "expiring_2021_2022" in plow
        ):
            continue

        try:
            cols = list(
                pd.read_csv(
                    p,
                    nrows=0
                ).columns
            )
        except Exception:
            continue

        has_id = has_alias(
            cols,
            ID_ALIASES
        )

        has_lineage = has_alias(
            cols,
            LINEAGE_ALIASES
        )

        has_award_id = has_alias(
            cols,
            AWARD_ID_ALIASES
        )

        has_title = has_alias(
            cols,
            AWARD_TITLE_ALIASES
        )

        has_evidence = has_alias(
            cols,
            AWARD_EVIDENCE_ALIASES
        )

        if not has_id:
            continue

        if not (
            has_lineage
            or has_award_id
            or has_title
        ):
            continue

        score = 0

        score += 5
        score += 5 if has_evidence else 0
        score += 4 if has_lineage else 0
        score += 3 if has_award_id else 0
        score += 2 if has_title else 0

        name = p.name.lower()

        if "award" in name:
            score += 3

        if "degree" in name:
            score += 3

        if "graduat" in name:
            score += 2

        if "official" in name:
            score += 4

        if "modeled" in name:
            score -= 5

        if "candidate" in name:
            score -= 3

        candidates.append(
            (
                score,
                p,
                cols,
                has_evidence,
                has_lineage,
                has_award_id,
                has_title
            )
        )


candidates.sort(
    key=lambda x: x[0],
    reverse=True
)

print("=" * 105)
print("AWARD SOURCE DISCOVERY")
print("=" * 105)

for i, c in enumerate(
    candidates[:10],
    start=1
):
    score, p, cols, ev, lin, aid, title = c

    print(
        f"{i:2d}. score={score:2d} | "
        f"{p.relative_to(ROOT)}"
    )
    print(
        f"    evidence={ev} "
        f"lineage={lin} "
        f"award_id={aid} "
        f"title={title}"
    )

if not candidates:
    fail(
        "No plausible official award-history CSV "
        "was found."
    )

best = candidates[0]

# Require clear separation from runner-up.
if len(candidates) > 1:
    if best[0] - candidates[1][0] < 2:
        fail(
            "\nAward source discovery is ambiguous. "
            "I am refusing to guess.\n"
            "Paste the AWARD SOURCE DISCOVERY block "
            "back to ChatGPT and we will pin the "
            "authoritative source explicitly."
        )

award_source = best[1]

print()
print(
    "SELECTED OFFICIAL-AWARD SOURCE:"
)
print(
    award_source.relative_to(ROOT)
)
print()


# =====================================================================
# 3. LOAD OFFICIAL AWARDS
# =====================================================================

awards = pd.read_csv(
    award_source,
    dtype=str,
    low_memory=False
).fillna("")

sid = find_col(
    awards.columns,
    ID_ALIASES,
    True,
    "award-history student ID"
)

a_lineage = find_col(
    awards.columns,
    LINEAGE_ALIASES,
    False
)

a_id = find_col(
    awards.columns,
    AWARD_ID_ALIASES,
    False
)

a_title = find_col(
    awards.columns,
    AWARD_TITLE_ALIASES,
    False
)

a_term = find_col(
    awards.columns,
    AWARD_EVIDENCE_ALIASES,
    False
)

awards["_student_id"] = (
    awards[sid].map(norm_id)
)

need_students = set(
    prior_expiring["student_id"]
)

awards = awards[
    awards["_student_id"]
    .isin(need_students)
].copy()

if awards.empty:
    fail(
        "Selected award-history source has zero "
        "rows for the prior expiring students."
    )

if a_lineage:
    awards["_award_lineage"] = (
        awards[a_lineage].map(norm)
    )
else:
    awards["_award_lineage"] = ""

if a_id:
    awards["_award_id"] = (
        awards[a_id].map(norm)
    )
else:
    awards["_award_id"] = ""

if a_title:
    awards["_award_title"] = (
        awards[a_title]
        .astype(str)
        .str.strip()
    )
    awards["_award_title_norm"] = (
        awards[a_title].map(norm)
    )
else:
    awards["_award_title"] = ""
    awards["_award_title_norm"] = ""


# =====================================================================
# 4. TARGET CREDENTIAL FIELDS
# =====================================================================

t_title_col = find_col(
    prior_expiring.columns,
    [
        "credential_title",
        "program_title",
        "title"
    ],
    False
)

t_id_col = find_col(
    prior_expiring.columns,
    [
        "credential_id",
        "credential_base_id"
    ],
    False
)

if t_title_col:
    prior_expiring["_target_title"] = (
        prior_expiring[t_title_col]
        .astype(str)
        .str.strip()
    )
else:
    prior_expiring["_target_title"] = (
        prior_expiring["_lineage"]
        .str.replace(
            "_",
            " ",
            regex=False
        )
        .str.title()
    )

prior_expiring[
    "_target_title_norm"
] = prior_expiring[
    "_target_title"
].map(norm)

if t_id_col:
    prior_expiring["_target_id"] = (
        prior_expiring[t_id_col].map(norm)
    )
else:
    prior_expiring["_target_id"] = ""


# =====================================================================
# 5. COMPARE TARGET TO RAW OFFICIAL AWARDS
#
# NO FUZZY MATCHING.
#
# Match methods:
# 1. canonical lineage exact
# 2. credential/degree ID exact
# 3. normalized official title exact
# =====================================================================

results = []

for _, t in prior_expiring.iterrows():

    sid_value = t["student_id"]

    student_awards = awards[
        awards["_student_id"].eq(
            sid_value
        )
    ].copy()

    all_titles = []

    for _, a in student_awards.iterrows():

        label = (
            a["_award_title"]
            if a["_award_title"]
            else (
                a["_award_id"]
                if a["_award_id"]
                else a["_award_lineage"]
            )
        )

        if a_term and str(
            a[a_term]
        ).strip():
            label = (
                f"{label} "
                f"[{str(a[a_term]).strip()}]"
            )

        if label:
            all_titles.append(label)

    all_titles = sorted(
        set(all_titles)
    )

    matches = []

    for _, a in student_awards.iterrows():

        method = None

        if (
            t["_lineage"]
            and a["_award_lineage"]
            and t["_lineage"]
            == a["_award_lineage"]
        ):
            method = "CANONICAL_LINEAGE_EXACT"

        elif (
            t["_target_id"]
            and a["_award_id"]
            and t["_target_id"]
            == a["_award_id"]
        ):
            method = "CREDENTIAL_ID_EXACT"

        elif (
            t["_target_title_norm"]
            and a["_award_title_norm"]
            and t["_target_title_norm"]
            == a["_award_title_norm"]
        ):
            method = "TITLE_EXACT"

        if method:
            matches.append(
                (
                    method,
                    a["_award_title"],
                    a["_award_id"],
                    a["_award_lineage"]
                )
            )

    same_award_found = (
        len(matches) > 0
    )

    match_methods = sorted(
        {
            m[0]
            for m in matches
        }
    )

    results.append(
        {
            "student_id":
                sid_value,
            "_lineage":
                t["_lineage"],
            "Target Credential":
                t["_target_title"],
            "Target Credential ID":
                t["_target_id"],
            "Catalog Year":
                t["catalog_year"],
            "Official Awarded Credentials":
                " | ".join(all_titles),
            "Same-Credential Official Award Found":
                "YES"
                if same_award_found
                else "NO",
            "Official Match Method":
                " | ".join(match_methods),
            "Truly Unawarded Credential":
                ""
                if same_award_found
                else t["_target_title"]
        }
    )


recon = pd.DataFrame(results)


# =====================================================================
# 6. LOCAL NAME + MAJOR
# =====================================================================

ih = pd.read_csv(
    IDENTITY_SOURCE,
    nrows=0
)

i_id = find_col(
    ih.columns,
    [
        "StudenID",
        "StudentID",
        "student_id"
    ]
)

i_first = find_col(
    ih.columns,
    [
        "FirstName",
        "first_name"
    ]
)

i_last = find_col(
    ih.columns,
    [
        "LastName",
        "last_name"
    ]
)

i_major = find_col(
    ih.columns,
    [
        "StudentMajor",
        "student_major",
        "major"
    ]
)

i_term = find_col(
    ih.columns,
    [
        "Term",
        "term",
        "term_taken"
    ],
    False
)

ident = pd.read_csv(
    IDENTITY_SOURCE,
    usecols=[
        x for x in [
            i_id,
            i_first,
            i_last,
            i_major,
            i_term
        ]
        if x
    ],
    dtype=str,
    low_memory=False
).fillna("")

ident["_student_id"] = (
    ident[i_id].map(norm_id)
)

ident = ident[
    ident["_student_id"].isin(
        need_students
    )
].copy()

latest_first = strongest_nonblank(
    ident,
    "_student_id",
    i_first,
    i_term
)

latest_last = strongest_nonblank(
    ident,
    "_student_id",
    i_last,
    i_term
)

latest_major = strongest_nonblank(
    ident,
    "_student_id",
    i_major,
    i_term
)

identity = (
    pd.DataFrame(
        {
            "_student_id":
                sorted(need_students)
        }
    )
    .merge(
        latest_first,
        on="_student_id",
        how="left"
    )
    .merge(
        latest_last,
        on="_student_id",
        how="left"
    )
    .merge(
        latest_major,
        on="_student_id",
        how="left"
    )
    .fillna("")
)

recon = (
    recon
    .merge(
        identity,
        left_on="student_id",
        right_on="_student_id",
        how="left"
    )
)

recon = pd.DataFrame(
    {
        "Banner ID":
            recon["student_id"],
        "Last Name":
            recon[i_last],
        "First Name":
            recon[i_first],
        "Declared Major":
            recon[i_major],
        "Catalog Year":
            recon["Catalog Year"],
        "Target Credential":
            recon["Target Credential"],
        "Target Credential ID":
            recon["Target Credential ID"],
        "Official Awarded Credentials":
            recon[
                "Official Awarded Credentials"
            ],
        "Same-Credential Official Award Found":
            recon[
                "Same-Credential Official Award Found"
            ],
        "Official Match Method":
            recon["Official Match Method"],
        "Truly Unawarded Credential":
            recon[
                "Truly Unawarded Credential"
            ]
    }
)

recon = recon.sort_values(
    [
        "Same-Credential Official Award Found",
        "Target Credential",
        "Last Name",
        "First Name"
    ],
    ascending=[
        True,
        True,
        True,
        True
    ]
).reset_index(drop=True)


# =====================================================================
# 7. CHECKS
# =====================================================================

already_awarded = int(
    recon[
        "Same-Credential Official Award Found"
    ].eq("YES").sum()
)

truly_unawarded = int(
    recon[
        "Truly Unawarded Credential"
    ]
    .astype(str)
    .str.strip()
    .ne("")
    .sum()
)

students_any_award = int(
    recon[
        "Official Awarded Credentials"
    ]
    .astype(str)
    .str.strip()
    .ne("")
    .sum()
)

if (
    already_awarded
    + truly_unawarded
    != len(recon)
):
    fail(
        "Reconciliation does not partition "
        "the prior candidate set."
    )

# Hard invariant.
bad = recon[
    recon[
        "Same-Credential Official Award Found"
    ].eq("YES")
    &
    recon[
        "Truly Unawarded Credential"
    ].astype(str).str.strip().ne("")
]

if not bad.empty:
    fail(
        "Hard invariant failed: a matched "
        "official award was also labeled unawarded."
    )


# =====================================================================
# 8. WORKBOOK — TWO SHEETS
# =====================================================================

GREEN = "00573F"
ORANGE = "E57200"
WHITE = "FFFFFF"
LIGHT = "E8F1ED"

wb = Workbook()

ws = wb.active
ws.title = "CONTROL_SUMMARY"
ws.sheet_view.showGridLines = False

detail = wb.create_sheet(
    "AWARD_RECONCILIATION"
)
detail.sheet_view.showGridLines = False


ws.merge_cells("A1:D1")
ws["A1"] = (
    "2021–2022 Expiring Award Reconciliation"
)
ws["A1"].fill = PatternFill(
    "solid",
    fgColor=GREEN
)
ws["A1"].font = Font(
    bold=True,
    color=WHITE,
    size=16
)

ws.merge_cells("A2:D2")
ws["A2"] = (
    "Diagnostic rebuild using raw/official award "
    "history. Downstream 'unawarded' status is not "
    "accepted without direct reconciliation."
)
ws["A2"].font = Font(
    italic=True
)

rows = [
    (
        "Run Timestamp",
        datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    ),
    (
        "Target Catalog",
        TARGET_CATALOG
    ),
    (
        "Prior Expiring Candidate Rows",
        len(recon)
    ),
    (
        "Candidates With Any Official Award",
        students_any_award
    ),
    (
        "Same Credential Already Officially Awarded",
        already_awarded
    ),
    (
        "TRUE Remaining Unawarded",
        truly_unawarded
    ),
    (
        "Partition Check",
        (
            f"{already_awarded} + "
            f"{truly_unawarded} = "
            f"{len(recon)}"
        )
    ),
    (
        "Award Source",
        str(
            award_source.relative_to(ROOT)
        )
    ),
    (
        "Award Source SHA-256",
        sha(award_source)
    ),
    (
        "Frozen Final Source SHA-256",
        sha(FINAL_DETAIL)
    )
]

r = 4

for label, value in rows:
    ws.cell(r, 1).value = label
    ws.cell(r, 1).font = Font(
        bold=True
    )
    ws.cell(r, 2).value = value

    if label in [
        "Same Credential Already Officially Awarded",
        "TRUE Remaining Unawarded"
    ]:
        ws.cell(r, 1).fill = PatternFill(
            "solid",
            fgColor=LIGHT
        )
        ws.cell(r, 2).fill = PatternFill(
            "solid",
            fgColor=LIGHT
        )

    r += 1

r += 2

ws.cell(r, 1).value = (
    "CONTROL INTERPRETATION"
)
ws.cell(r, 1).font = Font(
    bold=True,
    color=GREEN
)

r += 1

ws.merge_cells(
    start_row=r,
    start_column=1,
    end_row=r + 3,
    end_column=4
)

ws.cell(r, 1).value = (
    "A target credential is removed from the "
    "'Truly Unawarded Credential' column only "
    "when the official award history contains "
    "an exact canonical-lineage, credential-ID, "
    "or normalized-title match. No fuzzy or "
    "semantic matching is used."
)

ws.cell(r, 1).alignment = Alignment(
    wrap_text=True,
    vertical="top"
)

ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 75
ws.column_dimensions["C"].width = 25
ws.column_dimensions["D"].width = 25


# Detail sheet
detail.append(
    list(recon.columns)
)

for row in recon.itertuples(
    index=False,
    name=None
):
    detail.append(list(row))

for cell in detail[1]:
    cell.fill = PatternFill(
        "solid",
        fgColor=GREEN
    )
    cell.font = Font(
        bold=True,
        color=WHITE
    )
    cell.alignment = Alignment(
        wrap_text=True
    )

detail.freeze_panes = "A2"

if len(recon):
    ref = (
        f"A1:"
        f"{get_column_letter(len(recon.columns))}"
        f"{len(recon) + 1}"
    )

    table = Table(
        displayName="AwardReconciliation",
        ref=ref
    )

    table.tableStyleInfo = (
        TableStyleInfo(
            name="TableStyleMedium4",
            showRowStripes=True
        )
    )

    detail.add_table(table)

widths = [
    15, 20, 18, 32, 14,
    38, 28, 70, 22, 28, 38
]

for i, width in enumerate(
    widths,
    start=1
):
    detail.column_dimensions[
        get_column_letter(i)
    ].width = width

for row in detail.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

wb.save(OUTFILE)


# =====================================================================
# 9. STUDENT-BLIND CONSOLE
# =====================================================================

print()
print("=" * 105)
print("RECONCILIATION FUNNEL — STUDENT-BLIND")
print("=" * 105)

print(
    f"Prior expiring candidates: "
    f"{len(recon):,}"
)

print(
    f"Candidates with any official award history: "
    f"{students_any_award:,}"
)

print(
    f"Same target credential already awarded: "
    f"{already_awarded:,}"
)

print(
    f"TRUE remaining unawarded: "
    f"{truly_unawarded:,}"
)

print(
    "Partition check: "
    f"{already_awarded:,} + "
    f"{truly_unawarded:,} = "
    f"{len(recon):,}"
)

print()
print(
    "Selected official award source:"
)
print(
    award_source.relative_to(ROOT)
)

print()
print(
    "Award source SHA-256:"
)
print(
    sha(award_source)
)

print()
print(
    "Workbook:"
)
print(
    OUTFILE.resolve()
)

print()
print(
    "FERPA: workbook contains student identities. "
    "Do not upload it here."
)

