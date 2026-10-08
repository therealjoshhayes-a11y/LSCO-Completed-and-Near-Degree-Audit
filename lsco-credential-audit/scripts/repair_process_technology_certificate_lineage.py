from pathlib import Path
from datetime import datetime
import shutil
import re
import pandas as pd
from openpyxl import load_workbook

ROOT = Path.cwd()
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

COMPLETE = (
    ROOT
    / "data/processed/full_actual_audit/president_report/"
      "complete_rows_with_canonical_lineage.csv"
)

CROSSWALK = (
    ROOT
    / "data/interim/institutional_awards/"
      "award_program_crosswalk_curated.xlsx"
)

TARGET = "PROCESS_TECHNOLOGY_CERTIFICATE"
OLD = "PROCESS_TECHNOLOGY_BASIC"

CERT_IDS = {
    f"PROCESS_TECHNOLOGY_{year}"
    for year in range(2021, 2026)
}

AAS_IDS = {
    f"PROCESS_OPERATING_TECHNOLOGY_{year}"
    for year in range(2021, 2026)
}

BUILDERS = [
    ROOT / "scripts/build_president_report_data.py",
    ROOT / "scripts/build_president_report_data_v2.py",
    ROOT / "scripts/build_president_report_data_v3.py",
    ROOT / "scripts/build_president_report_data_corrected.py",
]


def backup(path: Path, label: str) -> Path:
    dst = path.with_name(
        f"{path.stem}_PRE_{label}_{STAMP}{path.suffix}"
    )
    shutil.copy2(path, dst)
    return dst


print("=" * 100)
print("PROCESS TECHNOLOGY CERTIFICATE — CANONICAL LINEAGE REPAIR")
print("=" * 100)


# =============================================================================
# 1. CURRENT MODELED COMPLETION LINEAGE
# =============================================================================

if not COMPLETE.exists():
    raise FileNotFoundError(COMPLETE)

df = pd.read_csv(
    COMPLETE,
    dtype=str,
    low_memory=False,
).fillna("")

required = {
    "credential_id",
    "canonical_lineage",
}

missing = required - set(df.columns)

if missing:
    raise RuntimeError(
        f"Completion file missing columns: {sorted(missing)}"
    )

cert_mask = df["credential_id"].isin(CERT_IDS)
aas_mask = df["credential_id"].isin(AAS_IDS)

found_cert_ids = set(
    df.loc[cert_mask, "credential_id"].unique()
)

if found_cert_ids != CERT_IDS:
    raise RuntimeError(
        "Expected all five Process Technology certificate IDs. "
        f"Found: {sorted(found_cert_ids)}"
    )

current_cert_lineages = set(
    df.loc[
        cert_mask,
        "canonical_lineage"
    ].str.strip().unique()
)

allowed = {
    OLD,
    TARGET,
}

unexpected = current_cert_lineages - allowed

if unexpected:
    raise RuntimeError(
        "Unexpected current Process Technology certificate lineage(s): "
        + ", ".join(sorted(unexpected))
    )

aas_before = (
    df.loc[
        aas_mask,
        ["credential_id", "canonical_lineage"]
    ]
    .drop_duplicates()
    .sort_values("credential_id")
)

if not aas_before.empty:
    bad_aas = aas_before[
        ~aas_before["canonical_lineage"]
        .eq("PROCESS_TECHNOLOGY_AAS")
    ]

    if not bad_aas.empty:
        raise RuntimeError(
            "Process Operating Technology AAS is not in expected "
            "PROCESS_TECHNOLOGY_AAS lineage:\n"
            + bad_aas.to_string(index=False)
        )

complete_backup = backup(
    COMPLETE,
    "PROCESS_TECH_CERT_LINEAGE_FIX"
)

changed_rows = int(
    (
        cert_mask
        & df["canonical_lineage"].eq(OLD)
    ).sum()
)

df.loc[
    cert_mask,
    "canonical_lineage"
] = TARGET

# Mechanical lineage intentionally remains PROCESS_TECHNOLOGY.

df.to_csv(
    COMPLETE,
    index=False,
)

check = (
    df.loc[
        cert_mask | aas_mask,
        [
            "credential_id",
            "catalog_year",
            "mechanical_lineage",
            "canonical_lineage",
        ]
    ]
    .drop_duplicates()
    .sort_values(
        [
            "canonical_lineage",
            "credential_id",
        ]
    )
)

if set(
    df.loc[
        cert_mask,
        "canonical_lineage"
    ].unique()
) != {TARGET}:
    raise RuntimeError(
        "Certificate lineage validation failed."
    )

if not df.loc[
    aas_mask,
    "canonical_lineage"
].eq(
    "PROCESS_TECHNOLOGY_AAS"
).all():
    raise RuntimeError(
        "AAS lineage changed unexpectedly."
    )


# =============================================================================
# 2. OFFICIAL AWARD CROSSWALK
# =============================================================================

if not CROSSWALK.exists():
    raise FileNotFoundError(CROSSWALK)

crosswalk_backup = backup(
    CROSSWALK,
    "PROCESS_TECH_CERT_LINEAGE_FIX"
)

wb = load_workbook(CROSSWALK)

if "Crosswalk Draft" not in wb.sheetnames:
    raise RuntimeError(
        "Crosswalk Draft sheet not found."
    )

ws = wb["Crosswalk Draft"]

header_row = None
headers = {}

for r in range(1, min(ws.max_row, 15) + 1):
    row_headers = {
        str(ws.cell(r, c).value or "").strip(): c
        for c in range(1, ws.max_column + 1)
    }

    if "Curr1ProgramCode" in row_headers:
        header_row = r
        headers = row_headers
        break

if header_row is None:
    raise RuntimeError(
        "Could not locate crosswalk header row."
    )

needed = {
    "Curr1ProgramCode",
    "Major1Code",
    "DegreeCode",
    "canonical_lineage",
    "match_status",
    "curation_note",
    "modeled_lineage_present",
    "reconciliation_key",
}

missing = needed - set(headers)

if missing:
    raise RuntimeError(
        "Crosswalk missing fields: "
        + ", ".join(sorted(missing))
    )

target_codes = {
    "CERT1_PTAA": "PTAA",
    "CERT1_PTAC": "PTAC",
}

matched = []

for r in range(header_row + 1, ws.max_row + 1):

    code = str(
        ws.cell(
            r,
            headers["Curr1ProgramCode"]
        ).value or ""
    ).strip()

    if code not in target_codes:
        continue

    major = str(
        ws.cell(
            r,
            headers["Major1Code"]
        ).value or ""
    ).strip()

    degree = str(
        ws.cell(
            r,
            headers["DegreeCode"]
        ).value or ""
    ).strip()

    if major != target_codes[code]:
        raise RuntimeError(
            f"{code}: expected major "
            f"{target_codes[code]}, found {major}"
        )

    if degree != "CERT1":
        raise RuntimeError(
            f"{code}: expected CERT1, found {degree}"
        )

    ws.cell(
        r,
        headers["canonical_lineage"]
    ).value = TARGET

    ws.cell(
        r,
        headers["match_status"]
    ).value = "CURATED_MATCH"

    ws.cell(
        r,
        headers["curation_note"]
    ).value = (
        "Catalog-validated Process Technology Basic Certificate "
        "alias. PTAA/PTAC represent the modeled 30-SCH "
        "Process Technology certificate lineage."
    )

    ws.cell(
        r,
        headers["modeled_lineage_present"]
    ).value = True

    ws.cell(
        r,
        headers["reconciliation_key"]
    ).value = TARGET

    matched.append(code)

if set(matched) != set(target_codes):
    raise RuntimeError(
        "Did not locate exactly the expected PTAA/PTAC "
        f"crosswalk rows. Found: {matched}"
    )

wb.save(CROSSWALK)


# =============================================================================
# 3. REMOVE THE OLD SOURCE-LEVEL AAS MISCLASSIFICATION
# =============================================================================

old_line = (
    '"PROCESS_TECHNOLOGY": '
    '"PROCESS_TECHNOLOGY_AAS",'
)

new_line = (
    '"PROCESS_TECHNOLOGY": '
    '"PROCESS_TECHNOLOGY_CERTIFICATE",'
)

patched_builders = []

for path in BUILDERS:

    if not path.exists():
        continue

    text = path.read_text(
        encoding="utf-8"
    )

    if new_line in text:
        continue

    if old_line not in text:
        print(
            f"Builder unchanged — exact old mapping "
            f"not present: {path.name}"
        )
        continue

    backup(
        path,
        "PROCESS_TECH_CERT_LINEAGE_FIX"
    )

    patched = text.replace(
        old_line,
        new_line,
        1,
    )

    path.write_text(
        patched,
        encoding="utf-8",
    )

    patched_builders.append(
        path.name
    )


# =============================================================================
# 4. FINAL QA
# =============================================================================

print()
print("Completion rows changed:", f"{changed_rows:,}")
print("Completion backup:", complete_backup)
print("Crosswalk backup:", crosswalk_backup)

print()
print("MODELED PROCESS TECHNOLOGY LINEAGES")
print(check.to_string(index=False))

print()
print("PTAA/PTAC crosswalk rows updated:")
print("  " + ", ".join(sorted(matched)))

print()
print("Reporting builders patched:")
if patched_builders:
    for name in patched_builders:
        print(" ", name)
else:
    print("  none required")

print()
print("PASS")
print(
    "PROCESS_TECHNOLOGY_*             -> "
    "PROCESS_TECHNOLOGY_CERTIFICATE"
)
print(
    "PROCESS_OPERATING_TECHNOLOGY_*   -> "
    "PROCESS_TECHNOLOGY_AAS"
)
