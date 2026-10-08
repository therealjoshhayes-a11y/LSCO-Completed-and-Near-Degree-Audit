from pathlib import Path
from datetime import datetime
import shutil

import pandas as pd
from openpyxl import load_workbook

ROOT = Path.cwd()

MODELED = (
    ROOT
    / "data/processed/full_actual_audit/president_report/"
      "complete_rows_with_canonical_lineage.csv"
)

CROSSWALK = (
    ROOT
    / "data/interim/institutional_awards/"
      "award_program_crosswalk_curated.xlsx"
)

OLD_CREDENTIAL = (
    "INFORMATION_TECHNOLOGY_SUPPORT_ASSISTANT_"
    "CYBERSECURITY_SPECIALIST_2021"
)

OLD_LINEAGE = (
    "INFORMATION_TECHNOLOGY_SUPPORT_ASSISTANT_"
    "CYBERSECURITY_SPECIALIST"
)

NEW_LINEAGE = "CYBERSECURITY_SPECIALIST"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")


def backup(path: Path) -> Path:
    dst = path.with_name(
        f"{path.stem}_PRE_CYBERSECURITY_LINEAGE_FIX_{STAMP}{path.suffix}"
    )
    shutil.copy2(path, dst)
    return dst


print("=" * 100)
print("CYBERSECURITY 2021->2022 LINEAGE REPAIR")
print("=" * 100)

# ------------------------------------------------------------------
# 1. MODELED CREDENTIAL-YEAR -> CANONICAL LINEAGE
# ------------------------------------------------------------------

modeled_backup = backup(MODELED)

df = pd.read_csv(MODELED, dtype=str, low_memory=False).fillna("")

mask = df["credential_id"].eq(OLD_CREDENTIAL)

if not mask.any():
    raise RuntimeError(
        f"No rows found for modeled credential: {OLD_CREDENTIAL}"
    )

before = sorted(df.loc[mask, "canonical_lineage"].unique())

bad_values = [
    x for x in before
    if x not in {OLD_LINEAGE, NEW_LINEAGE}
]

if bad_values:
    raise RuntimeError(
        f"Unexpected existing canonical lineage(s): {bad_values}"
    )

rows_updated = int(
    (mask & ~df["canonical_lineage"].eq(NEW_LINEAGE)).sum()
)

df.loc[mask, "canonical_lineage"] = NEW_LINEAGE

df.to_csv(MODELED, index=False)

# ------------------------------------------------------------------
# 2. OFFICIAL-AWARD CROSSWALK
# ------------------------------------------------------------------

crosswalk_backup = backup(CROSSWALK)

wb = load_workbook(CROSSWALK)

matching_rows = []

for ws in wb.worksheets:
    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }

    required = {
        "Curr1ProgramCode",
        "Major1Code",
        "DegreeCode",
        "canonical_lineage",
    }

    if not required.issubset(headers):
        continue

    for r in range(2, ws.max_row + 1):
        program = str(
            ws.cell(r, headers["Curr1ProgramCode"]).value or ""
        ).strip().upper()

        major = str(
            ws.cell(r, headers["Major1Code"]).value or ""
        ).strip().upper()

        degree = str(
            ws.cell(r, headers["DegreeCode"]).value or ""
        ).strip().upper()

        if (
            program == "ZCONV CERT"
            and major == "ITSA"
            and degree == "CERT1"
        ):
            matching_rows.append((ws, r, headers))

if len(matching_rows) != 1:
    raise RuntimeError(
        "Expected exactly one ZCONV / ITSA / CERT1 crosswalk row; "
        f"found {len(matching_rows)}."
    )

ws, r, headers = matching_rows[0]

old_award_lineage = str(
    ws.cell(r, headers["canonical_lineage"]).value or ""
).strip()

if old_award_lineage not in {OLD_LINEAGE, NEW_LINEAGE}:
    raise RuntimeError(
        "Unexpected crosswalk lineage for ZCONV / ITSA / CERT1: "
        f"{old_award_lineage}"
    )

ws.cell(
    r,
    headers["canonical_lineage"]
).value = NEW_LINEAGE

if "reconciliation_key" in headers:
    ws.cell(
        r,
        headers["reconciliation_key"]
    ).value = NEW_LINEAGE

wb.save(CROSSWALK)

# ------------------------------------------------------------------
# 3. VERIFY
# ------------------------------------------------------------------

check = pd.read_csv(
    MODELED,
    dtype=str,
    low_memory=False
).fillna("")

remaining_old = check.loc[
    check["credential_id"].eq(OLD_CREDENTIAL)
    & ~check["canonical_lineage"].eq(NEW_LINEAGE)
]

if not remaining_old.empty:
    raise RuntimeError(
        "Modeled-lineage verification failed."
    )

wb_check = load_workbook(
    CROSSWALK,
    read_only=True,
    data_only=True
)

verified_crosswalk = False

for ws2 in wb_check.worksheets:
    headers2 = {
        str(cell.value).strip(): cell.column
        for cell in ws2[1]
        if cell.value is not None
    }

    required2 = {
        "Curr1ProgramCode",
        "Major1Code",
        "DegreeCode",
        "canonical_lineage",
    }

    if not required2.issubset(headers2):
        continue

    for rr in range(2, ws2.max_row + 1):
        vals = (
            str(ws2.cell(
                rr, headers2["Curr1ProgramCode"]
            ).value or "").strip().upper(),
            str(ws2.cell(
                rr, headers2["Major1Code"]
            ).value or "").strip().upper(),
            str(ws2.cell(
                rr, headers2["DegreeCode"]
            ).value or "").strip().upper(),
        )

        if vals == ("ZCONV CERT", "ITSA", "CERT1"):
            lineage = str(
                ws2.cell(
                    rr,
                    headers2["canonical_lineage"]
                ).value or ""
            ).strip()

            verified_crosswalk = (
                lineage == NEW_LINEAGE
            )

if not verified_crosswalk:
    raise RuntimeError(
        "Official-award crosswalk verification failed."
    )

print()
print("PASS")
print(f"Modeled rows updated: {rows_updated:,}")
print(f"Canonical lineage: {NEW_LINEAGE}")
print()
print("Backups:")
print(modeled_backup.relative_to(ROOT))
print(crosswalk_backup.relative_to(ROOT))

