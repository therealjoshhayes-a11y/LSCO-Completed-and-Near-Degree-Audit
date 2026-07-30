from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ======================================================================================
# CONFIGURATION
# ======================================================================================

ROOT = Path(__file__).resolve().parents[1]

COURSES_PATH = ROOT / "data" / "processed" / "normalized_actual_student_course_history.csv"
MULTICATALOG_PATH = (
    ROOT / "data" / "processed" / "catalogs" / "requirements_master_multicatalog.csv"
)
MULTIYEAR_PATH = (
    ROOT / "data" / "processed" / "catalogs" / "requirements_master_multiyear.csv"
)
COMPOUND_PATH = (
    ROOT
    / "data"
    / "processed"
    / "catalogs"
    / "semantic_policies"
    / "compound_requirement_alternatives.csv"
)
CHUNK_DIR = ROOT / "data" / "processed" / "full_actual_audit" / "chunks"
OUTPUT_DIR = (
    ROOT
    / "data"
    / "processed"
    / "reporting"
    / "industrial_one_course_2025_2026"
)
WORKBOOK_PATH = OUTPUT_DIR / "industrial_one_course_report.xlsx"

# Fall 2025 through Summer 2026. Fall 2026 (202690) is intentionally excluded.
ACTIVE_TERM_MIN = 202590
ACTIVE_TERM_MAX = 202689

# Controlled credential-family scope. These are substring matches against credential_family.
# Basic, intermediate, management, certificate, and AAS variants remain distinct credentials.
INCLUDE_SCOPE_TOKENS = (
    "AUTOMOTIVE",
    "BUILDING_CONSTRUCTION",
    "BUSINESS_CONSTRUCTION_MANAGEMENT",
    "CONSTRUCTION_MANAGEMENT",
    "ELECTRICIAN",
    "ELECTRO_MECHANICAL",
    "ELECTROMECHANICAL",
    "HEATING_VENTILATION",
    "HVAC",
    "INDUSTRIAL_TECHNOLOGY",
    "INSTRUMENTATION",
    "LAYOUT_AND_FABRICATION",
    "MACHINING",
    "PIPEFITTING",
    "PIPE_WELDING",
    "PLUMBING",
    "PROCESS_OPERATING_TECHNOLOGY",
    "PROCESS_TECHNOLOGY",
    "PRODUCTION_WELDER",
    "SAFETY_HEALTH_AND_ENVIRONMENT",
    "SAFETY_HEALTH_AND_ENVIRONMENTAL",
    "WELDING",
)

# Adjacent transportation programs are excluded from this industrial/manufacturing slice.
EXCLUDE_SCOPE_TOKENS = (
    "LOGISTICS",
    "MARITIME",
    "ORDINARY_SEAMAN",
)

GREEN = "00573F"
LIGHT_GREEN = "E7F0EC"
LIGHT_GOLD = "F3EEDC"
LIGHT_GRAY = "E7E7E7"
MID_GRAY = "777777"
DARK_GRAY = "333333"
WHITE = "FFFFFF"


# ======================================================================================
# GENERIC HELPERS
# ======================================================================================


def normalize_text(value: object) -> str:
    return " ".join(str(value).strip().split())


def normalize_course(value: object) -> str:
    return normalize_text(value).upper()


def split_semicolon(value: object) -> list[str]:
    return [
        normalize_course(part)
        for part in str(value).split(";")
        if normalize_course(part)
    ]


def boolish(series: pd.Series) -> pd.Series:
    return series.astype(str).str.strip().str.upper().isin({"TRUE", "T", "1", "YES", "Y"})


def catalog_start(value: object) -> int:
    match = re.match(r"^(\d{4})", normalize_text(value))
    return int(match.group(1)) if match else -1


def course_credit_hours(course_code: str) -> int | None:
    parts = normalize_course(course_code).split()
    if len(parts) != 2:
        return None
    number = parts[1]
    if len(number) != 4 or not number.isdigit():
        return None
    return int(number[1])


def scope_match(credential_family: object) -> bool:
    family = normalize_text(credential_family).upper()
    if any(token in family for token in EXCLUDE_SCOPE_TOKENS):
        return False
    return any(token in family for token in INCLUDE_SCOPE_TOKENS)


def strip_level_suffixes(credential_family: object) -> str:
    family = normalize_text(credential_family).upper()
    suffixes = (
        "_CERTIFICATE_OF_COMPLETION",
        "_CERTIFICATE",
        "_AAS",
    )
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if family.endswith(suffix):
                family = family[: -len(suffix)]
                changed = True
    return family


def classify_award_level(
    credential_family: object,
    credential_id: object,
    credential_title: object,
    program_hours: float,
) -> str:
    combined = " ".join(
        [
            normalize_text(credential_family).upper(),
            normalize_text(credential_id).upper(),
            normalize_text(credential_title).upper(),
        ]
    )

    if re.search(r"(^|[_\s])AAS($|[_\s])", combined):
        return "AAS"
    if "ASSOCIATE OF APPLIED SCIENCE" in combined:
        return "AAS"

    # LSCO certificates in this slice are ordinarily 42 SCH or fewer; AAS plans are ~60 SCH.
    return "AAS" if program_hours >= 45 else "Certificate"


def format_course_choice(courses: Iterable[str], title_lookup: dict[str, str]) -> str:
    values = sorted({normalize_course(course) for course in courses if normalize_course(course)})
    if len(values) == 1:
        code = values[0]
        title = title_lookup.get(code, "")
        return f"{code} — {title}" if title else code
    return "One of: " + " / ".join(values)


def safe_filename_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    present = [column for column in columns if column in df.columns]
    return df[present].copy()


# ======================================================================================
# CONTROL TABLES
# ======================================================================================


def build_course_title_lookup(multicatalog: pd.DataFrame) -> dict[str, str]:
    candidates: dict[str, list[tuple[int, str]]] = {}

    for row in multicatalog.itertuples(index=False):
        row_dict = row._asdict()
        raw_text = normalize_text(row_dict.get("raw_requirement_text", ""))
        catalog_year = catalog_start(row_dict.get("catalog_year", ""))
        course_codes = split_semicolon(row_dict.get("course_codes", ""))

        if len(course_codes) != 1 or not raw_text:
            continue

        code = course_codes[0]
        match = re.search(rf"\b{re.escape(code)}\b\s*[-–—:]?\s*(.*)$", raw_text, flags=re.I)
        if not match:
            continue

        title = normalize_text(match.group(1))
        title = re.sub(r"\s+\d+(?:\.\d+)?\s*$", "", title).strip(" -–—:;,")
        if not title or title.upper() == code:
            continue

        candidates.setdefault(code, []).append((catalog_year, title))

    lookup: dict[str, str] = {}
    for code, values in candidates.items():
        values.sort(key=lambda item: (item[0], len(item[1])), reverse=True)
        lookup[code] = values[0][1]

    return lookup


def load_credential_metadata() -> tuple[pd.DataFrame, dict[str, dict], dict[str, str]]:
    if not MULTICATALOG_PATH.exists():
        raise FileNotFoundError(MULTICATALOG_PATH)

    multi = pd.read_csv(MULTICATALOG_PATH, dtype=str, low_memory=False).fillna("")
    required = {
        "requirement_id",
        "catalog_year",
        "credential_family",
        "credential_id",
        "credential_title",
        "credit_hours",
        "raw_requirement_text",
        "course_codes",
    }
    missing = required - set(multi.columns)
    if missing:
        raise KeyError("Multicatalog requirements missing columns: " + ", ".join(sorted(missing)))

    multi["credit_hours_numeric"] = pd.to_numeric(multi["credit_hours"], errors="coerce").fillna(0.0)
    multi["catalog_year_start_numeric"] = multi["catalog_year"].map(catalog_start)

    credential_meta = (
        multi.groupby(
            ["catalog_year", "credential_family", "credential_id", "credential_title"],
            as_index=False,
        )["credit_hours_numeric"]
        .sum()
        .rename(columns={"credit_hours_numeric": "program_hours"})
    )
    credential_meta["in_scope"] = credential_meta["credential_family"].map(scope_match)
    credential_meta = credential_meta[credential_meta["in_scope"]].copy()
    credential_meta["catalog_year_start_numeric"] = credential_meta["catalog_year"].map(catalog_start)
    credential_meta["award_level"] = credential_meta.apply(
        lambda row: classify_award_level(
            row["credential_family"],
            row["credential_id"],
            row["credential_title"],
            float(row["program_hours"]),
        ),
        axis=1,
    )
    credential_meta["credential_reporting_key"] = credential_meta.apply(
        lambda row: f"{strip_level_suffixes(row['credential_family'])}|{row['award_level']}",
        axis=1,
    )

    sort_cols = ["catalog_year_start_numeric"]
    if "requirement_sequence" in multi.columns:
        multi["requirement_sequence_numeric"] = pd.to_numeric(
            multi["requirement_sequence"], errors="coerce"
        ).fillna(999999)
        sort_cols.append("requirement_sequence_numeric")

    requirement_meta = (
        multi.sort_values(sort_cols, kind="mergesort")
        .drop_duplicates("requirement_id", keep="last")
        .set_index("requirement_id")
        .to_dict("index")
    )

    title_lookup = build_course_title_lookup(multi)
    return credential_meta, requirement_meta, title_lookup


def load_multiyear_options() -> dict[str, pd.DataFrame]:
    if not MULTIYEAR_PATH.exists():
        raise FileNotFoundError(MULTIYEAR_PATH)

    options = pd.read_csv(MULTIYEAR_PATH, dtype=str, low_memory=False).fillna("")
    required = {"requirement_id", "rule_type", "option_type", "option_value", "credit_hours"}
    missing = required - set(options.columns)
    if missing:
        raise KeyError("Multiyear requirements missing columns: " + ", ".join(sorted(missing)))

    if "min_required" not in options.columns:
        options["min_required"] = "1"

    return {
        str(requirement_id): group.copy()
        for requirement_id, group in options.groupby("requirement_id", sort=False)
    }


def load_compound_alternatives() -> dict[str, list[list[str]]]:
    if not COMPOUND_PATH.exists():
        return {}

    rows = pd.read_csv(COMPOUND_PATH, dtype=str, low_memory=False).fillna("")
    required = {"requirement_id", "alternative_group", "option_value"}
    missing = required - set(rows.columns)
    if missing:
        print(
            "WARNING: compound alternatives file exists but does not contain expected columns; "
            "compound rules will use the standard option resolver."
        )
        return {}

    order_column = (
        "alternative_course_order"
        if "alternative_course_order" in rows.columns
        else None
    )

    rows["alternative_group_numeric"] = pd.to_numeric(
        rows["alternative_group"], errors="coerce"
    ).fillna(0).astype(int)
    if order_column:
        rows["alternative_course_order_numeric"] = pd.to_numeric(
            rows[order_column], errors="coerce"
        ).fillna(0).astype(int)

    result: dict[str, list[list[str]]] = {}

    for requirement_id, requirement_rows in rows.groupby("requirement_id", sort=False):
        alternatives: list[list[str]] = []
        for _, alternative_rows in requirement_rows.groupby("alternative_group_numeric", sort=True):
            if order_column:
                alternative_rows = alternative_rows.sort_values("alternative_course_order_numeric")
            courses = [
                normalize_course(value)
                for value in alternative_rows["option_value"]
                if normalize_course(value)
            ]
            if courses:
                alternatives.append(courses)
        if alternatives:
            result[str(requirement_id)] = alternatives

    return result


# ======================================================================================
# 2025-2026 ACTIVE COHORT
# ======================================================================================


def load_active_students_and_passed_courses() -> tuple[set[str], dict[str, set[str]]]:
    if not COURSES_PATH.exists():
        raise FileNotFoundError(COURSES_PATH)

    usecols = [
        "student_id",
        "course_code",
        "term_sort",
        "passed",
        "grade",
        "final_grade",
    ]
    header = pd.read_csv(COURSES_PATH, nrows=0).columns.tolist()
    usecols = [column for column in usecols if column in header]
    courses = pd.read_csv(
        COURSES_PATH,
        dtype=str,
        low_memory=False,
        usecols=usecols,
    ).fillna("")

    required = {"student_id", "course_code", "term_sort"}
    missing = required - set(courses.columns)
    if missing:
        raise KeyError("Normalized course history missing columns: " + ", ".join(sorted(missing)))

    courses["student_id"] = courses["student_id"].astype(str).str.strip()
    courses["term_sort_numeric"] = pd.to_numeric(courses["term_sort"], errors="coerce")

    active_mask = courses["term_sort_numeric"].between(
        ACTIVE_TERM_MIN,
        ACTIVE_TERM_MAX,
        inclusive="both",
    )
    active_students = set(courses.loc[active_mask, "student_id"])
    active_students.discard("")

    if "passed" in courses.columns:
        passed_mask = boolish(courses["passed"])
    else:
        grade_column = "grade" if "grade" in courses.columns else "final_grade"
        passing_grades = {"A", "B", "C", "D", "S", "E", "T", "P", "CR"}
        passed_mask = courses[grade_column].astype(str).str.strip().str.upper().isin(passing_grades)

    passed = courses[passed_mask & courses["student_id"].isin(active_students)].copy()
    passed["course_code"] = passed["course_code"].map(normalize_course)

    passed_lookup = {
        str(student_id): set(group["course_code"])
        for student_id, group in passed.groupby("student_id", sort=False)
    }

    return active_students, passed_lookup


# ======================================================================================
# SUMMARY + DETAIL EXTRACTION
# ======================================================================================


def collect_summary_candidates(
    active_students: set[str],
    credential_meta: pd.DataFrame,
) -> pd.DataFrame:
    summary_files = sorted(CHUNK_DIR.glob("chunk_*_credential_summary.csv"))
    if not summary_files:
        raise FileNotFoundError(f"No summary chunks found in {CHUNK_DIR}")

    scope_ids = set(credential_meta["credential_id"])
    usecols = [
        "student_id",
        "catalog_year",
        "credential_id",
        "requirements_missing",
        "requirements_unresolved",
        "audit_status",
        "catalog_eligible",
        "eligibility_reason",
    ]
    frames: list[pd.DataFrame] = []

    for index, summary_path in enumerate(summary_files, start=1):
        header = pd.read_csv(summary_path, nrows=0).columns.tolist()
        present = [column for column in usecols if column in header]
        summary = pd.read_csv(
            summary_path,
            dtype=str,
            low_memory=False,
            usecols=present,
        ).fillna("")

        required = {
            "student_id",
            "catalog_year",
            "credential_id",
            "requirements_missing",
            "requirements_unresolved",
            "catalog_eligible",
        }
        missing = required - set(summary.columns)
        if missing:
            raise KeyError(f"{summary_path.name} missing columns: " + ", ".join(sorted(missing)))

        summary["requirements_missing_numeric"] = pd.to_numeric(
            summary["requirements_missing"], errors="coerce"
        )
        summary["requirements_unresolved_numeric"] = pd.to_numeric(
            summary["requirements_unresolved"], errors="coerce"
        )

        keep = summary[
            summary["student_id"].isin(active_students)
            & summary["credential_id"].isin(scope_ids)
            & summary["requirements_missing_numeric"].eq(1)
            & summary["requirements_unresolved_numeric"].eq(0)
            & boolish(summary["catalog_eligible"])
        ].copy()

        if not keep.empty:
            keep["summary_chunk"] = summary_path.name
            frames.append(keep)

        if index % 100 == 0 or index == len(summary_files):
            print(f"Read summary chunks: {index:,}/{len(summary_files):,}")

    if not frames:
        return pd.DataFrame()

    return pd.concat(frames, ignore_index=True)


def collect_missing_detail_rows(summary_candidates: pd.DataFrame) -> pd.DataFrame:
    if summary_candidates.empty:
        return pd.DataFrame()

    usecols = [
        "student_id",
        "catalog_year",
        "credential_id",
        "requirement_id",
        "rule_type",
        "status",
        "matched_options",
        "required_options",
        "option_types",
        "used_course_count",
    ]
    key_columns = ["student_id", "catalog_year", "credential_id"]
    frames: list[pd.DataFrame] = []

    grouped_chunks = list(summary_candidates.groupby("summary_chunk", sort=True))

    for index, (summary_chunk, chunk_candidates) in enumerate(grouped_chunks, start=1):
        detail_name = str(summary_chunk).replace(
            "_credential_summary.csv",
            "_audit_results.csv",
        )
        detail_path = CHUNK_DIR / detail_name
        if not detail_path.exists():
            raise FileNotFoundError(detail_path)

        header = pd.read_csv(detail_path, nrows=0).columns.tolist()
        present = [column for column in usecols if column in header]
        detail = pd.read_csv(
            detail_path,
            dtype=str,
            low_memory=False,
            usecols=present,
        ).fillna("")

        required = {
            "student_id",
            "catalog_year",
            "credential_id",
            "requirement_id",
            "rule_type",
            "status",
            "matched_options",
            "required_options",
            "option_types",
        }
        missing = required - set(detail.columns)
        if missing:
            raise KeyError(f"{detail_path.name} missing columns: " + ", ".join(sorted(missing)))

        candidate_keys = chunk_candidates[key_columns].drop_duplicates()
        detail = detail.merge(candidate_keys, on=key_columns, how="inner")

        status_upper = detail["status"].astype(str).str.upper()
        detail = detail[
            ~status_upper.eq("MET")
            & ~status_upper.str.startswith("UNRESOLVED")
        ].copy()

        if not detail.empty:
            detail["detail_chunk"] = detail_path.name
            frames.append(detail)

        if index % 25 == 0 or index == len(grouped_chunks):
            print(f"Read candidate detail chunks: {index:,}/{len(grouped_chunks):,}")

    if not frames:
        return pd.DataFrame()

    return pd.concat(frames, ignore_index=True)


# ======================================================================================
# ONE-COURSE RESOLUTION
# ======================================================================================


def resolve_one_course_need(
    detail_row: pd.Series,
    option_groups: dict[str, pd.DataFrame],
    requirement_meta: dict[str, dict],
    compound_alternatives: dict[str, list[list[str]]],
    passed_courses: set[str],
    title_lookup: dict[str, str],
) -> tuple[str, str, str, str]:
    """
    Returns:
        resolution_status, needed_label, needed_kind, resolution_note
    """

    requirement_id = str(detail_row["requirement_id"])
    group = option_groups.get(requirement_id)
    if group is None or group.empty:
        return "EXCLUDE_REQUIREMENT_NOT_FOUND", "", "", "No multiyear option rows found."

    matched = set(split_semicolon(detail_row.get("matched_options", "")))
    option_types = {
        normalize_text(value).upper()
        for value in group["option_type"]
        if normalize_text(value)
    }
    rule_type = normalize_text(group.iloc[0]["rule_type"]).upper()

    min_required_value = pd.to_numeric(
        pd.Series([group.iloc[0].get("min_required", "1")]),
        errors="coerce",
    ).iloc[0]
    min_required = 1 if pd.isna(min_required_value) else int(min_required_value)

    hours_value = pd.to_numeric(
        pd.Series([group.iloc[0].get("credit_hours", "")]),
        errors="coerce",
    ).iloc[0]
    required_credit_hours = 0.0 if pd.isna(hours_value) else float(hours_value)

    # Compound alternatives preserve structures such as (A+B) OR (C+D).
    alternatives = compound_alternatives.get(requirement_id, [])
    if alternatives:
        one_course_alternatives: list[str] = []
        passed_allocation_conflict = False

        for alternative in alternatives:
            normalized_alternative = [normalize_course(course) for course in alternative]
            remaining = [course for course in normalized_alternative if course not in matched]
            if len(remaining) == 1:
                if remaining[0] in passed_courses:
                    passed_allocation_conflict = True
                else:
                    one_course_alternatives.append(remaining[0])

        one_course_alternatives = sorted(set(one_course_alternatives))
        if one_course_alternatives:
            label = format_course_choice(one_course_alternatives, title_lookup)
            kind = "EXACT_COURSE" if len(one_course_alternatives) == 1 else "COMPOUND_COURSE_CHOICE"
            return "KEEP", label, kind, "Compound alternative is exactly one unpassed course away."

        if passed_allocation_conflict:
            return (
                "EXCLUDE_ALREADY_PASSED_ALLOCATION_CONFLICT",
                "",
                "",
                "A one-course compound option is already passed but allocated elsewhere.",
            )

        return (
            "EXCLUDE_COMPOUND_REQUIRES_MULTIPLE_COURSES",
            "",
            "",
            "No compound alternative can be completed with one additional course.",
        )

    # Ordinary exact/ANY_N course options.
    if option_types and option_types <= {"COURSE"}:
        all_options = sorted(
            {
                normalize_course(value)
                for value in group["option_value"]
                if normalize_course(value)
            }
        )
        matched_in_group = matched.intersection(all_options)
        remaining_slots = min_required - len(matched_in_group)

        if remaining_slots != 1:
            return (
                "EXCLUDE_REQUIREMENT_NOT_ONE_COURSE",
                "",
                "",
                f"Requirement has {remaining_slots} remaining course slots.",
            )

        available_options = [course for course in all_options if course not in matched_in_group]
        if not available_options:
            return "EXCLUDE_NO_REMAINING_COURSE_OPTION", "", "", "No unused option remains."

        passed_but_unmatched = [course for course in available_options if course in passed_courses]
        if passed_but_unmatched:
            return (
                "EXCLUDE_ALREADY_PASSED_ALLOCATION_CONFLICT",
                "",
                "",
                "Already-passed option(s) were unavailable because of audit allocation: "
                + "; ".join(passed_but_unmatched),
            )

        matched_hours = sum(
            course_credit_hours(course) or 0
            for course in matched_in_group
        )
        remaining_required_hours = max(required_credit_hours - matched_hours, 0.0)

        plausible_options: list[str] = []
        for course in available_options:
            hours = course_credit_hours(course)
            # Test the remaining hours, not the full ANY_N requirement. For example,
            # a 6-SCH ANY_2 rule with one 3-SCH match is one ordinary 3-SCH course away.
            if remaining_required_hours <= 4 or (
                hours is not None and hours >= remaining_required_hours
            ):
                plausible_options.append(course)

        if not plausible_options:
            return (
                "EXCLUDE_LISTED_OPTIONS_CANNOT_CLOSE_HOURS_IN_ONE_COURSE",
                "",
                "",
                f"Remaining displayed hours are {remaining_required_hours:g} SCH.",
            )

        label = format_course_choice(plausible_options, title_lookup)
        kind = "EXACT_COURSE" if len(plausible_options) == 1 else "COURSE_CHOICE"
        return "KEEP", label, kind, "One listed unpassed course slot remains."

    meta = requirement_meta.get(requirement_id, {})
    raw_text = normalize_text(meta.get("raw_requirement_text", ""))
    if not raw_text:
        raw_text = normalize_text(detail_row.get("required_options", ""))

    if "CORE_BUCKET" in option_types or rule_type == "CORE_BUCKET":
        if 0 < required_credit_hours <= 4 and min_required <= 1:
            return (
                "KEEP",
                f"One approved core course — {raw_text}",
                "CORE_CHOICE",
                "One core-course slot remains; advisor selects the specific approved course.",
            )
        return (
            "EXCLUDE_CORE_REQUIRES_MULTIPLE_COURSES_OR_HOURS",
            "",
            "",
            f"Core requirement is {required_credit_hours:g} SCH with min_required={min_required}.",
        )

    if "ELECTIVE" in option_types or rule_type == "ELECTIVE":
        if 0 < required_credit_hours <= 4:
            return (
                "KEEP",
                f"One approved elective — {raw_text}",
                "ELECTIVE_CHOICE",
                "One elective course can close the displayed hours.",
            )
        return (
            "EXCLUDE_ELECTIVE_REQUIRES_MULTIPLE_COURSES_OR_HOURS",
            "",
            "",
            f"Elective requirement is {required_credit_hours:g} SCH.",
        )

    return (
        "EXCLUDE_UNSUPPORTED_REQUIREMENT_SHAPE",
        "",
        "",
        f"rule_type={rule_type}; option_types={'; '.join(sorted(option_types))}",
    )


def resolve_candidates(
    summary_candidates: pd.DataFrame,
    missing_detail: pd.DataFrame,
    credential_meta: pd.DataFrame,
    option_groups: dict[str, pd.DataFrame],
    requirement_meta: dict[str, dict],
    compound_alternatives: dict[str, list[list[str]]],
    passed_lookup: dict[str, set[str]],
    title_lookup: dict[str, str],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    key_columns = ["student_id", "catalog_year", "credential_id"]

    detail_counts = (
        missing_detail.groupby(key_columns)
        .size()
        .rename("missing_detail_row_count")
        .reset_index()
        if not missing_detail.empty
        else pd.DataFrame(columns=key_columns + ["missing_detail_row_count"])
    )

    candidate_base = summary_candidates.merge(detail_counts, on=key_columns, how="left")
    bad_detail_counts = candidate_base[
        ~candidate_base["missing_detail_row_count"].fillna(0).eq(1)
    ].copy()
    if not bad_detail_counts.empty:
        bad_detail_counts["exclusion_reason"] = "MISSING_DETAIL_ROW_COUNT_NOT_ONE"
        bad_detail_counts["resolution_note"] = (
            "A one-requirement-short summary must resolve to exactly one non-MET detail row."
        )

    valid_keys = candidate_base[
        candidate_base["missing_detail_row_count"].fillna(0).eq(1)
    ][key_columns].drop_duplicates()

    detail = missing_detail.merge(valid_keys, on=key_columns, how="inner")
    detail = detail.merge(
        credential_meta[
            [
                "catalog_year",
                "credential_id",
                "credential_family",
                "credential_title",
                "award_level",
                "program_hours",
                "catalog_year_start_numeric",
                "credential_reporting_key",
            ]
        ],
        on=["catalog_year", "credential_id"],
        how="left",
        validate="many_to_one",
    )

    resolved_rows: list[dict] = []
    excluded_rows: list[dict] = []

    for _, row in detail.iterrows():
        student_id = str(row["student_id"])
        status, needed_label, needed_kind, note = resolve_one_course_need(
            detail_row=row,
            option_groups=option_groups,
            requirement_meta=requirement_meta,
            compound_alternatives=compound_alternatives,
            passed_courses=passed_lookup.get(student_id, set()),
            title_lookup=title_lookup,
        )

        output = row.to_dict()
        output["resolution_status"] = status
        output["needed_label"] = needed_label
        output["needed_kind"] = needed_kind
        output["resolution_note"] = note

        if status == "KEEP":
            resolved_rows.append(output)
        else:
            output["exclusion_reason"] = status
            excluded_rows.append(output)

    resolved = pd.DataFrame(resolved_rows)
    excluded = pd.DataFrame(excluded_rows)

    if not bad_detail_counts.empty:
        excluded = pd.concat([excluded, bad_detail_counts], ignore_index=True, sort=False)

    return resolved, excluded


# ======================================================================================
# CATALOG DEDUPLICATION
# ======================================================================================


def select_canonical_catalog(resolved: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    if resolved.empty:
        return resolved, pd.DataFrame()

    resolved = resolved.copy()
    grain = ["student_id", "credential_reporting_key"]

    conflict = (
        resolved.groupby(grain)["needed_label"]
        .nunique()
        .rename("distinct_needed_labels_across_catalogs")
        .reset_index()
    )
    conflict = conflict[conflict["distinct_needed_labels_across_catalogs"] > 1].copy()

    kind_priority = {
        "EXACT_COURSE": 0,
        "COURSE_CHOICE": 1,
        "COMPOUND_COURSE_CHOICE": 2,
        "CORE_CHOICE": 3,
        "ELECTIVE_CHOICE": 4,
    }
    resolved["kind_priority"] = resolved["needed_kind"].map(kind_priority).fillna(99)

    # Prefer the latest eligible catalog; use the more specific course resolution as a tie-breaker.
    canonical = (
        resolved.sort_values(
            grain + ["catalog_year_start_numeric", "kind_priority", "needed_label"],
            ascending=[True, True, False, True, True],
            kind="mergesort",
        )
        .drop_duplicates(grain, keep="first")
        .copy()
    )

    canonical = canonical.merge(conflict, on=grain, how="left")
    canonical["catalog_need_conflict"] = canonical[
        "distinct_needed_labels_across_catalogs"
    ].fillna(0).gt(1)

    conflict_detail = resolved.merge(conflict[grain], on=grain, how="inner")
    conflict_detail["qa_reason"] = "DIFFERENT_ONE_COURSE_NEEDS_ACROSS_ELIGIBLE_CATALOGS"

    return canonical, conflict_detail


# ======================================================================================
# REPORT TABLES
# ======================================================================================


def build_reporting_tables(
    canonical: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if canonical.empty:
        empty_course = pd.DataFrame(
            columns=[
                "credential_reporting_key",
                "credential_title",
                "award_level",
                "needed_label",
                "needed_kind",
                "student_count",
            ]
        )
        empty_catalog = pd.DataFrame(
            columns=[
                "credential_reporting_key",
                "credential_title",
                "award_level",
                "catalog_year",
                "student_count",
            ]
        )
        return empty_course, empty_catalog, canonical

    course_counts = (
        canonical.groupby(
            [
                "credential_reporting_key",
                "credential_title",
                "award_level",
                "needed_label",
                "needed_kind",
            ],
            as_index=False,
        )["student_id"]
        .nunique()
        .rename(columns={"student_id": "student_count"})
    )

    credential_totals = (
        canonical.groupby(
            ["credential_reporting_key", "credential_title", "award_level"],
            as_index=False,
        )["student_id"]
        .nunique()
        .rename(columns={"student_id": "credential_student_total"})
    )
    course_counts = course_counts.merge(
        credential_totals,
        on=["credential_reporting_key", "credential_title", "award_level"],
        how="left",
    )

    catalog_counts = (
        canonical.groupby(
            [
                "credential_reporting_key",
                "credential_title",
                "award_level",
                "catalog_year",
            ],
            as_index=False,
        )["student_id"]
        .nunique()
        .rename(columns={"student_id": "student_count"})
    )
    catalog_counts = catalog_counts.merge(
        credential_totals,
        on=["credential_reporting_key", "credential_title", "award_level"],
        how="left",
    )

    course_counts = course_counts.sort_values(
        ["award_level", "credential_title", "student_count", "needed_label"],
        ascending=[True, True, False, True],
        kind="mergesort",
    )
    catalog_counts["catalog_year_start_numeric"] = catalog_counts["catalog_year"].map(catalog_start)
    catalog_counts = catalog_counts.sort_values(
        ["award_level", "credential_title", "catalog_year_start_numeric"],
        kind="mergesort",
    )

    student_columns = [
        "student_id",
        "credential_reporting_key",
        "credential_family",
        "credential_title",
        "award_level",
        "catalog_year",
        "program_hours",
        "needed_label",
        "needed_kind",
        "requirement_id",
        "rule_type",
        "required_options",
        "option_types",
        "resolution_note",
        "catalog_need_conflict",
        "summary_chunk",
        "detail_chunk",
    ]
    student_detail = safe_filename_columns(canonical, student_columns).sort_values(
        ["award_level", "credential_title", "needed_label", "student_id"],
        kind="mergesort",
    )

    return course_counts, catalog_counts, student_detail


# ======================================================================================
# EXCEL OUTPUT — TUFTE-STYLE OUTLINE
# ======================================================================================


def apply_sheet_defaults(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Lamar State College Orange | Credential Completion Audit"
    ws.oddFooter.right.text = "Page &P of &N"


def write_title_block(ws, title: str, subtitle: str) -> None:
    ws.merge_cells("A1:B1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=GREEN)
    ws["A1"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A2:B2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=10, color=DARK_GRAY)

    ws.merge_cells("A3:B3")
    ws["A3"] = (
        "Population: students enrolled in 2025–2026; any eligible catalog; "
        "industrial/manufacturing certificates and AAS credentials; one additional course required."
    )
    ws["A3"].font = Font(name="Aptos", size=9, italic=True, color=MID_GRAY)
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 30

    ws["A4"] = "Credential / needed course"
    ws["B4"] = "Students"
    for cell in ws[4]:
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="right" if cell.column == 2 else "left")


def write_course_outline(ws, course_counts: pd.DataFrame) -> None:
    write_title_block(
        ws,
        "Industrial & Manufacturing: One-Course Completion Opportunities",
        "Needed courses organized by credential",
    )

    row = 5
    thin_green = Side(style="thin", color=GREEN)

    if course_counts.empty:
        ws.cell(row, 1, "No qualifying one-course completion candidates were found.")
        ws.cell(row, 1).font = Font(italic=True, color=MID_GRAY)
        return

    group_cols = ["award_level", "credential_title", "credential_reporting_key"]
    for (award_level, credential_title, _), group in course_counts.groupby(group_cols, sort=False):
        total = int(group["credential_student_total"].iloc[0])
        header_row = row
        ws.cell(row, 1, f"{credential_title} · {award_level}")
        ws.cell(row, 2, total)
        ws.cell(row, 1).font = Font(name="Aptos", size=11, bold=True, color=GREEN)
        ws.cell(row, 2).font = Font(name="Aptos", size=11, bold=True, color=GREEN)
        ws.cell(row, 2).alignment = Alignment(horizontal="right")
        ws.cell(row, 1).border = Border(bottom=thin_green)
        ws.cell(row, 2).border = Border(bottom=thin_green)
        row += 1

        detail_start = row
        for item in group.itertuples(index=False):
            ws.cell(row, 1, item.needed_label)
            ws.cell(row, 2, int(item.student_count))
            ws.cell(row, 1).alignment = Alignment(indent=2, wrap_text=True)
            ws.cell(row, 2).alignment = Alignment(horizontal="right")
            ws.cell(row, 1).font = Font(name="Aptos", size=10, color=DARK_GRAY)
            ws.cell(row, 2).font = Font(name="Aptos", size=10, color=DARK_GRAY)
            ws.row_dimensions[row].outlineLevel = 1
            row += 1

        detail_end = row - 1
        if detail_end >= detail_start:
            ws.row_dimensions.group(detail_start, detail_end, outline_level=1, hidden=False)

        ws.row_dimensions[header_row].height = 20
        row += 1

    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 12
    ws.auto_filter.ref = f"A4:B{max(row - 1, 4)}"


def write_catalog_outline(ws, catalog_counts: pd.DataFrame) -> None:
    write_title_block(
        ws,
        "Industrial & Manufacturing: Catalog Distribution",
        "Selected eligible catalog for each student–credential opportunity",
    )

    row = 5
    thin_green = Side(style="thin", color=GREEN)

    if catalog_counts.empty:
        ws.cell(row, 1, "No qualifying one-course completion candidates were found.")
        ws.cell(row, 1).font = Font(italic=True, color=MID_GRAY)
        return

    group_cols = ["award_level", "credential_title", "credential_reporting_key"]
    for (award_level, credential_title, _), group in catalog_counts.groupby(group_cols, sort=False):
        total = int(group["credential_student_total"].iloc[0])
        header_row = row
        ws.cell(row, 1, f"{credential_title} · {award_level}")
        ws.cell(row, 2, total)
        ws.cell(row, 1).font = Font(name="Aptos", size=11, bold=True, color=GREEN)
        ws.cell(row, 2).font = Font(name="Aptos", size=11, bold=True, color=GREEN)
        ws.cell(row, 2).alignment = Alignment(horizontal="right")
        ws.cell(row, 1).border = Border(bottom=thin_green)
        ws.cell(row, 2).border = Border(bottom=thin_green)
        row += 1

        detail_start = row
        for item in group.itertuples(index=False):
            ws.cell(row, 1, str(item.catalog_year))
            ws.cell(row, 2, int(item.student_count))
            ws.cell(row, 1).alignment = Alignment(indent=2)
            ws.cell(row, 2).alignment = Alignment(horizontal="right")
            ws.cell(row, 1).font = Font(name="Aptos", size=10, color=DARK_GRAY)
            ws.cell(row, 2).font = Font(name="Aptos", size=10, color=DARK_GRAY)
            ws.row_dimensions[row].outlineLevel = 1
            row += 1

        detail_end = row - 1
        if detail_end >= detail_start:
            ws.row_dimensions.group(detail_start, detail_end, outline_level=1, hidden=False)

        ws.row_dimensions[header_row].height = 20
        row += 1

    ws.column_dimensions["A"].width = 65
    ws.column_dimensions["B"].width = 12
    ws.auto_filter.ref = f"A4:B{max(row - 1, 4)}"


def write_flat_table(ws, title: str, df: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=GREEN)

    if df.empty:
        ws["A3"] = "No rows."
        ws["A3"].font = Font(italic=True, color=MID_GRAY)
        return

    start_row = 3
    for column_index, column in enumerate(df.columns, start=1):
        cell = ws.cell(start_row, column_index, column)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(wrap_text=True)

    for row_offset, values in enumerate(df.itertuples(index=False, name=None), start=1):
        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(start_row + row_offset, column_index, value)
            cell.font = Font(name="Aptos", size=9, color=DARK_GRAY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_offset % 2 == 0:
            for column_index in range(1, len(df.columns) + 1):
                ws.cell(start_row + row_offset, column_index).fill = PatternFill(
                    "solid", fgColor="F7F7F7"
                )

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(df.columns))}{start_row + len(df)}"

    for column_index, column in enumerate(df.columns, start=1):
        values = [str(column)] + [normalize_text(value) for value in df[column].head(500)]
        max_len = min(max((len(value) for value in values), default=10) + 2, 45)
        ws.column_dimensions[get_column_letter(column_index)].width = max(10, max_len)


def write_scope_sheet(ws, credential_meta: pd.DataFrame) -> None:
    scope = credential_meta[
        [
            "catalog_year",
            "credential_family",
            "credential_id",
            "credential_title",
            "award_level",
            "program_hours",
        ]
    ].sort_values(["catalog_year", "award_level", "credential_title"], kind="mergesort")
    write_flat_table(ws, "Industrial & Manufacturing Credential Scope", scope)


def build_workbook(
    course_counts: pd.DataFrame,
    catalog_counts: pd.DataFrame,
    student_detail: pd.DataFrame,
    qa_exclusions: pd.DataFrame,
    catalog_conflicts: pd.DataFrame,
    credential_meta: pd.DataFrame,
) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    ws_courses = wb.create_sheet("Credential Course Needs")
    apply_sheet_defaults(ws_courses)
    write_course_outline(ws_courses, course_counts)

    ws_catalogs = wb.create_sheet("Catalog Counts")
    apply_sheet_defaults(ws_catalogs)
    write_catalog_outline(ws_catalogs, catalog_counts)

    ws_students = wb.create_sheet("Student Detail")
    write_flat_table(ws_students, "Student-Level Audit Detail", student_detail)

    qa_columns = [
        "student_id",
        "catalog_year",
        "credential_id",
        "credential_title",
        "award_level",
        "requirement_id",
        "exclusion_reason",
        "resolution_note",
        "required_options",
        "option_types",
        "summary_chunk",
        "detail_chunk",
    ]
    qa_display = safe_filename_columns(qa_exclusions, qa_columns)
    ws_qa = wb.create_sheet("QA Exclusions")
    write_flat_table(ws_qa, "Excluded Near-Completion Candidates", qa_display)

    conflict_columns = [
        "student_id",
        "credential_reporting_key",
        "credential_title",
        "award_level",
        "catalog_year",
        "needed_label",
        "needed_kind",
        "requirement_id",
        "qa_reason",
    ]
    conflict_display = safe_filename_columns(catalog_conflicts, conflict_columns)
    ws_conflicts = wb.create_sheet("Catalog Need Conflicts")
    write_flat_table(ws_conflicts, "Different One-Course Needs Across Eligible Catalogs", conflict_display)

    ws_scope = wb.create_sheet("Scope Credentials")
    write_scope_sheet(ws_scope, credential_meta)

    # Put the two executive-facing outlined sheets first.
    wb._sheets = [
        ws_courses,
        ws_catalogs,
        ws_students,
        ws_qa,
        ws_conflicts,
        ws_scope,
    ]

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK_PATH)


# ======================================================================================
# MAIN
# ======================================================================================


def main() -> None:
    print("=" * 100)
    print("INDUSTRIAL & MANUFACTURING — ONE-COURSE COMPLETION REPORT")
    print("=" * 100)

    credential_meta, requirement_meta, title_lookup = load_credential_metadata()
    option_groups = load_multiyear_options()
    compound_alternatives = load_compound_alternatives()
    active_students, passed_lookup = load_active_students_and_passed_courses()

    print(f"Active 2025-2026 students:          {len(active_students):,}")
    print(f"In-scope credential-years:         {len(credential_meta):,}")
    print(f"Multiyear requirement groups:      {len(option_groups):,}")
    print(f"Compound requirement groups:       {len(compound_alternatives):,}")

    summary_candidates = collect_summary_candidates(active_students, credential_meta)
    print(f"One-requirement-short audit rows:   {len(summary_candidates):,}")

    if summary_candidates.empty:
        resolved = pd.DataFrame()
        exclusions = pd.DataFrame()
        canonical = pd.DataFrame()
        conflicts = pd.DataFrame()
    else:
        missing_detail = collect_missing_detail_rows(summary_candidates)
        print(f"Non-MET candidate detail rows:      {len(missing_detail):,}")

        resolved, exclusions = resolve_candidates(
            summary_candidates=summary_candidates,
            missing_detail=missing_detail,
            credential_meta=credential_meta,
            option_groups=option_groups,
            requirement_meta=requirement_meta,
            compound_alternatives=compound_alternatives,
            passed_lookup=passed_lookup,
            title_lookup=title_lookup,
        )
        canonical, conflicts = select_canonical_catalog(resolved)

    course_counts, catalog_counts, student_detail = build_reporting_tables(canonical)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    course_counts.to_csv(OUTPUT_DIR / "credential_course_needs.csv", index=False)
    catalog_counts.to_csv(OUTPUT_DIR / "credential_catalog_counts.csv", index=False)
    student_detail.to_csv(OUTPUT_DIR / "student_detail.csv", index=False)
    exclusions.to_csv(OUTPUT_DIR / "qa_exclusions.csv", index=False)
    conflicts.to_csv(OUTPUT_DIR / "catalog_need_conflicts.csv", index=False)
    credential_meta.to_csv(OUTPUT_DIR / "scope_credentials.csv", index=False)

    build_workbook(
        course_counts=course_counts,
        catalog_counts=catalog_counts,
        student_detail=student_detail,
        qa_exclusions=exclusions,
        catalog_conflicts=conflicts,
        credential_meta=credential_meta,
    )

    print("\n" + "=" * 100)
    print("REPORT RESULTS")
    print("=" * 100)
    print(f"Resolved audit rows before catalog selection: {len(resolved):,}")
    print(f"Canonical student-credential opportunities:   {len(canonical):,}")
    print(f"Unique students represented:                  {canonical['student_id'].nunique() if not canonical.empty else 0:,}")
    print(f"Credential/course count rows:                 {len(course_counts):,}")
    print(f"Catalog count rows:                           {len(catalog_counts):,}")
    print(f"QA exclusions:                                {len(exclusions):,}")
    print(f"Catalog-dependent need rows:                  {len(conflicts):,}")
    print(f"Workbook:                                     {WORKBOOK_PATH}")
    print("=" * 100)


if __name__ == "__main__":
    main()
