from pathlib import Path
import re
import shutil
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path.cwd()

SRC = (
    ROOT
    / "data/processed/reporting/"
    "certificate_lineage_bruteforce_review_20260811"
)

DST = (
    ROOT
    / "data/processed/reporting/"
    "certificate_lineage_bruteforce_review_20260811_CERT_ONLY"
)

WORKBOOK = DST / "CERTIFICATE_LINEAGE_BRUTE_FORCE_REVIEW_CERT_ONLY.xlsx"

ASSOC_LINEAGE = "LIBERAL_ARTS"

DST.mkdir(parents=True, exist_ok=True)


def normalize(value):
    text = str(value or "").strip().upper()
    text = re.sub(r"\s+", "_", text)
    return text


def relevant_columns(df):
    tokens = (
        "lineage",
        "credential",
        "title",
    )

    return [
        c for c in df.columns
        if any(t in c.lower() for t in tokens)
    ]


def liberal_arts_mask(df):
    cols = relevant_columns(df)

    if not cols:
        return pd.Series(False, index=df.index)

    mask = pd.Series(False, index=df.index)

    for col in cols:
        values = df[col].fillna("").map(normalize)

        mask |= (
            values.eq(ASSOC_LINEAGE)
            | values.str.startswith(ASSOC_LINEAGE + "_")
            | values.str.contains(
                r"(^|[|;,])LIBERAL_ARTS($|[|;,])",
                regex=True,
            )
        )

    return mask


audit_rows = []
filtered_files = []

csvs = sorted(SRC.glob("*.csv"))

if not csvs:
    raise RuntimeError(
        f"No cached CSV files found in {SRC}"
    )

for src in csvs:

    df = pd.read_csv(
        src,
        dtype=str,
        low_memory=False,
    ).fillna("")

    before = len(df)

    mask = liberal_arts_mask(df)

    removed = int(mask.sum())

    out = df.loc[~mask].copy()

    after = len(out)

    # Hard verification
    remaining = int(
        liberal_arts_mask(out).sum()
    )

    if remaining:
        raise RuntimeError(
            f"{src.name}: "
            f"{remaining} LIBERAL_ARTS rows remain."
        )

    dst = DST / src.name

    out.to_csv(
        dst,
        index=False,
    )

    filtered_files.append(dst)

    audit_rows.append(
        {
            "file": src.name,
            "rows_before": before,
            "rows_removed_liberal_arts": removed,
            "rows_after": after,
        }
    )


audit = pd.DataFrame(audit_rows)

audit.to_csv(
    DST / "00_CERT_ONLY_FILTER_AUDIT.csv",
    index=False,
)


# ---------------------------------------------------------
# Rebuild workbook from filtered cached CSVs
# ---------------------------------------------------------

wb = Workbook()

# Remove default sheet after first real sheet is created.
default = wb.active

used_sheet_names = set()


def safe_sheet_name(stem):
    name = re.sub(
        r"^\d+[_-]*",
        "",
        stem,
    )

    name = re.sub(
        r"[\[\]:*?/\\]",
        "_",
        name,
    )

    name = name[:31] or "Sheet"

    base = name
    n = 2

    while name in used_sheet_names:
        suffix = f"_{n}"
        name = (
            base[:31-len(suffix)]
            + suffix
        )
        n += 1

    used_sheet_names.add(name)
    return name


# Filter audit first
files_for_workbook = [
    DST / "00_CERT_ONLY_FILTER_AUDIT.csv"
] + [
    p for p in filtered_files
    if p.name != "00_CERT_ONLY_FILTER_AUDIT.csv"
]


for path in files_for_workbook:

    df = pd.read_csv(
        path,
        dtype=str,
        low_memory=False,
    ).fillna("")

    ws = wb.create_sheet(
        safe_sheet_name(path.stem)
    )

    # Header
    for c_idx, col in enumerate(
        df.columns,
        start=1,
    ):
        cell = ws.cell(
            row=1,
            column=c_idx,
            value=col,
        )

        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    # Body
    for r_idx, row in enumerate(
        df.itertuples(
            index=False,
            name=None,
        ),
        start=2,
    ):
        for c_idx, value in enumerate(
            row,
            start=1,
        ):
            ws.cell(
                row=r_idx,
                column=c_idx,
                value=value,
            )

    ws.freeze_panes = "A2"

    if df.columns.size:
        ws.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(len(df.columns))}"
            f"{len(df)+1}"
        )

    # Bounded widths
    for c_idx, col in enumerate(
        df.columns,
        start=1,
    ):
        values = [
            str(col)
        ]

        if len(df):
            values += (
                df[col]
                .astype(str)
                .head(250)
                .tolist()
            )

        width = min(
            max(
                max(
                    len(v)
                    for v in values
                ) + 2,
                10,
            ),
            42,
        )

        ws.column_dimensions[
            get_column_letter(c_idx)
        ].width = width


wb.remove(default)

wb.save(WORKBOOK)


# ---------------------------------------------------------
# Final QA
# ---------------------------------------------------------

total_removed = int(
    audit[
        "rows_removed_liberal_arts"
    ].sum()
)

files_changed = int(
    (
        audit[
            "rows_removed_liberal_arts"
        ] > 0
    ).sum()
)

print("=" * 100)
print("CERTIFICATE-ONLY CACHED REVIEW REBUILD")
print("=" * 100)

print(
    "Cached CSV files:",
    len(csvs),
)

print(
    "Files containing Liberal Arts:",
    files_changed,
)

print(
    "Total cached rows removed:",
    total_removed,
)

print()

print(
    audit[
        audit[
            "rows_removed_liberal_arts"
        ] > 0
    ].to_string(index=False)
)

print()
print("PASS: no LIBERAL_ARTS rows remain in filtered cached outputs.")

print()
print("WROTE:")
print(DST)

print()
print("WORKBOOK:")
print(WORKBOOK)
