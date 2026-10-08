from pathlib import Path
import re
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


# =============================================================================
# CASES
# =============================================================================

CASES = [
    {
        "student_id": "R80016075",
        "catalog_year": "2021-2022",
        "lineage": "PROCESS_TECHNOLOGY_AAS",
        "label": "Process Technology AAS",
    },
    {
        "student_id": "R80078992",
        "catalog_year": "2022-2023",
        "lineage": "LIBERAL_ARTS",
        "label": "Liberal Arts",
    },
]

TARGET_IDS = {case["student_id"] for case in CASES}

ROOT = Path(".")
COURSE_HISTORY = (
    ROOT
    / "data/processed/"
      "normalized_actual_student_course_history.csv"
)

AUDIT_ROOT = (
    ROOT
    / "data/processed/"
      "full_actual_audit"
)

OUT_DIR = (
    ROOT
    / "data/processed/reporting/"
      "degreeworks_discrepancy_qa_20260803"
)

OUT_DIR.mkdir(parents=True, exist_ok=True)

OUT_PATH = (
    OUT_DIR
    / "LSCO_DegreeWorks_Discrepancy_Detailed_Audits.xlsx"
)


# =============================================================================
# HELPERS
# =============================================================================

def clean(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalize(value):
    value = clean(value).upper()
    value = re.sub(r"[^A-Z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalize_column(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def find_column(columns, aliases, required=False):
    lookup = {
        normalize_column(column): column
        for column in columns
    }

    for alias in aliases:
        key = normalize_column(alias)
        if key in lookup:
            return lookup[key]

    if required:
        raise KeyError(
            f"Could not find required column from aliases {aliases}.\n"
            f"Available columns:\n{list(columns)}"
        )

    return None


def numeric(value):
    return pd.to_numeric(value, errors="coerce")


def safe_sheet_name(value):
    value = re.sub(r"[\[\]:*?/\\]", "_", value)
    return value[:31]


def extract_chunk_token(path):
    numbers = re.findall(r"\d+", path.stem)

    if not numbers:
        return None

    return numbers[-1].lstrip("0") or "0"


def find_candidate_files(root, include_words):
    matches = []

    for path in root.rglob("*.csv"):
        text = path.name.lower()

        if all(word in text for word in include_words):
            matches.append(path)

    return sorted(matches)


# =============================================================================
# LOAD COURSE HISTORIES
# =============================================================================

if not COURSE_HISTORY.exists():
    raise FileNotFoundError(
        f"Course-history file not found:\n{COURSE_HISTORY}"
    )

course_header = pd.read_csv(
    COURSE_HISTORY,
    nrows=0,
).columns

course_student_col = find_column(
    course_header,
    ["student_id", "id", "studentid"],
    required=True,
)

course_parts = []

for chunk in pd.read_csv(
    COURSE_HISTORY,
    dtype="string",
    low_memory=False,
    chunksize=50_000,
):
    ids = (
        chunk[course_student_col]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    matched = chunk.loc[ids.isin(TARGET_IDS)].copy()

    if not matched.empty:
        matched[course_student_col] = ids.loc[matched.index]
        course_parts.append(matched)

if not course_parts:
    raise RuntimeError(
        "No course-history rows were found for either target student."
    )

course_history = pd.concat(
    course_parts,
    ignore_index=True,
).fillna("")

print("=" * 100)
print("COURSE HISTORY")
print("=" * 100)

for student_id in sorted(TARGET_IDS):
    count = (
        course_history[course_student_col]
        .eq(student_id)
        .sum()
    )
    print(f"{student_id}: {count:,} course rows")


# =============================================================================
# LOCATE AUDIT FILES
# =============================================================================

if not AUDIT_ROOT.exists():
    raise FileNotFoundError(
        f"Audit output directory not found:\n{AUDIT_ROOT}"
    )

summary_files = find_candidate_files(
    AUDIT_ROOT,
    ["summary"],
)

detail_files = [
    path
    for path in AUDIT_ROOT.rglob("*.csv")
    if "detail" in path.name.lower()
]

print()
print("=" * 100)
print("AUDIT FILE DISCOVERY")
print("=" * 100)
print(f"Summary files found: {len(summary_files):,}")
print(f"Detail files found:  {len(detail_files):,}")

if not detail_files:
    raise RuntimeError(
        "No audit detail CSV files were found under "
        f"{AUDIT_ROOT}."
    )


# =============================================================================
# FIND TARGET STUDENTS' CHUNKS FROM SUMMARY FILES
# =============================================================================

student_chunk_tokens = {
    student_id: set()
    for student_id in TARGET_IDS
}

for path in summary_files:
    try:
        header = pd.read_csv(path, nrows=0).columns
    except Exception:
        continue

    student_col = find_column(
        header,
        ["student_id", "id", "studentid"],
    )

    if not student_col:
        continue

    try:
        summary = pd.read_csv(
            path,
            usecols=[student_col],
            dtype="string",
            low_memory=False,
        )
    except Exception:
        continue

    ids = (
        summary[student_col]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    present = TARGET_IDS.intersection(set(ids))

    if present:
        token = extract_chunk_token(path)

        for student_id in present:
            if token is not None:
                student_chunk_tokens[student_id].add(token)

for student_id, tokens in student_chunk_tokens.items():
    print(
        f"{student_id} summary chunk token(s): "
        f"{sorted(tokens) if tokens else 'not resolved'}"
    )


# =============================================================================
# SELECT LIKELY DETAIL FILES
# =============================================================================

candidate_detail_files = set()

all_known_tokens = set().union(
    *student_chunk_tokens.values()
)

if all_known_tokens:
    for path in detail_files:
        token = extract_chunk_token(path)

        if token in all_known_tokens:
            candidate_detail_files.add(path)

# If token matching failed, fall back to all detail files.
if not candidate_detail_files:
    print(
        "Chunk-to-detail filename mapping was not resolved; "
        "using bounded fallback scan of all detail files."
    )
    candidate_detail_files = set(detail_files)

candidate_detail_files = sorted(candidate_detail_files)

print(
    f"Candidate detail files to inspect: "
    f"{len(candidate_detail_files):,}"
)


# =============================================================================
# EXTRACT TARGET DETAIL ROWS
# =============================================================================

detail_parts = []
files_read = 0

for path in candidate_detail_files:
    try:
        header = pd.read_csv(path, nrows=0).columns
    except Exception:
        continue

    student_col = find_column(
        header,
        ["student_id", "id", "studentid"],
    )

    if not student_col:
        continue

    files_read += 1

    try:
        for chunk in pd.read_csv(
            path,
            dtype="string",
            low_memory=False,
            chunksize=100_000,
        ):
            ids = (
                chunk[student_col]
                .fillna("")
                .astype(str)
                .str.strip()
            )

            matched = chunk.loc[
                ids.isin(TARGET_IDS)
            ].copy()

            if not matched.empty:
                matched[student_col] = ids.loc[
                    matched.index
                ]
                matched["_source_detail_file"] = str(path)
                detail_parts.append(matched)

    except Exception as exc:
        print(
            f"WARNING: Could not fully read {path.name}: {exc}"
        )

if not detail_parts and len(candidate_detail_files) < len(detail_files):
    print(
        "No rows found in inferred detail files. "
        "Running fallback scan across all detail files."
    )

    already_read = set(candidate_detail_files)

    for path in detail_files:
        if path in already_read:
            continue

        try:
            header = pd.read_csv(path, nrows=0).columns
        except Exception:
            continue

        student_col = find_column(
            header,
            ["student_id", "id", "studentid"],
        )

        if not student_col:
            continue

        for chunk in pd.read_csv(
            path,
            dtype="string",
            low_memory=False,
            chunksize=100_000,
        ):
            ids = (
                chunk[student_col]
                .fillna("")
                .astype(str)
                .str.strip()
            )

            matched = chunk.loc[
                ids.isin(TARGET_IDS)
            ].copy()

            if not matched.empty:
                matched[student_col] = ids.loc[
                    matched.index
                ]
                matched["_source_detail_file"] = str(path)
                detail_parts.append(matched)

if not detail_parts:
    raise RuntimeError(
        "No existing detailed-audit rows were found for "
        "either target student."
    )

audit_detail = pd.concat(
    detail_parts,
    ignore_index=True,
).fillna("")

audit_student_col = find_column(
    audit_detail.columns,
    ["student_id", "id", "studentid"],
    required=True,
)

audit_detail[audit_student_col] = (
    audit_detail[audit_student_col]
    .astype(str)
    .str.strip()
)

print()
print("=" * 100)
print("AUDIT DETAIL EXTRACTED")
print("=" * 100)

for student_id in sorted(TARGET_IDS):
    count = (
        audit_detail[audit_student_col]
        .eq(student_id)
        .sum()
    )
    print(f"{student_id}: {count:,} raw detail rows")


# =============================================================================
# IDENTIFY IMPORTANT AUDIT COLUMNS
# =============================================================================

catalog_col = find_column(
    audit_detail.columns,
    [
        "catalog_year",
        "catalog",
        "eligible_catalog_year",
    ],
)

lineage_col = find_column(
    audit_detail.columns,
    [
        "canonical_lineage",
        "detected_lineage",
        "credential_lineage",
        "lineage",
    ],
)

credential_col = find_column(
    audit_detail.columns,
    [
        "credential_id",
        "credential_name",
        "program_name",
        "degree_name",
        "credential",
    ],
)

requirement_col = find_column(
    audit_detail.columns,
    [
        "requirement_id",
        "requirement",
        "requirement_name",
        "requirement_label",
        "requirement_group",
    ],
)

rule_type_col = find_column(
    audit_detail.columns,
    [
        "rule_type",
        "requirement_type",
        "option_type",
    ],
)

required_course_col = find_column(
    audit_detail.columns,
    [
        "required_course",
        "requirement_course",
        "course_required",
        "required_course_code",
        "option_course",
    ],
)

matched_course_col = find_column(
    audit_detail.columns,
    [
        "matched_course",
        "used_course",
        "course_used",
        "applied_course",
        "matched_course_code",
        "course_code",
    ],
)

status_col = find_column(
    audit_detail.columns,
    [
        "status",
        "requirement_status",
        "match_status",
        "completion_status",
        "result",
    ],
)

required_hours_col = find_column(
    audit_detail.columns,
    [
        "required_hours",
        "hours_required",
        "requirement_hours",
        "required_credits",
    ],
)

applied_hours_col = find_column(
    audit_detail.columns,
    [
        "applied_hours",
        "matched_hours",
        "hours_applied",
        "earned_hours",
        "credits_applied",
    ],
)

missing_hours_col = find_column(
    audit_detail.columns,
    [
        "missing_hours",
        "hours_missing",
        "remaining_hours",
        "credits_remaining",
    ],
)

grade_col = find_column(
    audit_detail.columns,
    [
        "grade",
        "grade_code",
        "final_grade",
    ],
)

term_col = find_column(
    audit_detail.columns,
    [
        "term",
        "term_code",
        "course_term",
        "term_numeric",
    ],
)


# =============================================================================
# FILTER EACH REQUESTED AUDIT
# =============================================================================

case_outputs = []

for case in CASES:
    student_id = case["student_id"]

    rows = audit_detail.loc[
        audit_detail[audit_student_col].eq(student_id)
    ].copy()

    if catalog_col:
        catalog_mask = (
            rows[catalog_col]
            .astype(str)
            .str.strip()
            .eq(case["catalog_year"])
        )

        if catalog_mask.any():
            rows = rows.loc[catalog_mask].copy()

    search_columns = [
        column
        for column in [
            lineage_col,
            credential_col,
        ]
        if column
    ]

    if search_columns:
        row_text = (
            rows[search_columns]
            .astype(str)
            .agg(" | ".join, axis=1)
            .map(normalize)
        )

        target_words = normalize(
            case["lineage"]
        ).split()

        lineage_mask = pd.Series(
            True,
            index=rows.index,
        )

        for word in target_words:
            lineage_mask &= row_text.str.contains(
                rf"\b{re.escape(word)}\b",
                regex=True,
                na=False,
            )

        if lineage_mask.any():
            rows = rows.loc[lineage_mask].copy()

    if rows.empty:
        raise RuntimeError(
            f"No matching detailed audit remained for "
            f"{student_id}, {case['label']}, "
            f"{case['catalog_year']}."
        )

    # Remove exact duplicated rows caused by multiple file reads.
    rows = rows.drop_duplicates().copy()

    sort_columns = [
        column
        for column in [
            requirement_col,
            required_course_col,
            matched_course_col,
            term_col,
        ]
        if column
    ]

    if sort_columns:
        rows = rows.sort_values(
            sort_columns,
            na_position="last",
        )

    case_outputs.append(
        {
            **case,
            "detail": rows,
            "courses": course_history.loc[
                course_history[course_student_col]
                .eq(student_id)
            ].copy(),
        }
    )


# =============================================================================
# BUILD A COMPACT REQUIREMENT VIEW
# =============================================================================

def build_compact_view(detail):
    mapping = [
        ("Requirement", requirement_col),
        ("Rule Type", rule_type_col),
        ("Required Course / Option", required_course_col),
        ("Matched Course", matched_course_col),
        ("Grade", grade_col),
        ("Term", term_col),
        ("Required Hours", required_hours_col),
        ("Applied Hours", applied_hours_col),
        ("Missing Hours", missing_hours_col),
        ("Status", status_col),
        ("Catalog", catalog_col),
        ("Credential / Lineage", lineage_col or credential_col),
        ("Source Detail File", "_source_detail_file"),
    ]

    available = [
        (label, column)
        for label, column in mapping
        if column and column in detail.columns
    ]

    if not available:
        return detail.copy()

    compact = pd.DataFrame()

    for label, column in available:
        compact[label] = detail[column].astype(str)

    return compact.drop_duplicates().reset_index(drop=True)


for case in case_outputs:
    case["compact"] = build_compact_view(
        case["detail"]
    )


# =============================================================================
# WRITE CSV EVIDENCE FILES
# =============================================================================

for case in case_outputs:
    student_id = case["student_id"]
    slug = (
        f"{student_id}_"
        f"{case['lineage']}_"
        f"{case['catalog_year']}"
    )

    case["courses"].to_csv(
        OUT_DIR / f"{slug}_course_history.csv",
        index=False,
        encoding="utf-8-sig",
    )

    case["detail"].to_csv(
        OUT_DIR / f"{slug}_raw_audit_detail.csv",
        index=False,
        encoding="utf-8-sig",
    )

    case["compact"].to_csv(
        OUT_DIR / f"{slug}_requirement_review.csv",
        index=False,
        encoding="utf-8-sig",
    )


# =============================================================================
# EXCEL FORMATTING
# =============================================================================

DARK_GREEN = "1F4E3D"
MEDIUM_GREEN = "548235"
LIGHT_GREEN = "E2F0D9"
LIGHT_BLUE = "DDEBF7"
LIGHT_RED = "FCE4D6"
LIGHT_GOLD = "FFF2CC"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
BORDER_COLOR = "B7B7B7"

thin_border = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)


def write_dataframe(
    ws,
    dataframe,
    start_row,
    start_col=1,
    table_name=None,
    max_width=45,
):
    if dataframe.empty:
        ws.cell(
            row=start_row,
            column=start_col,
            value="No rows found",
        )
        return start_row

    for offset, column in enumerate(
        dataframe.columns,
        start=start_col,
    ):
        cell = ws.cell(
            row=start_row,
            column=offset,
            value=str(column),
        )
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill(
            "solid",
            fgColor=MEDIUM_GREEN,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    for row_offset, (_, row) in enumerate(
        dataframe.iterrows(),
        start=start_row + 1,
    ):
        for col_offset, value in enumerate(
            row.tolist(),
            start=start_col,
        ):
            cell = ws.cell(
                row=row_offset,
                column=col_offset,
                value=clean(value),
            )
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            text = normalize(value)

            if any(
                word in text
                for word in [
                    "MISSING",
                    "INCOMPLETE",
                    "UNMET",
                    "FAIL",
                    "SHORT",
                ]
            ):
                cell.fill = PatternFill(
                    "solid",
                    fgColor=LIGHT_RED,
                )

            elif any(
                word in text
                for word in [
                    "COMPLETE",
                    "SATISFIED",
                    "MATCHED",
                    "PASS",
                ]
            ):
                cell.fill = PatternFill(
                    "solid",
                    fgColor=LIGHT_GREEN,
                )

    last_row = start_row + len(dataframe)
    last_col = start_col + len(dataframe.columns) - 1

    for idx, column in enumerate(
        dataframe.columns,
        start=start_col,
    ):
        lengths = [
            len(str(column)),
        ]

        sample = dataframe[column].astype(str).head(100)

        lengths.extend(
            len(value)
            for value in sample
        )

        width = min(
            max(max(lengths) + 2, 11),
            max_width,
        )

        ws.column_dimensions[
            ws.cell(row=1, column=idx).column_letter
        ].width = width

    if table_name:
        table = Table(
            displayName=table_name,
            ref=(
                f"{ws.cell(start_row, start_col).coordinate}:"
                f"{ws.cell(last_row, last_col).coordinate}"
            ),
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        ws.add_table(table)

    return last_row


# =============================================================================
# CREATE WORKBOOK
# =============================================================================

wb = Workbook()
cover = wb.active
cover.title = "QA Summary"

cover.merge_cells("A1:G1")
cover["A1"] = (
    "LSCO Credential Audit — DegreeWorks Discrepancy QA"
)
cover["A1"].font = Font(
    bold=True,
    color=WHITE,
    size=16,
)
cover["A1"].fill = PatternFill(
    "solid",
    fgColor=DARK_GREEN,
)
cover["A1"].alignment = Alignment(
    horizontal="center",
    vertical="center",
)
cover.row_dimensions[1].height = 28

cover.merge_cells("A2:G2")
cover["A2"] = (
    "Both students were classified COMPLETE by the LSCO audit "
    "while DegreeWorks reportedly shows one unmet course. "
    "This workbook preserves the requirement-level and "
    "course-history evidence for reconciliation."
)
cover["A2"].fill = PatternFill(
    "solid",
    fgColor=LIGHT_GREEN,
)
cover["A2"].alignment = Alignment(
    wrap_text=True,
    vertical="center",
)
cover.row_dimensions[2].height = 48

summary_headers = [
    "R#",
    "Audit Tested",
    "Catalog",
    "LSCO Audit Result",
    "DegreeWorks Result",
    "QA Disposition",
    "Reviewer Notes",
]

for col, header in enumerate(summary_headers, start=1):
    cell = cover.cell(
        row=4,
        column=col,
        value=header,
    )
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill(
        "solid",
        fgColor=MEDIUM_GREEN,
    )
    cell.border = thin_border
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

for row_number, case in enumerate(
    case_outputs,
    start=5,
):
    values = [
        case["student_id"],
        case["label"],
        case["catalog_year"],
        "COMPLETE — existing modeled result",
        "One course missing — reported by reviewer",
        "Pending detailed reconciliation",
        "",
    ]

    for col, value in enumerate(values, start=1):
        cell = cover.cell(
            row=row_number,
            column=col,
            value=value,
        )
        cell.border = thin_border
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        if col in [1, 6, 7]:
            cell.fill = PatternFill(
                "solid",
                fgColor=LIGHT_BLUE,
            )

    cover.row_dimensions[row_number].height = 55

cover_widths = {
    "A": 15,
    "B": 34,
    "C": 15,
    "D": 30,
    "E": 34,
    "F": 30,
    "G": 50,
}

for column, width in cover_widths.items():
    cover.column_dimensions[column].width = width

cover.freeze_panes = "A5"
cover.sheet_view.showGridLines = False


# Individual audit sheets
for case_index, case in enumerate(
    case_outputs,
    start=1,
):
    student_id = case["student_id"]

    audit_ws = wb.create_sheet(
        safe_sheet_name(
            f"{student_id} Audit"
        )
    )

    audit_ws.merge_cells("A1:M1")
    audit_ws["A1"] = (
        f"{student_id} — {case['label']} "
        f"({case['catalog_year']} Catalog)"
    )
    audit_ws["A1"].font = Font(
        bold=True,
        color=WHITE,
        size=15,
    )
    audit_ws["A1"].fill = PatternFill(
        "solid",
        fgColor=DARK_GREEN,
    )
    audit_ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    audit_ws.merge_cells("A2:M2")
    audit_ws["A2"] = (
        "Requirement-level evidence from the existing LSCO "
        "audit. Red cells indicate text suggesting an unmet "
        "or incomplete condition; green cells indicate "
        "matched or satisfied conditions."
    )
    audit_ws["A2"].fill = PatternFill(
        "solid",
        fgColor=LIGHT_GREEN,
    )
    audit_ws["A2"].alignment = Alignment(
        wrap_text=True,
        vertical="center",
    )

    write_dataframe(
        audit_ws,
        case["compact"],
        start_row=4,
        table_name=f"AuditCase{case_index}",
    )

    audit_ws.freeze_panes = "A5"
    audit_ws.sheet_view.showGridLines = False

    course_ws = wb.create_sheet(
        safe_sheet_name(
            f"{student_id} Courses"
        )
    )

    course_ws.merge_cells("A1:Z1")
    course_ws["A1"] = (
        f"{student_id} — Complete Course History"
    )
    course_ws["A1"].font = Font(
        bold=True,
        color=WHITE,
        size=15,
    )
    course_ws["A1"].fill = PatternFill(
        "solid",
        fgColor=DARK_GREEN,
    )
    course_ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    write_dataframe(
        course_ws,
        case["courses"],
        start_row=3,
        table_name=f"CoursesCase{case_index}",
        max_width=35,
    )

    course_ws.freeze_panes = "A4"
    course_ws.sheet_view.showGridLines = False

    raw_ws = wb.create_sheet(
        safe_sheet_name(
            f"{student_id} Raw"
        )
    )

    raw_ws["A1"] = (
        f"{student_id} — Raw Existing Audit Detail"
    )
    raw_ws["A1"].font = Font(
        bold=True,
        color=WHITE,
        size=14,
    )
    raw_ws["A1"].fill = PatternFill(
        "solid",
        fgColor=DARK_GREEN,
    )

    write_dataframe(
        raw_ws,
        case["detail"],
        start_row=3,
        table_name=f"RawCase{case_index}",
        max_width=35,
    )

    raw_ws.freeze_panes = "A4"
    raw_ws.sheet_view.showGridLines = False


# =============================================================================
# SAVE
# =============================================================================

wb.save(OUT_PATH)

print()
print("=" * 100)
print("DEGREEWORKS DISCREPANCY QA WORKBOOK CREATED")
print("=" * 100)

for case in case_outputs:
    print(
        f"{case['student_id']} | "
        f"{case['label']} | "
        f"{case['catalog_year']} | "
        f"{len(case['compact']):,} requirement-detail rows | "
        f"{len(case['courses']):,} course-history rows"
    )

print()
print(f"Workbook:\n{OUT_PATH}")
print(f"\nEvidence directory:\n{OUT_DIR}")
