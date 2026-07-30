"""
PHASE 3 — build the outlined "Credential Course Needs" workbook from the
Phase 2 aggregate CSV. Reads/writes aggregate data only; no student_id
anywhere in this script.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IN_PATH = r"R:\data\processed\reporting\course_needs_aggregate.csv"
OUT_PATH = r"R:\data\processed\reporting\credential_course_needs_report.xlsx"

BOLD = Font(bold=True)
NORMAL = Font(bold=False)
THIN_BOTTOM = Border(bottom=Side(style="thin", color="BBBBBB"))
RIGHT = Alignment(horizontal="right")


def credential_display(title, level):
    return f"{title} · {level}" if level else title


def _setup_sheet(ws):
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.column_dimensions["A"].width = 65
    ws.column_dimensions["B"].width = 10


def build_rollup_sheet(ws, df, breakdown_col, sort_breakdown_by_index=False):
    """
    breakdown_col: 'course_label' for the course-needs view,
                   'catalog_year' for the catalog-year view.
    """
    _setup_sheet(ws)
    row = 1

    for family, fam_df in df.groupby("credential_family", sort=True):
        ws.cell(row=row, column=1, value=family.replace("_", " ").title()).font = BOLD
        ws.row_dimensions[row].outlineLevel = 0
        row += 1

        cred_keys = fam_df[["credential_id", "credential_title", "credential_level"]].drop_duplicates()
        cred_keys = cred_keys.sort_values(["credential_title", "credential_level"])

        for _, cred in cred_keys.iterrows():
            cred_df = fam_df[fam_df["credential_id"] == cred["credential_id"]]
            total = int(cred_df["unmet_count"].sum())
            label = credential_display(cred["credential_title"], cred["credential_level"])

            c1 = ws.cell(row=row, column=1, value=label)
            c1.font = BOLD
            c1.border = THIN_BOTTOM
            c2 = ws.cell(row=row, column=2, value=total)
            c2.font = BOLD
            c2.alignment = RIGHT
            c2.border = THIN_BOTTOM
            ws.row_dimensions[row].outlineLevel = 1
            row += 1

            rollup = cred_df.groupby(breakdown_col)["unmet_count"].sum()
            rollup = rollup.sort_index() if sort_breakdown_by_index else rollup.sort_values(ascending=False)

            for label_text, count in rollup.items():
                if int(count) == 0:
                    continue  # zero rows omitted, not printed
                ws.cell(row=row, column=1, value=f"    {label_text}").font = NORMAL
                cc = ws.cell(row=row, column=2, value=int(count))
                cc.alignment = RIGHT
                ws.row_dimensions[row].outlineLevel = 2
                row += 1
        row += 1  # blank spacer between families


def build_detail_sheet(ws, df):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = BOLD
    for _, r in df.iterrows():
        ws.append(list(r))
    for i, col in enumerate(df.columns, start=1):
        width = max(14, min(45, len(str(col)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def main():
    df = pd.read_csv(IN_PATH, dtype=str)
    df["unmet_count"] = df["unmet_count"].astype(int)

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Course Needs"
    build_rollup_sheet(ws1, df, breakdown_col="course_label", sort_breakdown_by_index=False)

    ws2 = wb.create_sheet("Catalog Year Counts")
    build_rollup_sheet(ws2, df, breakdown_col="catalog_year", sort_breakdown_by_index=True)

    ws3 = wb.create_sheet("Detail")
    build_detail_sheet(ws3, df)

    wb.save(OUT_PATH)
    print(f"Wrote report: {OUT_PATH}")
    print(f"Sheets: {wb.sheetnames}")
    print(f"Total credentials represented: {df['credential_id'].nunique()}")
    print(f"Total unmet-requirement rows rolled up: {int(df['unmet_count'].sum())}")


if __name__ == "__main__":
    main()
