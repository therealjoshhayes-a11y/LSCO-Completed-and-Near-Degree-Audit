from __future__ import annotations
r"""READ-ONLY Fall 2026 reporting from previously computed audit evidence.

No engine imports/calls, reallocation, curriculum parsing or new audits.
Commands:
 python -u .\scripts\report_fall2026_from_saved_evidence.py
"""
import csv
from collections import Counter, defaultdict
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent.parent
PROCESSED=ROOT/"data/processed"
OUT=PROCESSED/"reporting/fall26_saved_evidence_reports"
COHORT_DIR=PROCESSED/"reporting"
INCREMENTAL=PROCESSED/"incremental_audit/fall_2026_refresh_20261007"
HIST_SUMMARY=PROCESSED/"full_actual_audit/full_actual_credential_summary.csv"
HIST_DETAIL=PROCESSED/"full_actual_audit/full_actual_audit_results.csv"
REQUIREMENTS=PROCESSED/"catalogs/staging_six_year/requirements_master_multiyear.csv"
IDENTITY=ROOT/"data/raw/student_exports/banner_course_history_actual.csv"
COLS=["Banner ID","Last Name","First Name","Declared Major","Catalog Year",
      "Credential ID","Evidence Status","Remaining Requirement Count (not whole classes)",
      "Outstanding Catalog Rules","Fall Registered Courses","Qualification"]

def require(condition,message):
    if not condition: raise RuntimeError(message)

def read(path,columns=None):
    require(path.is_file(),f"Required saved evidence missing: {path}")
    return pd.read_csv(path,dtype=str,usecols=columns,low_memory=False).fillna("")

def concat_csv(files):
    frames=[]
    for file in files:
        frames.append(read(file))
    return pd.concat(frames,ignore_index=True) if frames else pd.DataFrame()

def sheet(wb,name,records,headers=COLS):
    sh=wb.create_sheet(name);sh.append(headers)
    for row in records:sh.append(row)
    sh.freeze_panes="A2";sh.sheet_view.showGridLines=False
    sh.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{len(records)+1}"
    for c in sh[1]:
        c.fill=PatternFill("solid",fgColor="145A32")
        c.font=Font(bold=True,color="FFFFFF")
    for i,heading in enumerate(headers,1):
        sh.column_dimensions[get_column_letter(i)].width=65 if "Rule" in heading or "Courses" in heading else 24

def output_book(path,title,rows,notes):
    wb=Workbook();ws=wb.active;ws.title="CONTROL SUMMARY"
    for pair in [("Report",title),("Student scope","Fall 2026 term 202690"),
                 ("Rows",len(rows)),("Distinct students",len({x[0] for x in rows})),
                 ("Method","SAVED EVIDENCE ONLY: no degree audits or hypothetical course reallocation"),
                 ("Qualification",notes)]:ws.append(pair)
    byyear=defaultdict(list)
    for row in rows:byyear[row[4]].append(row)
    for year,items in sorted(byyear.items()):
        sheet(wb,year,sorted(items,key=lambda x:(x[1],x[2],x[0],x[5])))
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def main():
    source=sorted(COHORT_DIR.glob("fall_2026_products_4_5_inputs_*/RESTRICTED_fall2026_registered_course_inventory.csv"))
    require(source,"Run the Fall enrollment preflight first; no stored Fall cohort inventory found.")
    fall=read(source[-1])
    fall["Student_Id"]=fall.Student_Id.astype(str).str.strip()
    require(len(set(fall.Student_Id))>=8000,"Saved Fall enrollment cohort unexpectedly small.")
    registered=fall.groupby("Student_Id").course_code.apply(lambda x:"; ".join(sorted(set(x)))).to_dict()
    cohort=set(registered)
    # Reuse actual previously completed summaries. Delta only replaces matching
    # student/catalog/credential keys; never conflate an existing COMPLETE
    # status with a Fall-projected COMPLETE status.
    require(HIST_SUMMARY.is_file(),"Missing original empirical audit summary.")
    current_path=INCREMENTAL/"RESTRICTED_current_eligible_empirical_credential_summary.csv"
    if current_path.is_file():
        summaries=read(current_path)
        print("Using authoritative saved current empirical summary:",current_path)
    else:
        hist=read(HIST_SUMMARY)
        delta_path=INCREMENTAL/"RESTRICTED_delta_credential_summary.csv"
        require(delta_path.is_file(),
                "Current empirical summary and incremental delta absent. Refusing stale July-only report.")
        delta=read(delta_path)
        key=["student_id","catalog_year","credential_id"]
        hist=hist.set_index(key,drop=False)
        delta=delta.set_index(key,drop=False)
        hist=hist.loc[~hist.index.isin(delta.index)].reset_index(drop=True)
        summaries=pd.concat([hist,delta.reset_index(drop=True)],ignore_index=True)
    summaries=summaries[summaries.student_id.isin(cohort)].copy()
    require(not summaries.empty,"No precomputed audited combinations for Fall cohort.")
    for col in ("requirements_missing","requirements_unresolved","requirements_met"):
        if col not in summaries:summaries[col]=""
    # Read ONLY historical per-student chunks. Never open the enormous
    # full_actual_audit_results.csv combined file.
    wanted=set(zip(summaries.student_id,summaries.catalog_year,summaries.credential_id))
    detail_columns=["student_id","catalog_year","credential_id","requirement_id",
                    "status","path_group_id","path_selected"]
    indexed=defaultdict(list)
    chunk_dir=INCREMENTAL/"chunks"
    delta_files=sorted(chunk_dir.glob("chunk_*_detail.csv"))
    historical_chunks=PROCESSED/"full_actual_audit/chunks"
    historical_summary_files=sorted(historical_chunks.glob("chunk_*_credential_summary.csv"))
    historical_detail_files={p.name.replace("_audit_results.csv",""):
       p for p in historical_chunks.glob("chunk_*_audit_results.csv")}
    require(historical_summary_files or delta_files,
            "No saved audit chunks available; refusing scan of combined detail.")
    def key_of(row):
        return (row["student_id"],row["catalog_year"],row["credential_id"])
    def read_detail_chunk(path,allowed,override=None):
        with path.open("r",encoding="utf-8-sig",newline="") as handle:
            reader=csv.DictReader(handle)
            missing=set(detail_columns)-set(reader.fieldnames or [])
            require(not missing,
                    f"Missing selected-path provenance in {path.name}: {sorted(missing)}")
            for row in reader:
                key=key_of(row)
                if key not in allowed or (override is not None and key in override):
                    continue
                if row["path_group_id"].strip() and row["path_selected"].strip().upper()!="TRUE":
                    continue
                indexed[key].append({c:row[c] for c in detail_columns})
    # Delta precedence: identify exact keys from SMALL existing chunk summaries.
    delta_keys=set()
    for f in sorted(chunk_dir.glob("chunk_*_summary.csv")):
        with f.open("r",encoding="utf-8-sig",newline="") as handle:
            for row in csv.DictReader(handle):
                key=key_of(row)
                if key in wanted:delta_keys.add(key)
    # Only attempt historic chunk reads whose matching summary contains a
    # relevant student. This is a cheap index over 25-student files.
    relevant_historical=[]
    for sf in historical_summary_files:
        with sf.open("r",encoding="utf-8-sig",newline="") as handle:
            if any(row.get("student_id","") in cohort for row in csv.DictReader(handle)):
                df=historical_chunks/sf.name.replace("_credential_summary.csv","_audit_results.csv")
                require(df.is_file(),f"Missing historical detail chunk {df.name}")
                relevant_historical.append(df)
    print(f"Saved historical detail chunks selected: {len(relevant_historical)} of {len(historical_summary_files)}",flush=True)
    for f in relevant_historical:read_detail_chunk(f,wanted,override=delta_keys)
    for f in delta_files:read_detail_chunk(f,delta_keys)
    print(f"Matched saved student/catalog/credential detail keys: {len(indexed)}",flush=True)
    req=read(REQUIREMENTS)
    defs=req.groupby(["catalog_year","credential_id","requirement_id"],sort=False).agg(
       rule_text=("option_value",lambda x:" OR ".join(dict.fromkeys(y.strip() for y in x if y.strip()))),
       rule_kind=("option_type",lambda x:" / ".join(sorted(set(x))))
    ).reset_index()
    # Only read named/major student metadata, no requirements engine.
    names={}
    if IDENTITY.is_file():
        with IDENTITY.open(newline="",encoding="utf-8-sig") as stream:
            reader=csv.DictReader(stream)
            for r in reader:
                sid=(r.get("StudenID") or r.get("StudentID") or "").strip()
                if sid not in cohort:continue
                names.setdefault(sid,[(r.get("LastName") or ""),(r.get("FirstName") or ""),(r.get("StudentMajor") or "")])
    desc={(x.catalog_year,x.credential_id,x.requirement_id):(x.rule_text,x.rule_kind)
          for x in defs.itertuples(index=False)}
    precompleted=[];review=[];gap_counts=Counter()
    for rec in summaries.to_dict("records"):
        sid,year,cred=rec["student_id"],rec["catalog_year"],rec["credential_id"]
        key=(sid,year,cred)
        rows=indexed.get(key,[])
        if not rows:continue
        unmet=[x for x in rows if x["status"]!="MET"]
        if rec.get("audit_status")=="COMPLETE":
            # Existing empirical result, NOT a verified Fall projection.
            target=precompleted
            finding="EMPIRICALLY COMPLETE IN SAVED AUDIT"
        elif not unmet:
            continue
        else:
            # One unmet requirement ROW need not be one whole-class obligation.
            # Carry the actual rule and do not falsely label as 1/2/3 classes.
            target=review
            finding="UNKNOWN WHOLE-CLASS COUNT; REQUIRES EXISTING ALLOCATION EVIDENCE"
        rule_texts=[]
        for d in unmet:
            label,kind=desc.get((year,cred,d["requirement_id"]),("SOURCE RULE NOT FOUND",""))
            rule_texts.append(f'{d["requirement_id"]}: {label} [{kind}]')
        who=names.get(sid,["","",""])
        try:missing=int(rec.get("requirements_missing",""))
        except (ValueError,TypeError):missing=-1
        if target is review:gap_counts[missing]+=1
        target.append([sid,*who,year,cred,finding,
          "" if missing<0 else missing," | ".join(rule_texts),
          registered[sid],"Fall course-to-requirement projected allocation NOT in saved empirical audit"])
    OUT.mkdir(parents=True,exist_ok=True)
    output_book(OUT/"F26_Existing_Complete.xlsx",
                "Previously computed complete combinations among Fall enrollees",
                precompleted,"NOT newly projected Fall graduates or screened official awards")
    output_book(OUT/"F26_Unmet_Rules_Review.xlsx",
                "Saved unmet catalog rules among Fall enrollees",
                review,"Unmet requirement ROWS are not whole-class counts; do not publish as 1–3 remaining classes")
    print("READ-ONLY FALL EVIDENCE EXTRACTION")
    print("Fall students:",len(cohort))
    print("Saved summary combinations in cohort:",len(summaries))
    print("Saved empirical complete combinations:",len(precompleted))
    print("Combinations with unmet rules requiring whole-class/Fall allocation evidence:",len(review))
    print("Unmet requirement ROW distribution (not class counts):",dict(sorted(gap_counts.items())[:10]))
    print("Output:",OUT)
    print("Product 4 newly projected graduations: NOT CLAIMED.")
    print("Product 5 one-to-three whole classes: NOT CLAIMED.")
    print("No audit engine was imported or invoked.")

if __name__=="__main__":main()
