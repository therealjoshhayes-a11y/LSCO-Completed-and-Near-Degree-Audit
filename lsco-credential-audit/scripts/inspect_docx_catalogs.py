from pathlib import Path
import csv
import re
from docx import Document

from lsco_audit.catalog.registry import load_catalog_registry

COURSE_RE = re.compile(r"\b[A-Z]{3,4}\s+\d{4}\b")
TOTAL_RE = re.compile(r"\b(Total Program Hours|Semester Hours|Credit Hours)\b", re.I)

out_dir = Path("data/interim/catalogs")
out_dir.mkdir(parents=True, exist_ok=True)

summary_rows = []
toc_rows = []
table_rows = []

for record in load_catalog_registry():
    docx_path = Path(record.source_docx_path)
    doc = Document(docx_path)

    year_dir = out_dir / record.catalog_year
    year_dir.mkdir(parents=True, exist_ok=True)

    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)

    heading_count = 0
    course_para_count = 0

    for i, p in enumerate(doc.paragraphs):
        text = " ".join(p.text.split())
        style = p.style.name if p.style else ""

        if style.lower().startswith("heading"):
            heading_count += 1

        if COURSE_RE.search(text):
            course_para_count += 1

        if text and ("." * 5 in text or re.search(r"\s+\d{1,3}$", text)):
            if len(toc_rows) < 20000:
                toc_rows.append({
                    "catalog_year": record.catalog_year,
                    "paragraph_index": i,
                    "style": style,
                    "text": text,
                })

    course_table_rows = 0
    total_table_rows = 0

    for table_i, table in enumerate(doc.tables):
        for row_i, row in enumerate(table.rows):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            row_text = " | ".join(cells)

            has_course = bool(COURSE_RE.search(row_text))
            has_total = bool(TOTAL_RE.search(row_text))

            if has_course:
                course_table_rows += 1
            if has_total:
                total_table_rows += 1

            if has_course or has_total:
                table_rows.append({
                    "catalog_year": record.catalog_year,
                    "table_index": table_i,
                    "row_index": row_i,
                    "cell_count": len(cells),
                    "has_course": has_course,
                    "has_total": has_total,
                    "row_text": row_text,
                })

    summary_rows.append({
        "catalog_year": record.catalog_year,
        "docx_path": str(docx_path),
        "paragraph_count": para_count,
        "table_count": table_count,
        "heading_count": heading_count,
        "course_paragraph_count": course_para_count,
        "course_table_row_count": course_table_rows,
        "total_table_row_count": total_table_rows,
    })

def write_csv(path, rows, fieldnames):
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

write_csv(
    out_dir / "docx_structure_summary.csv",
    summary_rows,
    [
        "catalog_year",
        "docx_path",
        "paragraph_count",
        "table_count",
        "heading_count",
        "course_paragraph_count",
        "course_table_row_count",
        "total_table_row_count",
    ],
)

write_csv(
    out_dir / "docx_toc_candidates.csv",
    toc_rows,
    ["catalog_year", "paragraph_index", "style", "text"],
)

write_csv(
    out_dir / "docx_course_table_rows.csv",
    table_rows,
    [
        "catalog_year",
        "table_index",
        "row_index",
        "cell_count",
        "has_course",
        "has_total",
        "row_text",
    ],
)

print("Wrote:")
print(out_dir / "docx_structure_summary.csv")
print(out_dir / "docx_toc_candidates.csv")
print(out_dir / "docx_course_table_rows.csv")
