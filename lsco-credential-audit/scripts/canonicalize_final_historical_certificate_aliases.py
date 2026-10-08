from pathlib import Path
from datetime import datetime
import shutil
import pandas as pd
from openpyxl import load_workbook

ROOT = Path.cwd()
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

COMPLETE = (
    ROOT
    / "data/processed/full_actual_audit/president_report"
    / "complete_rows_with_canonical_lineage.csv"
)

CROSSWALK = (
    ROOT
    / "data/interim/institutional_awards"
    / "award_program_crosswalk_curated.xlsx"
)

GOV = (
    ROOT
    / "data/interim/institutional_awards"
    / "governed_lineage_relationships_final.csv"
)

ALIASES = {
    "GENERAL_STUDIES":
        "GENERAL_STUDIES_CORE_CURRICULUM",

    "ORDINARY_SEAMAN_BASIC_SAFETY_TRAINING":
        "ORDINARY_SEAMAN_I",

    "PHARMACY_TECHNOLOGY":
        "HOSPITAL_PHARMACY_TECHNOLOGY",

    "PHARMACY_TECHNOLOGY_BASIC":
        "RETAIL_PHARMACY_TECHNOLOGY_BASIC",

    "SAFETY_HEALTH_ENVIRONMENT_CERTIFICATE":
        "SAFETY_HEALTH_AND_ENVIRONMENT",

    "WELDING_TECHNOLOGY":
        "WELDING_TECHNOLOGY_CERTIFICATE_OF_COMPLETION",
}


def backup(path: Path, label: str):
    dst = path.with_name(
        f"{path.stem}_PRE_{label}_{STAMP}{path.suffix}"
    )
    shutil.copy2(path, dst)
    return dst


def yes(series):
    return (
        series.astype(str)
        .str.strip()
        .str.lower()
        .isin(["true", "1", "yes", "y"])
    )


print("=" * 110)
print("FINAL HISTORICAL CERTIFICATE CANONICALIZATION")
print("=" * 110)


# =============================================================================
# 1. MODELED COMPLETION ARTIFACT
# =============================================================================

complete = pd.read_csv(
    COMPLETE,
    dtype=str,
    low_memory=False,
).fillna("")

if "canonical_lineage" not in complete.columns:
    raise RuntimeError(
        "canonical_lineage missing from completion artifact."
    )

complete_backup = backup(
    COMPLETE,
    "FINAL_CERT_ALIAS_FIX"
)

before_counts = (
    complete["canonical_lineage"]
    .value_counts()
)

changed_complete = 0

for old, new in ALIASES.items():

    mask = complete[
        "canonical_lineage"
    ].eq(old)

    n = int(mask.sum())

    if n:
        print(
            f"MODELED: {old} -> {new}: "
            f"{n:,} rows"
        )

        complete.loc[
            mask,
            "canonical_lineage"
        ] = new

        changed_complete += n

complete.to_csv(
    COMPLETE,
    index=False,
)


# =============================================================================
# 2. CURATED OFFICIAL-AWARD CROSSWALK
# =============================================================================

crosswalk_backup = backup(
    CROSSWALK,
    "FINAL_CERT_ALIAS_FIX"
)

wb = load_workbook(CROSSWALK)

if "Crosswalk Draft" not in wb.sheetnames:
    raise RuntimeError(
        "Crosswalk Draft sheet missing."
    )

ws = wb["Crosswalk Draft"]

header_row = None
headers = {}

for r in range(
    1,
    min(ws.max_row, 20) + 1,
):

    probe = {
        str(
            ws.cell(r, c).value or ""
        ).strip(): c

        for c in range(
            1,
            ws.max_column + 1
        )
    }

    if "canonical_lineage" in probe:
        header_row = r
        headers = probe
        break

if header_row is None:
    raise RuntimeError(
        "Could not locate crosswalk headers."
    )

changed_crosswalk = 0

for r in range(
    header_row + 1,
    ws.max_row + 1,
):

    current = str(
        ws.cell(
            r,
            headers["canonical_lineage"]
        ).value or ""
    ).strip()

    if current not in ALIASES:
        continue

    target = ALIASES[current]

    ws.cell(
        r,
        headers["canonical_lineage"]
    ).value = target

    if "reconciliation_key" in headers:

        key = str(
            ws.cell(
                r,
                headers["reconciliation_key"]
            ).value or ""
        ).strip()

        if key == current:
            ws.cell(
                r,
                headers["reconciliation_key"]
            ).value = target

    if "match_status" in headers:
        ws.cell(
            r,
            headers["match_status"]
        ).value = "CURATED_MATCH"

    if "curation_note" in headers:

        old_note = str(
            ws.cell(
                r,
                headers["curation_note"]
            ).value or ""
        ).strip()

        addition = (
            "Catalog-validated historical canonical "
            f"identity: {current} -> {target}."
        )

        ws.cell(
            r,
            headers["curation_note"]
        ).value = (
            f"{old_note} {addition}".strip()
        )

    changed_crosswalk += 1

wb.save(CROSSWALK)

print()
print(
    "Official crosswalk rows canonicalized:",
    changed_crosswalk,
)


# =============================================================================
# 3. GOVERNANCE
#
# Normalize aliases anywhere they still occur. If that normalization turns
# a relationship into A -> A, the relationship is redundant and gets retired.
# =============================================================================

gov_backup = backup(
    GOV,
    "FINAL_CERT_ALIAS_FIX"
)

gov = pd.read_csv(
    GOV,
    dtype=str,
).fillna("")

required = {
    "detected_lineage",
    "institutional_lineage",
    "awarded_represents_detected",
    "detected_remains_additional_award",
}

missing = required - set(gov.columns)

if missing:
    raise RuntimeError(
        f"Governance missing columns: {sorted(missing)}"
    )

original_detected = (
    gov["detected_lineage"].copy()
)

original_institutional = (
    gov["institutional_lineage"].copy()
)

gov["detected_lineage"] = (
    gov["detected_lineage"]
    .replace(ALIASES)
)

gov["institutional_lineage"] = (
    gov["institutional_lineage"]
    .replace(ALIASES)
)

alias_changed = (
    original_detected.ne(
        gov["detected_lineage"]
    )
    |
    original_institutional.ne(
        gov["institutional_lineage"]
    )
)

became_self = (
    alias_changed
    &
    gov["detected_lineage"].eq(
        gov["institutional_lineage"]
    )
)

retired = gov.loc[
    became_self
].copy()

if len(retired):

    retired_path = GOV.with_name(
        "retired_canonical_alias_governance_"
        f"{STAMP}.csv"
    )

    retired.to_csv(
        retired_path,
        index=False,
    )

    print()
    print(
        "Redundant alias-governance rows retired:",
        len(retired),
    )
    print(
        "Retired-row audit:",
        retired_path,
    )

gov = gov.loc[
    ~became_self
].copy()

gov.to_csv(
    GOV,
    index=False,
)


# =============================================================================
# 4. QA
# =============================================================================

check_complete = pd.read_csv(
    COMPLETE,
    dtype=str,
    low_memory=False,
).fillna("")

old_remaining = sorted(
    set(
        check_complete[
            "canonical_lineage"
        ]
    )
    &
    set(ALIASES)
)

if old_remaining:
    raise RuntimeError(
        "Old aliases remain in modeled completion "
        f"artifact: {old_remaining}"
    )


check_cw = pd.read_excel(
    CROSSWALK,
    sheet_name="Crosswalk Draft",
    dtype=str,
).fillna("")

cw_old = sorted(
    set(
        check_cw[
            "canonical_lineage"
        ].astype(str)
    )
    &
    set(ALIASES)
)

if cw_old:
    raise RuntimeError(
        "Old aliases remain in award crosswalk: "
        f"{cw_old}"
    )


check_gov = pd.read_csv(
    GOV,
    dtype=str,
).fillna("")

consuming = check_gov[
    yes(
        check_gov[
            "awarded_represents_detected"
        ]
    )
    &
    ~yes(
        check_gov[
            "detected_remains_additional_award"
        ]
    )
].copy()


print()
print("=" * 110)
print("POST-REPAIR CONSUMING GOVERNANCE")
print("=" * 110)

print(
    "Count:",
    len(consuming),
)

if len(consuming):

    cols = [
        c for c in [
            "detected_lineage",
            "institutional_lineage",
            "relationship_type",
            "governance_note",
        ]
        if c in consuming.columns
    ]

    print()
    print(
        consuming[
            cols
        ].to_string(
            index=False
        )
    )

    raise RuntimeError(
        "Consuming governance remains."
    )

print("NONE")

print()
print("Modeled completion rows canonicalized:")
print(f"{changed_complete:,}")

print()
print("Completion backup:")
print(complete_backup)

print()
print("Crosswalk backup:")
print(crosswalk_backup)

print()
print("Governance backup:")
print(gov_backup)

print()
print("=" * 110)
print("PASS — CONSUMING GOVERNANCE LAYER CLOSED")
print("=" * 110)
