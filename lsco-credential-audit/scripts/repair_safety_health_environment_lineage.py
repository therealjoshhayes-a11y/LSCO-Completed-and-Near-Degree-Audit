from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
import os
import shutil
import tempfile

import pandas as pd
from openpyxl import load_workbook


ROOT = Path.cwd()

MODELED_PATH = (
    ROOT
    / "data/processed/full_actual_audit/president_report/"
    / "complete_rows_with_canonical_lineage.csv"
)

CROSSWALK_PATH = (
    ROOT
    / "data/interim/institutional_awards/"
    / "award_program_crosswalk_curated.xlsx"
)

TARGET = "SAFETY_HEALTH_AND_ENVIRONMENT"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")


def norm(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip().upper()


def fail(msg):
    raise RuntimeError(msg)


print("=" * 110)
print("SAFETY / HEALTH / ENVIRONMENT — CERTIFICATE LINEAGE REPAIR")
print("=" * 110)

for path in [MODELED_PATH, CROSSWALK_PATH]:
    if not path.exists():
        fail(f"Missing required file: {path}")


# =============================================================================
# 1. MODELED COMPLETION LINEAGE
# =============================================================================

modeled = pd.read_csv(
    MODELED_PATH,
    dtype=str,
    low_memory=False,
).fillna("")

required = {
    "credential_id",
    "mechanical_lineage",
    "canonical_lineage",
}

missing = required - set(modeled.columns)

if missing:
    fail(
        "Modeled lineage file missing columns: "
        + ", ".join(sorted(missing))
    )


def modeled_is_safety_certificate(row):
    values = " | ".join(
        norm(row.get(c, ""))
        for c in [
            "credential_id",
            "mechanical_lineage",
            "canonical_lineage",
        ]
    )

    if "SAFETY_HEALTH" not in values:
        return False

    # Absolutely do not collapse the associate degree.
    if "_AAS" in values or "ASSOCIATE" in values:
        return False

    return True


modeled_mask = modeled.apply(
    modeled_is_safety_certificate,
    axis=1,
)

modeled_targets = modeled.loc[
    modeled_mask,
    [
        "catalog_year",
        "credential_id",
        "mechanical_lineage",
        "canonical_lineage",
    ],
].drop_duplicates().sort_values(
    ["catalog_year", "credential_id"]
)

print()
print("MODELED CERTIFICATE GENERATIONS BEFORE REPAIR")
print(modeled_targets.to_string(index=False))

if modeled_targets.empty:
    fail(
        "No modeled Safety/Health/Environment certificate "
        "rows were detected. Aborting."
    )

aas_before = modeled.loc[
    modeled["canonical_lineage"]
    .str.upper()
    .str.contains("SAFETY_HEALTH", na=False)
    & modeled["canonical_lineage"]
    .str.upper()
    .str.contains("AAS", na=False)
].copy()

modeled_before_non_target = (
    modeled.loc[~modeled_mask]
    .copy()
    .reset_index(drop=True)
)

modeled.loc[
    modeled_mask,
    "canonical_lineage"
] = TARGET


# =============================================================================
# 2. AWARD CROSSWALK
# =============================================================================

wb = load_workbook(CROSSWALK_PATH)

if "Crosswalk Draft" not in wb.sheetnames:
    fail(
        "Workbook does not contain 'Crosswalk Draft'."
    )

ws = wb["Crosswalk Draft"]

headers = {
    norm(cell.value): cell.column
    for cell in ws[1]
    if cell.value is not None
}

needed_headers = {
    "CURR1PROGRAMCODE",
    "MAJOR1CODE",
    "DEGREECODE",
    "MAJORDESC",
    "DEGREEDESC",
    "CANONICAL_LINEAGE",
    "RECONCILIATION_KEY",
}

missing_headers = needed_headers - set(headers)

if missing_headers:
    fail(
        "Crosswalk sheet missing columns: "
        + ", ".join(sorted(missing_headers))
    )

crosswalk_target_rows = []

for r in range(2, ws.max_row + 1):

    program = norm(
        ws.cell(
            r,
            headers["CURR1PROGRAMCODE"]
        ).value
    )

    major = norm(
        ws.cell(
            r,
            headers["MAJOR1CODE"]
        ).value
    )

    degree = norm(
        ws.cell(
            r,
            headers["DEGREECODE"]
        ).value
    )

    major_desc = norm(
        ws.cell(
            r,
            headers["MAJORDESC"]
        ).value
    )

    degree_desc = norm(
        ws.cell(
            r,
            headers["DEGREEDESC"]
        ).value
    )

    old_lineage = norm(
        ws.cell(
            r,
            headers["CANONICAL_LINEAGE"]
        ).value
    )

    code_hit = any(
        token in {
            program,
            major,
        }
        or token in program
        or token in major
        for token in [
            "SHEC",
            "SHEP",
            "OSEH",
        ]
    )

    description_hit = (
        "SAFETY" in major_desc
        and "HEALTH" in major_desc
        and (
            "ENV" in major_desc
            or "ENVIRONMENT" in major_desc
        )
    )

    lineage_hit = (
        old_lineage.startswith(
            "SAFETY_HEALTH"
        )
    )

    is_certificate = (
        degree.startswith("CERT")
        or "CERTIFICATE" in degree_desc
    )

    is_aas = (
        degree == "AAS"
        or old_lineage.endswith("_AAS")
        or "ASSOCIATE" in degree_desc
    )

    if (
        (code_hit or description_hit or lineage_hit)
        and is_certificate
        and not is_aas
    ):
        crosswalk_target_rows.append(
            {
                "excel_row": r,
                "Curr1ProgramCode": program,
                "Major1Code": major,
                "DegreeCode": degree,
                "MajorDesc": major_desc,
                "old_canonical_lineage":
                    old_lineage,
            }
        )


crosswalk_targets = pd.DataFrame(
    crosswalk_target_rows
)

print()
print("OFFICIAL-AWARD CROSSWALK ROWS BEFORE REPAIR")

if crosswalk_targets.empty:
    fail(
        "No Safety/Health/Environment certificate "
        "crosswalk rows detected. Aborting."
    )

print(
    crosswalk_targets.to_string(
        index=False
    )
)

for item in crosswalk_target_rows:

    r = item["excel_row"]

    ws.cell(
        r,
        headers["CANONICAL_LINEAGE"]
    ).value = TARGET

    ws.cell(
        r,
        headers["RECONCILIATION_KEY"]
    ).value = TARGET


# =============================================================================
# 3. KEEP THE MODELED-INVENTORY TAB INTERNALLY CONSISTENT
# =============================================================================

inventory_updates = 0

if "Modeled Inventory" in wb.sheetnames:

    inv = wb["Modeled Inventory"]

    inv_headers = {
        norm(cell.value): cell.column
        for cell in inv[1]
        if cell.value is not None
    }

    if {
        "CREDENTIAL_ID",
        "CANONICAL_LINEAGE",
    }.issubset(inv_headers):

        for r in range(
            2,
            inv.max_row + 1,
        ):

            credential = norm(
                inv.cell(
                    r,
                    inv_headers["CREDENTIAL_ID"]
                ).value
            )

            lineage = norm(
                inv.cell(
                    r,
                    inv_headers["CANONICAL_LINEAGE"]
                ).value
            )

            text = (
                credential
                + " | "
                + lineage
            )

            if (
                "SAFETY_HEALTH" in text
                and "_AAS" not in text
                and "ASSOCIATE" not in text
            ):
                inv.cell(
                    r,
                    inv_headers[
                        "CANONICAL_LINEAGE"
                    ],
                ).value = TARGET

                inventory_updates += 1


# =============================================================================
# 4. BACKUPS
# =============================================================================

modeled_backup = MODELED_PATH.with_name(
    MODELED_PATH.stem
    + "_PRE_SHE_LINEAGE_FIX_"
    + STAMP
    + MODELED_PATH.suffix
)

crosswalk_backup = CROSSWALK_PATH.with_name(
    CROSSWALK_PATH.stem
    + "_PRE_SHE_LINEAGE_FIX_"
    + STAMP
    + CROSSWALK_PATH.suffix
)

shutil.copy2(
    MODELED_PATH,
    modeled_backup,
)

shutil.copy2(
    CROSSWALK_PATH,
    crosswalk_backup,
)


# =============================================================================
# 5. WRITE MODELED CSV ATOMICALLY
# =============================================================================

tmp_modeled = MODELED_PATH.with_suffix(
    ".tmp.csv"
)

modeled.to_csv(
    tmp_modeled,
    index=False,
)

os.replace(
    tmp_modeled,
    MODELED_PATH,
)


# =============================================================================
# 6. WRITE CROSSWALK WORKBOOK ATOMICALLY
# =============================================================================

tmp_crosswalk = CROSSWALK_PATH.with_name(
    CROSSWALK_PATH.stem
    + ".tmp.xlsx"
)

wb.save(tmp_crosswalk)

os.replace(
    tmp_crosswalk,
    CROSSWALK_PATH,
)


# =============================================================================
# 7. RELOAD + VALIDATE MODELED FILE
# =============================================================================

check = pd.read_csv(
    MODELED_PATH,
    dtype=str,
    low_memory=False,
).fillna("")

check_mask = check.apply(
    modeled_is_safety_certificate,
    axis=1,
)

remaining = check.loc[
    check_mask
    & check["canonical_lineage"].ne(
        TARGET
    )
]

if not remaining.empty:
    fail(
        "Post-write validation failed: "
        "Safety certificate rows remain "
        "outside target lineage."
    )

check_non_target = (
    check.loc[~check_mask]
    .copy()
    .reset_index(drop=True)
)

if not modeled_before_non_target.equals(
    check_non_target
):
    fail(
        "NON-TARGET MODELED ROWS CHANGED. "
        "Restore backup immediately."
    )


# =============================================================================
# 8. VERIFY AAS REMAINS SEPARATE
# =============================================================================

aas_after = check.loc[
    check["canonical_lineage"]
    .str.upper()
    .str.contains("SAFETY_HEALTH", na=False)
    & check["canonical_lineage"]
    .str.upper()
    .str.contains("AAS", na=False)
]

if len(aas_before) != len(aas_after):
    fail(
        "AAS row count changed. "
        "Certificate repair touched associate lineage."
    )

if (
    not aas_after.empty
    and
    aas_after["canonical_lineage"]
    .eq(TARGET)
    .any()
):
    fail(
        "AAS was incorrectly collapsed into "
        "certificate lineage."
    )


# =============================================================================
# 9. VERIFY CROSSWALK
# =============================================================================

wb_check = load_workbook(
    CROSSWALK_PATH,
    data_only=False,
)

ws_check = wb_check["Crosswalk Draft"]

check_headers = {
    norm(cell.value): cell.column
    for cell in ws_check[1]
    if cell.value is not None
}

for item in crosswalk_target_rows:

    r = item["excel_row"]

    actual = norm(
        ws_check.cell(
            r,
            check_headers[
                "CANONICAL_LINEAGE"
            ],
        ).value
    )

    recon = norm(
        ws_check.cell(
            r,
            check_headers[
                "RECONCILIATION_KEY"
            ],
        ).value
    )

    if actual != TARGET:
        fail(
            f"Crosswalk row {r} lineage "
            f"did not update: {actual}"
        )

    if recon != TARGET:
        fail(
            f"Crosswalk row {r} reconciliation "
            f"key did not update: {recon}"
        )


# =============================================================================
# 10. FINAL STUDENT-BLIND REPORT
# =============================================================================

after_targets = (
    check.loc[
        check_mask,
        [
            "catalog_year",
            "credential_id",
            "mechanical_lineage",
            "canonical_lineage",
        ],
    ]
    .drop_duplicates()
    .sort_values(
        [
            "catalog_year",
            "credential_id",
        ]
    )
)

print()
print("=" * 110)
print("AFTER REPAIR")
print("=" * 110)

print(
    after_targets.to_string(
        index=False
    )
)

print()
print("Modeled COMPLETE rows updated:",
      int(modeled_mask.sum()))

print(
    "Official crosswalk rows unified:",
    len(crosswalk_target_rows)
)

print(
    "Modeled Inventory rows unified:",
    inventory_updates
)

print()
print("Canonical certificate lineage:")
print(TARGET)

print()
print("AAS remains separate:")
print(
    sorted(
        set(
            aas_after[
                "canonical_lineage"
            ]
        )
    )
)

print()
print("BACKUPS")
print(modeled_backup)
print(crosswalk_backup)

print()
print("PASS — Safety/Health/Environment certificate lineage unified.")
