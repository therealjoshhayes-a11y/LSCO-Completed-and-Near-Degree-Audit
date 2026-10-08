from __future__ import annotations

from pathlib import Path
import hashlib

import pandas as pd
from docx import Document


BASELINE = Path(
    "config/core_bucket_lookup_multicatalog.csv"
)

SOURCE_2025 = Path(
    "data/raw/catalogs/2025-2026/"
    "2025-2026_Catalog.docx"
)

SOURCE_2026 = Path(
    "data/raw/catalogs/2026-2027/"
    "2026-2027_Catalog.docx"
)

OUTPUT = Path(
    "data/processed/catalogs/staging_2026/"
    "core_bucket_lookup_2026_staged.csv"
)

CORE_TABLE_INDEX = 13


def normalized_table_text(
    path: Path,
    table_index: int,
) -> str:
    if not path.exists():
        raise SystemExit(
            f"ABORT: missing source catalog: {path}"
        )

    doc = Document(path)

    if table_index >= len(doc.tables):
        raise SystemExit(
            f"ABORT: table {table_index} missing "
            f"from {path}"
        )

    table = doc.tables[table_index]

    lines = []

    for row in table.rows:
        cells = [
            " ".join(cell.text.split())
            for cell in row.cells
        ]

        lines.append(
            "\t".join(cells)
        )

    return "\n".join(lines)


def main() -> None:
    if not BASELINE.exists():
        raise SystemExit(
            f"ABORT: missing baseline: {BASELINE}"
        )

    baseline = pd.read_csv(
        BASELINE,
        dtype=str,
        keep_default_na=False,
        low_memory=False,
    )

    expected_columns = [
        "catalog_year",
        "bucket_name",
        "course_code",
        "source_file",
        "source_table_index",
        "membership_type",
    ]

    if list(baseline.columns) != expected_columns:
        raise SystemExit(
            "ABORT: baseline schema mismatch."
        )

    prior = baseline[
        baseline["catalog_year"].eq(
            "2025-2026"
        )
    ].copy()

    if len(prior) != 484:
        raise SystemExit(
            "ABORT: expected 484 2025-2026 "
            f"rows; found {len(prior)}."
        )

    if prior["bucket_name"].nunique() != 30:
        raise SystemExit(
            "ABORT: expected 30 2025-2026 "
            "bucket names."
        )

    if prior["course_code"].nunique() != 77:
        raise SystemExit(
            "ABORT: expected 77 2025-2026 "
            "core courses."
        )

    canonical = prior[
        prior["membership_type"].eq(
            "CANONICAL"
        )
    ].copy()

    aliases = prior[
        prior["membership_type"].eq(
            "CONTROLLED_ALIAS"
        )
    ].copy()

    if len(canonical) != 121:
        raise SystemExit(
            "ABORT: expected 121 canonical "
            f"memberships; found {len(canonical)}."
        )

    if len(aliases) != 363:
        raise SystemExit(
            "ABORT: expected 363 controlled "
            f"aliases; found {len(aliases)}."
        )

    if not canonical[
        "source_table_index"
    ].astype(str).str.strip().eq(
        str(CORE_TABLE_INDEX)
    ).all():
        raise SystemExit(
            "ABORT: canonical memberships "
            "must all cite source table 13."
        )

    if not aliases[
        "source_table_index"
    ].astype(str).str.strip().eq(
        ""
    ).all():
        raise SystemExit(
            "ABORT: controlled aliases must "
            "have blank source_table_index."
        )

    text_2025 = normalized_table_text(
        SOURCE_2025,
        CORE_TABLE_INDEX,
    )

    text_2026 = normalized_table_text(
        SOURCE_2026,
        CORE_TABLE_INDEX,
    )

    hash_2025 = hashlib.sha256(
        text_2025.encode("utf-8")
    ).hexdigest()

    hash_2026 = hashlib.sha256(
        text_2026.encode("utf-8")
    ).hexdigest()

    if text_2025 != text_2026:
        raise SystemExit(
            "ABORT: 2025 and 2026 core "
            "source tables are not identical."
        )

    staged = prior.copy()

    staged["catalog_year"] = (
        "2026-2027"
    )

    staged["source_file"] = str(
        SOURCE_2026
    )

    if staged.duplicated(
        subset=[
            "catalog_year",
            "bucket_name",
            "course_code",
            "membership_type",
        ]
    ).any():
        raise SystemExit(
            "ABORT: duplicate staged "
            "core memberships."
        )

    if len(staged) != 484:
        raise SystemExit(
            "ABORT: staged row count changed."
        )

    if staged[
        "bucket_name"
    ].nunique() != 30:
        raise SystemExit(
            "ABORT: staged bucket count changed."
        )

    if staged[
        "course_code"
    ].nunique() != 77:
        raise SystemExit(
            "ABORT: staged course count changed."
        )

    if staged[
        "membership_type"
    ].value_counts().to_dict() != {
        "CONTROLLED_ALIAS": 363,
        "CANONICAL": 121,
    }:
        raise SystemExit(
            "ABORT: membership-type counts "
            "changed unexpectedly."
        )

    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    staged.to_csv(
        OUTPUT,
        index=False,
    )

    print("=" * 80)
    print(
        "STAGED 2026 CORE BUCKET LOOKUP"
    )
    print("=" * 80)

    print(
        "SOURCE TABLE INDEX:",
        CORE_TABLE_INDEX,
    )
    print(
        "2025 TABLE SHA256:",
        hash_2025,
    )
    print(
        "2026 TABLE SHA256:",
        hash_2026,
    )
    print(
        "SOURCE TABLES IDENTICAL:",
        text_2025 == text_2026,
    )

    print()
    print("ROWS:", len(staged))
    print(
        "UNIQUE BUCKETS:",
        staged["bucket_name"].nunique(),
    )
    print(
        "UNIQUE COURSES:",
        staged["course_code"].nunique(),
    )
    print(
        "MEMBERSHIP TYPES:",
        staged[
            "membership_type"
        ].value_counts().to_dict(),
    )

    print()
    print("OUTPUT:", OUTPUT)
    print(
        "CANONICAL CONFIG MODIFIED: NO"
    )
    print(
        "PRODUCTION FILES MODIFIED: NO"
    )


if __name__ == "__main__":
    main()
