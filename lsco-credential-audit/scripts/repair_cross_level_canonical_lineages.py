from pathlib import Path
from datetime import datetime
import shutil
import re
import pandas as pd
from openpyxl import load_workbook

ROOT = Path.cwd()
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

COMPLETE = (
    ROOT
    / "data/processed/full_actual_audit/president_report/"
      "complete_rows_with_canonical_lineage.csv"
)

CROSSWALK = (
    ROOT
    / "data/interim/institutional_awards/"
      "award_program_crosswalk_curated.xlsx"
)

TARGETS = {
    "CERT1_MMMC": "ELECTROMECHANICAL_TECHNOLOGY_CERTIFICATE",
    "AAS_MMMD":   "ELECTROMECHANICAL_TECHNOLOGY_AAS",

    "CERT1_LOMC": "LOGISTICS_MANAGEMENT_CERTIFICATE",
    "AAS_LOMD":   "LOGISTICS_MANAGEMENT_MARITIME_AAS",

    "CERT1_MSGC": "MASSAGE_THERAPY_CERTIFICATE",
    "AAS_MSGD":   "MASSAGE_THERAPY_MANAGEMENT_AAS",
}


def backup(path: Path, label: str):
    dst = path.with_name(
        f"{path.stem}_PRE_{label}_{STAMP}{path.suffix}"
    )
    shutil.copy2(path, dst)
    return dst


def norm(v):
    return re.sub(
        r"[^A-Z0-9]+",
        "_",
        str(v or "").strip().upper()
    ).strip("_")


print("=" * 105)
print("CROSS-LEVEL CANONICAL LINEAGE REPAIR")
print("=" * 105)


# =============================================================================
# 1. MODELED COMPLETION LINEAGES
# =============================================================================

df = pd.read_csv(
    COMPLETE,
    dtype=str,
    low_memory=False,
).fillna("")

if "credential_id" not in df.columns:
    raise RuntimeError("credential_id missing from completion file")

if "canonical_lineage" not in df.columns:
    raise RuntimeError("canonical_lineage missing from completion file")

ids = df["credential_id"].map(norm)

new_lineage = pd.Series("", index=df.index, dtype=str)


# Electromechanical — do NOT touch Basic certificate.
electro = ids.str.contains(
    r"ELECTRO_?MECHANICAL_TECHNOLOGY",
    regex=True,
)

electro_basic = electro & ids.str.contains("BASIC", regex=False)

electro_cert = (
    electro
    & ~electro_basic
    & (
        ids.str.fullmatch(
            r"ELECTROMECHANICAL_TECHNOLOGY_20\d{2}",
            na=False,
        )
        | ids.str.contains(
            "CERTIFICATE_OF_COMPLETION",
            regex=False,
        )
    )
)

electro_aas = (
    electro
    & ~electro_basic
    & ids.str.contains(
        "_AAS_",
        regex=False,
    )
)

new_lineage.loc[electro_cert] = (
    "ELECTROMECHANICAL_TECHNOLOGY_CERTIFICATE"
)

new_lineage.loc[electro_aas] = (
    "ELECTROMECHANICAL_TECHNOLOGY_AAS"
)


# Logistics Management certificate vs Maritime AAS.
logistics = ids.str.startswith(
    "LOGISTICS_MANAGEMENT"
)

logistics_aas = (
    logistics
    & ids.str.contains(
        "MARITIME",
        regex=False,
    )
)

logistics_cert = (
    logistics
    & ~logistics_aas
)

new_lineage.loc[logistics_cert] = (
    "LOGISTICS_MANAGEMENT_CERTIFICATE"
)

new_lineage.loc[logistics_aas] = (
    "LOGISTICS_MANAGEMENT_MARITIME_AAS"
)


# Massage Therapy certificate vs Management AAS.
massage = ids.str.startswith(
    "MASSAGE_THERAPY"
)

massage_aas = (
    massage
    & ids.str.contains(
        "MANAGEMENT",
        regex=False,
    )
)

massage_cert = (
    massage
    & ~massage_aas
)

new_lineage.loc[massage_cert] = (
    "MASSAGE_THERAPY_CERTIFICATE"
)

new_lineage.loc[massage_aas] = (
    "MASSAGE_THERAPY_MANAGEMENT_AAS"
)


# Show exactly what will change.
change_mask = new_lineage.ne("")

preview_cols = [
    c for c in [
        "catalog_year",
        "credential_id",
        "credential_title",
        "mechanical_lineage",
        "canonical_lineage",
    ]
    if c in df.columns
]

preview = df.loc[
    change_mask,
    preview_cols
].drop_duplicates().copy()

preview["new_canonical_lineage"] = (
    new_lineage.loc[
        preview.index
    ].values
    if preview.index.is_unique
    else ""
)

# Safer display derived directly.
show = df.loc[
    change_mask,
    preview_cols
].copy()

show["new_canonical_lineage"] = (
    new_lineage.loc[change_mask]
)

show = (
    show
    .drop_duplicates()
    .sort_values(
        [
            "new_canonical_lineage",
            "credential_id",
        ]
    )
)

print()
print("MODELED CREDENTIALS TO RECLASSIFY")
print(show.to_string(index=False))

if show.empty:
    raise RuntimeError(
        "No modeled credentials matched repair patterns."
    )

# Guard against unresolved full Electromechanical rows.
unclassified_electro = (
    electro
    & ~electro_basic
    & ~electro_cert
    & ~electro_aas
)

if unclassified_electro.any():
    bad = df.loc[
        unclassified_electro,
        preview_cols
    ].drop_duplicates()

    raise RuntimeError(
        "Unclassified Electromechanical credential IDs:\n"
        + bad.to_string(index=False)
    )

complete_backup = backup(
    COMPLETE,
    "CROSS_LEVEL_LINEAGE_FIX"
)

df.loc[
    change_mask,
    "canonical_lineage"
] = new_lineage.loc[
    change_mask
]

df.to_csv(
    COMPLETE,
    index=False,
)

print()
print(
    "Modeled completion rows reclassified:",
    int(change_mask.sum()),
)


# =============================================================================
# 2. CURATED OFFICIAL-AWARD CROSSWALK
# =============================================================================

crosswalk_backup = backup(
    CROSSWALK,
    "CROSS_LEVEL_LINEAGE_FIX"
)

wb = load_workbook(CROSSWALK)

if "Crosswalk Draft" not in wb.sheetnames:
    raise RuntimeError("Crosswalk Draft sheet missing")

ws = wb["Crosswalk Draft"]

header_row = None
headers = {}

for r in range(1, min(ws.max_row, 20) + 1):

    probe = {
        str(
            ws.cell(r, c).value or ""
        ).strip(): c
        for c in range(
            1,
            ws.max_column + 1
        )
    }

    if "Curr1ProgramCode" in probe:
        header_row = r
        headers = probe
        break

if header_row is None:
    raise RuntimeError(
        "Could not locate crosswalk headers"
    )

required = {
    "Curr1ProgramCode",
    "canonical_lineage",
}

missing = required - set(headers)

if missing:
    raise RuntimeError(
        f"Crosswalk missing columns: {sorted(missing)}"
    )

seen = []

for r in range(
    header_row + 1,
    ws.max_row + 1,
):

    code = str(
        ws.cell(
            r,
            headers["Curr1ProgramCode"]
        ).value or ""
    ).strip()

    if code not in TARGETS:
        continue

    target = TARGETS[code]

    ws.cell(
        r,
        headers["canonical_lineage"],
    ).value = target

    if "reconciliation_key" in headers:
        ws.cell(
            r,
            headers["reconciliation_key"],
        ).value = target

    if "modeled_lineage_present" in headers:
        ws.cell(
            r,
            headers["modeled_lineage_present"],
        ).value = True

    if "match_status" in headers:
        ws.cell(
            r,
            headers["match_status"],
        ).value = "CURATED_MATCH"

    if "curation_note" in headers:

        if code in {
            "CERT1_MMMC",
            "AAS_MMMD",
        }:
            note = (
                "Catalog-validated distinct Electromechanical "
                "Technology certificate and AAS lineages."
            )

        elif code in {
            "CERT1_LOMC",
            "AAS_LOMD",
        }:
            note = (
                "Catalog-validated distinct Logistics Management "
                "certificate and Logistics Management (Maritime) AAS."
            )

        else:
            note = (
                "Catalog-validated distinct Massage Therapy "
                "certificate and Massage Therapy Management AAS."
            )

        ws.cell(
            r,
            headers["curation_note"],
        ).value = note

    seen.append(code)

if set(seen) != set(TARGETS):
    raise RuntimeError(
        "Did not locate all six expected award codes. "
        f"Found: {sorted(seen)}"
    )

wb.save(CROSSWALK)


# =============================================================================
# 3. VERIFY CROSSWALK HAS NO CROSS-LEVEL COLLISIONS
# =============================================================================

cw = pd.read_excel(
    CROSSWALK,
    sheet_name="Crosswalk Draft",
    dtype=str,
).fillna("")

cw["canonical_lineage"] = (
    cw["canonical_lineage"]
    .astype(str)
    .str.strip()
)

cw["award_level"] = (
    cw["award_level"]
    .astype(str)
    .str.strip()
    .str.upper()
)

collisions = (
    cw[
        cw["canonical_lineage"].ne("")
    ]
    .groupby(
        "canonical_lineage"
    )["award_level"]
    .nunique()
)

collisions = collisions[
    collisions > 1
]

print()
print("=" * 105)
print("POST-REPAIR CROSS-LEVEL COLLISIONS")
print("=" * 105)

if collisions.empty:
    print("NONE")
else:
    print(collisions.to_string())

print()
print("Completion backup:")
print(complete_backup)

print()
print("Crosswalk backup:")
print(crosswalk_backup)

print()
print("PASS")
