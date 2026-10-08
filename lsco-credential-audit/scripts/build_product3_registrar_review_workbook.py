from __future__ import annotations
"""Build the LSCO restricted Product 3 registrar review workbook locally.

Run from lsco-credential-audit root. Never upload generated student-level data.
Requires: pip install openpyxl
"""
import csv
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path.cwd()
REPORTING = ROOT / "data/processed/reporting"
PREFIX = "final_unawarded_complete_"
FINAL = "RESTRICTED_FINAL_unawarded_complete_awards.csv"
EXPIRED = "RESTRICTED_historical_complete_expired_catalog_candidates.csv"
REVIEWS = "RESTRICTED_unresolved_awardability_reviews.csv"
TRACE = "RESTRICTED_final_selection_trace.csv"
IDENTITY_SOURCE = ROOT / "data/raw/student_exports/banner_course_history_actual.csv"
GREEN = "#145A32"
TEAL = "#146C59"
PALE = "#EAF3ED"
GOLD = "#FFF0CA"
WHITE = "#FFFFFF"
INK = "#172B22"

def fail(message: str):
    raise RuntimeError(message)

def read_rows(path: Path):
    if not path.is_file(): fail(f"Missing local source: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))

def val(row, *names):
    for name in names:
        if str(row.get(name, "")).strip():
            return str(row[name]).strip()
    return ""

def chosen_output():
    dirs = sorted(
        (d for d in REPORTING.glob(PREFIX + "*") if d.is_dir() and (d / FINAL).exists()
         and (d / EXPIRED).exists() and (d / TRACE).exists()),
        key=lambda d: d.name,
        reverse=True
    )
    if not dirs: fail("No final selector output exists. Run select_final_unawarded_complete_awards_v3.py first.")
    return dirs[0]

def identity_lookup(needed):
    """Replicate archived August operational identity extraction: latest nonblank major by Term."""
    if not IDENTITY_SOURCE.is_file():
        fail(f"Identity enrichment is REQUIRED for shipment: {IDENTITY_SOURCE}")
    result = {}
    with IDENTITY_SOURCE.open("r", encoding="utf-8-sig", newline="") as stream:
        reader=csv.DictReader(stream)
        names=set(reader.fieldnames or [])
        def pick(candidates):
            for name in candidates:
                if name in names: return name
            fail("Banner identity source missing column from: "+str(candidates))
        idcol=pick(["StudenID","StudentID","student_id"])
        first=pick(["FirstName","first_name"])
        last=pick(["LastName","last_name"])
        major=pick(["StudentMajor","student_major","major"])
        term=pick(["Term","term","term_taken","term_sort"])
        middle=next((c for c in ("MiddleInitial","MiddleName","MI","middle_initial") if c in names), None)
        seq=0
        for row in reader:
            sid=(row.get(idcol) or "").strip()
            if sid not in needed: continue
            seq+=1
            raw=(row.get(term) or "").strip()
            try: sortkey=(1,int(float(raw)),seq)
            except ValueError: sortkey=(0,raw,seq)
            existing=result.setdefault(sid,{"first":"","last":"","middle":"","major":"","_major_sort":(-1,0,0),"_name_sort":(-1,0,0)})
            if sortkey>=existing["_name_sort"]:
                existing["first"]=(row.get(first) or "").strip() or existing["first"]
                existing["last"]=(row.get(last) or "").strip() or existing["last"]
                if middle: existing["middle"]=(row.get(middle) or "").strip() or existing["middle"]
                existing["_name_sort"]=sortkey
            m=(row.get(major) or "").strip()
            if m and sortkey>=existing["_major_sort"]:
                existing["major"]=m
                existing["_major_sort"]=sortkey
    missing=needed-set(result)
    if missing: fail(f"Missing Banner identity for {len(missing)} student(s); first 5 IDs: {sorted(missing)[:5]}")
    return result


def add_sheet(wb, name, headers, data, widths=None):
    sh=wb.create_sheet(name)
    sh.append(headers)
    for record in data: sh.append(record)
    sh.freeze_panes="A2"
    sh.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{len(data)+1}"
    for cell in sh[1]:
        cell.fill=PatternFill("solid",fgColor=GREEN.lstrip("#"))
        cell.font=Font(name="Aptos",size=10,bold=True,color="FFFFFF")
        cell.alignment=Alignment(vertical="center",wrap_text=True)
    sh.row_dimensions[1].height=32
    for j,width in enumerate(widths or [21]*len(headers),1):
        sh.column_dimensions[get_column_letter(j)].width=min(width,75)
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name="Aptos",size=10,color=INK.lstrip("#"))
            cell.alignment=Alignment(vertical="center")
            if row[0].row%2==0: cell.fill=PatternFill("solid",fgColor="F4F8F5")
    sh.sheet_view.showGridLines=False
    return sh

def main():
    source=chosen_output()
    awards=read_rows(source/FINAL)
    if not awards: fail("Final awards file is empty.")
    ids=[val(x,"student_id") for x in awards]
    if any(not x for x in ids): fail("Final recommendation missing student_id.")
    if len({(val(x,"student_id"),val(x,"candidate_lineage")) for x in awards}) != len(awards):
        fail("Student+lineage must be unique in the final recommendation list.")
    if any(val(x,"catalog_year")=="2021-2022" for x in awards):
        fail("Expired 2021-2022 recommendations detected: refusing workbook release.")
    if len(awards)!=367 or len(set(ids))!=326:
        fail(f"Unexpected release population: {len(awards)} awards and {len(set(ids))} students. Verify baseline first.")
    variance=sum(1 for x in awards if val(x,"conferral_temporal_eligibility")=="PASS_ANALYST_APPROVED_VARIANCE")
    # Fail closed if variance metadata never propagated; all selected 2026-27
    # rows must carry the explicit analyst-approved flag.
    modern=[x for x in awards if val(x,"catalog_year")=="2026-2027"]
    if any(val(x,"conferral_temporal_eligibility")!="PASS_ANALYST_APPROVED_VARIANCE" for x in modern):
        fail("2026-2027 records missing analyst-approved variance designation.")
    if len(modern)!=variance: fail("Variance count mismatch.")

    identities=identity_lookup(set(ids))
    wb=Workbook()
    overview=wb.active
    overview.title="CONTROL SUMMARY"
    catalog_counts=Counter(val(x,"catalog_year") for x in awards)
    summary=[
        ["LAMAR STATE COLLEGE ORANGE — UNAWARDED COMPLETE RECOMMENDATIONS",""],
        ["Product 3 • Registrar operational report • October 8, 2026",""],
        ["Recommended awards",len(awards)],
        ["Distinct students",len(set(ids))],
        ["Certificates",sum(val(x,"candidate_award_category")=="CERTIFICATE" for x in awards)],
        ["Associate degrees",sum(val(x,"candidate_award_category")=="ASSOCIATE" for x in awards)],
        ["2026–2027 analyst-approved variance",variance],
        ["2026–2027 expiration source","Confirmation pending; analyst-approved variance"],
        ["Scope","ONLY final recommended awards; no incomplete/review/expired populations"],
        ["Conferral","Subject to Registrar verification"],
        ["Source",str(source/FINAL)],
        ["Catalog","Recommendations"],
    ]
    summary.extend([[year,n] for year,n in sorted(catalog_counts.items())])
    for rec in summary: overview.append(rec)
    overview.column_dimensions["A"].width=70
    overview.column_dimensions["B"].width=90
    overview.freeze_panes="A3"
    overview.sheet_view.showGridLines=False
    for row in overview.iter_rows():
        for cell in row:
            cell.alignment=Alignment(vertical="center",wrap_text=True)
        if row[0].row in (1,2,12):
            for cell in row:
                cell.fill=PatternFill("solid",fgColor=GREEN.lstrip("#"))
                cell.font=Font(name="Aptos",size=11,bold=True,color="FFFFFF")
        else:
            row[0].fill=PatternFill("solid",fgColor=PALE.lstrip("#"))
            row[0].font=Font(name="Aptos",bold=True,color=INK.lstrip("#"))
        overview.row_dimensions[row[0].row].height=29

    # Restore the original operational report's compact student-facing columns.
    headers=["Banner ID","Last Name","First Name","Middle Initial","Declared Major",
             "Unawarded Credential","Credential Level","Catalog Year",
             "Modeled Completion Term","Registrar Disposition","Registrar Notes"]
    by_year={}
    missing_completion=0
    for award in awards:
        sid=val(award,"student_id")
        info=identities[sid]
        completion=val(award,"award_term_taken","modeled_completion_term",
                       "completion_term","completion_term_taken","award_term")
        if not completion: missing_completion+=1
        year=val(award,"catalog_year")
        by_year.setdefault(year,[]).append([
            sid,info["last"],info["first"],info["middle"],info["major"],
            val(award,"credential_title","awardability_credential_title"),
            val(award,"candidate_award_category"),year,completion,"",""
        ])
    for year in sorted(by_year):
        data=sorted(by_year[year],key=lambda row:(row[5].upper(),row[1].upper(),row[2].upper(),row[0]))
        sheet=add_sheet(wb,year,headers,data,[18,22,20,16,27,46,20,17,25,23,45])
        dropdown=DataValidation(type="list",formula1='"APPROVE,HOLD,REJECT,FOLLOW UP"',allow_blank=True)
        sheet.add_data_validation(dropdown)
        dropdown.add(f"J2:J{len(data)+1}")
    if sum(len(data) for data in by_year.values()) != len(awards):
        fail("Catalog tab population does not reconcile to recommended awards.")
    output=source/"RESTRICTED_LSCO_Product3_Registrar_Review_20261008.xlsx"
    wb.save(output)
    check=load_workbook(output,read_only=True,data_only=True)
    if set(check.sheetnames) != {"CONTROL SUMMARY",*by_year.keys()}:
        fail("Unexpected worksheet detected in operational release.")
    if sum(check[year].max_row-1 for year in by_year)!=len(awards):
        fail("Export verification failed: catalog tabs do not sum to final recommendations.")
    check.close()
    if not output.is_file() or output.stat().st_size<5000:
        fail("Workbook export did not produce a valid-size file.")
    print("PRODUCT 3 REGISTRAR REVIEW WORKBOOK CREATED")
    print(f"Workbook: {output}")
    print(f"Students: {len(set(ids))}; Recommendations: {len(awards)}")
    print(f"Identity and major enriched: {len(set(ids))} students")
    print(f"Modeled completion term unavailable in final source: {missing_completion} recommendations")
    print("Only final recommended awards exported; no expired or human-review records.")
    print(f"Analyst-approved 2026-27 variance recommendations: {variance}")
    print("Release status: PENDING REGISTRAR REVIEW; FERPA RESTRICTED.")

if __name__=="__main__":
    main()
