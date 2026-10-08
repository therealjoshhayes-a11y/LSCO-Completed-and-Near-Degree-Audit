"""Create the LSCO Registrar Review workbook locally from restricted CSVs.
Run from the project root: python scripts/build_lsco_registrar_workbook.py
Requires openpyxl (already a dependency in this project). No audit runs or source writes.
"""
from __future__ import annotations
import csv
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path.cwd()
REPORTS = ROOT / 'data' / 'processed' / 'reporting'
SOURCE = REPORTS / 'final_unawarded_complete_20261008_124402' / 'RESTRICTED_FINAL_unawarded_complete_awards.csv'
OFFICIAL = REPORTS / 'official_award_screening_final_20261008_092747' / 'RESTRICTED_canonical_official_awards_final_mapping.csv'
DEST = ROOT / 'outputs' / 'LSCO_Registrar_New_Award_Review_20261008.xlsx'
EXPECTED = 430
GREEN, DARK, LIGHT, GOLD, WHITE = '006747', '123D32', 'EAF3EE', 'C5A15B', 'FFFFFF'


def read(path):
    if not path.is_file():
        raise FileNotFoundError(f'Required local file not found: {path}')
    with path.open(encoding='utf-8-sig', newline='') as fp:
        return list(csv.DictReader(fp))


def val(row, *keys):
    for key in keys:
        s = str(row.get(key, '') or '').strip()
        if s:
            return s
    return ''


def safe(value):
    value = str(value or '')
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value


def band(ws, row, start, end, color=GREEN):
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=start, max_col=end):
        for c in cells:
            c.fill = PatternFill('solid', fgColor=color)
            c.font = Font(name='Aptos', size=10, bold=True, color=WHITE)
            c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def title(ws, text, last='G'):
    ws.merge_cells(f'A1:{last}2')
    c = ws['A1']
    c.value = text
    c.fill = PatternFill('solid', fgColor=DARK)
    c.font = Font(name='Aptos Display', size=17, bold=True, color=WHITE)
    c.alignment = Alignment(vertical='center')
    for row in ws.iter_rows(min_row=1,max_row=2,min_col=1,max_col=ord(last)-64):
        for cell in row: cell.fill = PatternFill('solid',fgColor=DARK)
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20
    ws.sheet_view.showGridLines = False


def cols(ws, widths):
    for c,width in widths.items(): ws.column_dimensions[c].width=width


def main():
    records = read(SOURCE)
    if len(records) != EXPECTED:
        raise ValueError(f'Expected {EXPECTED} final awards; got {len(records)}. No workbook produced.')
    keys=[(val(r,'student_id'),val(r,'catalog_year'),val(r,'credential_id')) for r in records]
    if any(not all(k) for k in keys) or len(keys)!=len(set(keys)):
        raise ValueError('Duplicate or incomplete selected award keys.')
    if any(val(r,'candidate_award_category') not in ('ASSOCIATE','CERTIFICATE') for r in records):
        raise ValueError('Unexpected award category.')
    by_student=defaultdict(list)
    for r in records: by_student[val(r,'student_id')].append(r)
    official_by_student=defaultdict(list)
    official_keys=set()
    for r in read(OFFICIAL):
        sid, degree, major, term=(val(r,'ID','student_id','StudentID'),val(r,'DegreeCode'),val(r,'Major1Code'),val(r,'StudGradTerm'))
        if not (sid and degree and major): continue
        key=(sid,degree,major,term)
        if key in official_keys: continue
        official_keys.add(key)
        official_by_student[sid].append((degree,major,term,val(r,'suppression_family')))

    years=sorted({val(r,'catalog_year') for r in records})
    wb=Workbook()
    summary=wb.active
    summary.title='SUMMARY'
    title(summary,'LSCO  |  GRADUATION AUDIT — REGISTRAR REVIEW','H')
    summary.merge_cells('A3:H3')
    summary['A3']='October 8, 2026 governed selection  •  Recommendations, not conferred awards'
    summary['A3'].fill=PatternFill('solid',fgColor=LIGHT)
    summary['A3'].font=Font(name='Aptos',size=10,color=DARK)
    summary.row_dimensions[3].height=24
    metrics=[('A5','NEW RECOMMENDATIONS','B5',len(records)),('D5','DISTINCT STUDENTS','E5',len(by_student)),('G5','PRIOR OFFICIAL AWARDS','H5',len(official_keys))]
    for la,label,va,num in metrics:
        summary[la]=label;summary[va]=num
        for cell in (summary[la],summary[va]):
            cell.fill=PatternFill('solid',fgColor=GREEN)
            cell.font=Font(name='Aptos',color=WHITE,bold=True,size=11)
        summary[va].alignment=Alignment(horizontal='center')
    summary.row_dimensions[5].height=27
    summary.append([])
    for i,head in enumerate(('Catalog year','New associates','New certificates','New total','Distinct students'),1): summary.cell(8,i,head)
    band(summary,8,1,5)
    yr_rows=[]
    for y in years:
        group=[r for r in records if val(r,'catalog_year')==y]
        yr_rows.append((y,sum(val(r,'candidate_award_category')=='ASSOCIATE' for r in group),sum(val(r,'candidate_award_category')=='CERTIFICATE' for r in group),len(group),len({val(r,'student_id') for r in group})))
    for i,entry in enumerate(yr_rows,9):
        for j,v in enumerate(entry,1): summary.cell(i,j,v)
    total_row=9+len(years)
    for j,v in enumerate(('TOTAL',sum(r[1] for r in yr_rows),sum(r[2] for r in yr_rows),len(records),len(by_student)),1):
        c=summary.cell(total_row,j,v);c.fill=PatternFill('solid',fgColor=LIGHT);c.font=Font(bold=True,color=DARK)
    major_header=total_row+3
    for j,h in enumerate(('Catalog year','Award category','Credential / major','Recommended awards'),1): summary.cell(major_header,j,h)
    band(summary,major_header,1,4)
    majors=Counter((val(r,'catalog_year'),val(r,'candidate_award_category'),val(r,'credential_title')) for r in records)
    for i,((y,cat,name),count) in enumerate(sorted(majors.items(),key=lambda x:(x[0][0],x[0][1],-x[1],x[0][2])),major_header+1):
        for j,v in enumerate((y,cat,safe(name),count),1):summary.cell(i,j,v)
    end=major_header+len(majors)+2
    summary.merge_cells(start_row=end,start_column=1,end_row=end,end_column=8)
    summary.cell(end,1,'Prior official awards count distinct student + degree code + major code + graduation term in the canonical official ledger (institution-wide, not only these students).')
    summary.cell(end,1).alignment=Alignment(wrap_text=True,vertical='center')
    summary.cell(end,1).fill=PatternFill('solid',fgColor=LIGHT)
    summary.row_dimensions[end].height=34
    cols(summary,{'A':19,'B':23,'C':46,'D':21,'E':21,'F':4,'G':27,'H':20})
    summary.freeze_panes='A9'

    for year in years:
        ws=wb.create_sheet(year)
        title(ws,f'LSCO  |  {year}  •  NEW AWARD RECOMMENDATIONS')
        group=[r for r in records if val(r,'catalog_year')==year]
        ws.merge_cells('A3:G3')
        ws['A3']=f'{len(group)} selected awards  •  {len({val(r,"student_id") for r in group})} distinct students  •  Not yet conferred'
        ws['A3'].fill=PatternFill('solid',fgColor=LIGHT)
        ws['A3'].font=Font(name='Aptos',color=DARK,size=10)
        ws.row_dimensions[3].height=24
        heads=['Student ID','Catalog year','Recommended credential','Type','Other recommended credentials (all catalogs)','Previously conferred credentials (official)','Canonical lineage']
        for j,h in enumerate(heads,1):ws.cell(4,j,h)
        band(ws,4,1,7)
        for i,r in enumerate(sorted(group,key=lambda r:(val(r,'student_id'),val(r,'candidate_award_category'),val(r,'credential_title'))),5):
            sid=val(r,'student_id')
            others=[f'{val(x,"credential_title")} [{val(x,"catalog_year")}; {val(x,"candidate_award_category")}]' for x in by_student[sid] if x is not r]
            official=[f'{degree}/{major} [{term}]'+(f' — {lineage}' if lineage else '') for degree,major,term,lineage in sorted(official_by_student.get(sid,[]),key=lambda x:(x[2],x[1]))]
            values=[safe(sid),year,safe(val(r,'credential_title')),val(r,'candidate_award_category'),safe('; '.join(others) or '—'),safe('; '.join(official) or 'None in ledger'),safe(val(r,'candidate_lineage','credential_lineage'))]
            for j,v in enumerate(values,1):
                c=ws.cell(i,j,v)
                c.font=Font(name='Aptos',size=10,color=DARK)
                c.alignment=Alignment(vertical='top',wrap_text=j in (3,5,6,7))
                if i%2==0:c.fill=PatternFill('solid',fgColor='F3F7F4')
            ws.row_dimensions[i].height=42 if len(others)+len(official)>2 else 29
        cols(ws,{'A':17,'B':16,'C':40,'D':17,'E':68,'F':78,'G':34})
        ws.freeze_panes='C5'
        ws.auto_filter.ref=f'A4:G{max(5,4+len(group))}'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape'
        ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1
        ws.sheet_properties.tabColor=GREEN
    summary.sheet_properties.tabColor=GOLD
    DEST.parent.mkdir(parents=True,exist_ok=True)
    wb.save(DEST)
    print('SUCCESS:',DEST)
    print(f'NEW RECOMMENDATIONS: {len(records)}; UNIQUE STUDENTS: {len(by_student)}; PRIOR OFFICIAL AWARDS: {len(official_keys)}')
    for y,a,c,total,students in yr_rows: print(f'{y}: {total} awards (associate {a}, certificate {c}; {students} students)')

if __name__=='__main__': main()
