from pathlib import Path
from datetime import datetime
import shutil
import pandas as pd
from openpyxl import load_workbook

ROOT = Path.cwd()
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

COMPLETE = (
    ROOT
    / "data/processed/full_actual_audit/president_report"
    / "complete_rows_with_canonical_lineage.csv"
)

CROSSWALK = (
    ROOT
    / "data/interim/institutional_awards"
    / "award_program_crosswalk_curated.xlsx"
)

GOV = (
    ROOT
    / "data/interim/institutional_awards"
    / "governed_lineage_relationships_final.csv"
)

LEGACY_MAP = {
    "TEACHING_T038": "TEACHING_AAT1",
    "TEACHING_T074": "TEACHING_AAT1",
    "TEACHING_T039": "TEACHING_AAT2",
    "TEACHING_T075": "TEACHING_AAT2",
}

LEGACY_IDS = set(LEGACY_MAP)


def backup(path: Path, label: str):
    dst = path.with_name(
        f"{path.stem}_PRE_{label}_{STAMP}{path.suffix}"
    )
    shutil.copy2(path, dst)
    return dst


print("=" * 110)
print("TEACHING CANONICAL IDENTITY REPAIR")
print("=" * 110)


# =============================================================================
# 1. MODELED COMPLETION IDENTITIES
# =============================================================================

df = pd.read_csv(
    COMPLETE,
    dtype=str,
    low_memory=False,
).fillna("")

required = {
    "credential_id",
    "canonical_lineage",
}

missing = required - set(df.columns)

if missing:
    raise RuntimeError(
        f"Completion file missing columns: {sorted(missing)}"
    )

before = df[
    df["canonical_lineage"].isin(LEGACY_IDS)
].copy()

print()
print("LEGACY MODELED TEACHING ROWS")
print("-" * 110)

show_cols = [
    c for c in [
        "catalog_year",
        "credential_id",
        "mechanical_lineage",
        "canonical_lineage",
    ]
    if c in before.columns
]

if before.empty:
    print("NONE")
else:
    print(
        before[show_cols]
        .drop_duplicates()
        .sort_values("credential_id")
        .to_string(index=False)
    )

complete_backup = backup(
    COMPLETE,
    "TEACHING_IDENTITY_FIX"
)

for old, new in LEGACY_MAP.items():
    mask = df["canonical_lineage"].eq(old)
    df.loc[
        mask,
        "canonical_lineage"
    ] = new


# Future-proof current artifact if a 2025 AAT completion is present.
# Mechanical parser defect remains documented separately.
teacher_aat_mask = (
    df["credential_id"]
    .astype(str)
    .str.startswith(
        "TEACHER_EDUCATION_AAS_2025"
    )
)

df.loc[
    teacher_aat_mask,
    "canonical_lineage"
] = "TEACHER_EDUCATION_AAT"

df.to_csv(
    COMPLETE,
    index=False,
)


# =============================================================================
# 2. OFFICIAL AWARD CROSSWALK
# =============================================================================

crosswalk_backup = backup(
    CROSSWALK,
    "TEACHING_IDENTITY_FIX"
)

wb = load_workbook(CROSSWALK)

if "Crosswalk Draft" not in wb.sheetnames:
    raise RuntimeError(
        "Crosswalk Draft sheet missing."
    )

ws = wb["Crosswalk Draft"]

header_row = None
headers = {}

for r in range(1, min(ws.max_row, 20) + 1):

    probe = {
        str(ws.cell(r, c).value or "").strip(): c
        for c in range(1, ws.max_column + 1)
    }

    if "Curr1ProgramCode" in probe:
        header_row = r
        headers = probe
        break

if header_row is None:
    raise RuntimeError(
        "Could not locate crosswalk header row."
    )

needed = {
    "Curr1ProgramCode",
    "canonical_lineage",
}

missing = needed - set(headers)

if missing:
    raise RuntimeError(
        f"Crosswalk missing columns: {sorted(missing)}"
    )

expected = {
    "AAT_AAT1": "TEACHING_AAT1",
    "AAT_AAT2": "TEACHING_AAT2",
    "AAT_AATP": "TEACHER_EDUCATION_AAT",
    "CERT1_TEPC":
        "TEACHER_EDUCATION_CERTIFICATE_OF_COMPLETION",
}

seen = {}

for r in range(
    header_row + 1,
    ws.max_row + 1,
):

    code = str(
        ws.cell(
            r,
            headers["Curr1ProgramCode"]
        ).value or ""
    ).strip()

    if code not in expected:
        continue

    target = expected[code]

    current = str(
        ws.cell(
            r,
            headers["canonical_lineage"]
        ).value or ""
    ).strip()

    seen[code] = current

    # Only AATP needs reassignment.
    if code == "AAT_AATP":

        ws.cell(
            r,
            headers["canonical_lineage"]
        ).value = target

        if "reconciliation_key" in headers:
            ws.cell(
                r,
                headers["reconciliation_key"]
            ).value = target

        if "modeled_lineage_present" in headers:
            ws.cell(
                r,
                headers["modeled_lineage_present"]
            ).value = True

        if "match_status" in headers:
            ws.cell(
                r,
                headers["match_status"]
            ).value = "CURATED_MATCH"

        if "curation_note" in headers:
            ws.cell(
                r,
                headers["curation_note"]
            ).value = (
                "Catalog- and award-term-validated 2025-2026 "
                "Teacher Education Associate of Arts in Teaching. "
                "AATP awards occur beginning Fall 2025 and map to "
                "the consolidated Teacher Education AAT lineage."
            )

if set(seen) != set(expected):
    raise RuntimeError(
        "Expected all four Teaching award codes. "
        f"Found: {sorted(seen)}"
    )

# Guard existing established mappings.
if seen["AAT_AAT1"] != "TEACHING_AAT1":
    raise RuntimeError(
        "AAT_AAT1 was not mapped to TEACHING_AAT1."
    )

if seen["AAT_AAT2"] != "TEACHING_AAT2":
    raise RuntimeError(
        "AAT_AAT2 was not mapped to TEACHING_AAT2."
    )

if (
    seen["CERT1_TEPC"]
    != "TEACHER_EDUCATION_CERTIFICATE_OF_COMPLETION"
):
    raise RuntimeError(
        "CERT1_TEPC mapping unexpected."
    )

wb.save(CROSSWALK)


# =============================================================================
# 3. RETIRE GOVERNANCE USING INVALID T0xx CANONICAL IDENTITIES
# =============================================================================

gov_backup = backup(
    GOV,
    "TEACHING_IDENTITY_FIX"
)

g = pd.read_csv(
    GOV,
    dtype=str,
).fillna("")

required = {
    "detected_lineage",
    "institutional_lineage",
}

missing = required - set(g.columns)

if missing:
    raise RuntimeError(
        f"Governance file missing columns: {sorted(missing)}"
    )

legacy_mask = (
    g["detected_lineage"].isin(LEGACY_IDS)
    | g["institutional_lineage"].isin(LEGACY_IDS)
)

retired = g.loc[
    legacy_mask
].copy()

print()
print("LEGACY GOVERNANCE ROWS RETIRED:", len(retired))

if not retired.empty:

    audit_path = GOV.with_name(
        f"retired_teaching_T0xx_governance_{STAMP}.csv"
    )

    retired.to_csv(
        audit_path,
        index=False,
    )

    print("Retired-row audit:")
    print(audit_path)

g = g.loc[
    ~legacy_mask
].copy()

g.to_csv(
    GOV,
    index=False,
)


# =============================================================================
# 4. QA
# =============================================================================

check_complete = pd.read_csv(
    COMPLETE,
    dtype=str,
    low_memory=False,
).fillna("")

bad_modeled = check_complete[
    check_complete["canonical_lineage"].isin(
        LEGACY_IDS
    )
]

if not bad_modeled.empty:
    raise RuntimeError(
        "Legacy T0xx canonical lineages remain "
        "in completion artifact."
    )

check_cw = pd.read_excel(
    CROSSWALK,
    sheet_name="Crosswalk Draft",
    dtype=str,
).fillna("")

aatp = check_cw[
    check_cw["Curr1ProgramCode"].eq(
        "AAT_AATP"
    )
]

if len(aatp) != 1:
    raise RuntimeError(
        f"Expected one AAT_AATP row; found {len(aatp)}"
    )

if (
    aatp.iloc[0]["canonical_lineage"]
    != "TEACHER_EDUCATION_AAT"
):
    raise RuntimeError(
        "AATP remap failed."
    )

check_g = pd.read_csv(
    GOV,
    dtype=str,
).fillna("")

bad_gov = check_g[
    check_g["detected_lineage"].isin(
        LEGACY_IDS
    )
    | check_g["institutional_lineage"].isin(
        LEGACY_IDS
    )
]

if not bad_gov.empty:
    raise RuntimeError(
        "Legacy T0xx identities remain in active governance."
    )


print()
print("=" * 110)
print("FINAL TEACHING IDENTITY")
print("=" * 110)

print(
    "T038 / T074 -> TEACHING_AAT1"
)
print(
    "T039 / T075 -> TEACHING_AAT2"
)
print(
    "AAT_AATP     -> TEACHER_EDUCATION_AAT"
)
print(
    "CERT1_TEPC   -> "
    "TEACHER_EDUCATION_CERTIFICATE_OF_COMPLETION"
)

print()
print("Completion backup:")
print(complete_backup)

print()
print("Crosswalk backup:")
print(crosswalk_backup)

print()
print("Governance backup:")
print(gov_backup)

print()
print("PASS")
